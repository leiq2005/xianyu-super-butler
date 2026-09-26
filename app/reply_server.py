from fastapi import FastAPI, HTTPException, Depends, UploadFile, File, Form, Body, Query
from fastapi.staticfiles import StaticFiles
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from pydantic import BaseModel, Field
from typing import List, Tuple, Optional, Dict, Any
from pathlib import Path
import secrets
import time
import json
import os
import re
import sys
import pandas as pd
import io
import asyncio
import sqlite3
from collections import defaultdict, OrderedDict

# 同 app/config.py：打包成 exe 后以 exe 所在目录作为项目根。
PROJECT_ROOT = (
    Path(sys.executable).resolve().parent
    if getattr(sys, 'frozen', False)
    else Path(__file__).resolve().parent.parent
)

from app import cookie_manager
from app.db_manager import db_manager
from app.product_automation import ProductAutomationService
from app.file_log_collector import setup_file_logging, get_file_log_collector
from app.ai_reply_engine import ai_reply_engine
from app.routers.delivery_block import create_delivery_block_router
from utils.qr_login import qr_login_manager
from utils.xianyu_utils import trans_cookies
from utils.image_utils import image_manager
from utils.order_status_rules import (
    get_order_status,
    is_stable_order_status,
    normalize_order_status,
)

from loguru import logger

product_automation = ProductAutomationService(db_manager)

# 刮刮乐远程控制路由
try:
    from app.api_captcha_remote import router as captcha_router
    CAPTCHA_ROUTER_AVAILABLE = True
except ImportError:
    logger.warning("⚠️ api_captcha_remote 未找到，刮刮乐远程控制功能不可用")
    CAPTCHA_ROUTER_AVAILABLE = False

# 关键字文件路径
KEYWORDS_FILE = PROJECT_ROOT / "回复关键字.txt"

# 简单的用户认证配置
ADMIN_USERNAME = "admin"
DEFAULT_ADMIN_PASSWORD = "admin123"  # 系统初始化时的默认密码
SESSION_TOKENS = {}  # 存储会话token: {token: {'user_id': int, 'username': str, 'timestamp': float}}
TOKEN_EXPIRE_TIME = 24 * 60 * 60  # token过期时间：24小时

# HTTP Bearer认证
security = HTTPBearer(auto_error=False)

# 扫码登录检查锁 - 防止并发处理同一个session
qr_check_locks = defaultdict(lambda: asyncio.Lock())
qr_check_processed = {}  # 记录已处理的session: {session_id: {'processed': bool, 'timestamp': float}}
qr_check_tasks = {}  # 保持后台任务引用，避免扫码后的Cookie准备阻塞状态查询

# 账号密码登录会话管理
password_login_sessions = {}  # {session_id: {'account_id': str, 'account': str, 'password': str, 'show_browser': bool, 'status': str, 'verification_url': str, 'qr_code_url': str, 'slider_instance': object, 'task': asyncio.Task, 'timestamp': float}}
password_login_locks = defaultdict(lambda: asyncio.Lock())

# 不再需要单独的密码初始化，由数据库初始化时处理


def cleanup_qr_check_records():
    """清理过期的扫码检查记录"""
    current_time = time.time()
    expired_sessions = []

    for session_id, record in qr_check_processed.items():
        # 清理超过1小时的记录
        if current_time - record['timestamp'] > 3600:
            expired_sessions.append(session_id)

    for session_id in expired_sessions:
        if session_id in qr_check_processed:
            del qr_check_processed[session_id]
        if session_id in qr_check_locks:
            del qr_check_locks[session_id]
        task = qr_check_tasks.pop(session_id, None)
        if task and not task.done():
            task.cancel()


def load_keywords() -> List[Tuple[str, str]]:
    """读取关键字→回复映射表

    文件格式支持：
        关键字<空格/制表符/冒号>回复内容
    忽略空行和以 # 开头的注释行
    """
    mapping: List[Tuple[str, str]] = []
    if not KEYWORDS_FILE.exists():
        return mapping

    with KEYWORDS_FILE.open('r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith('#'):
                continue
            # 尝试用\t、空格、冒号分隔
            if '\t' in line:
                key, reply = line.split('\t', 1)
            elif ' ' in line:
                key, reply = line.split(' ', 1)
            elif ':' in line:
                key, reply = line.split(':', 1)
            else:
                # 无法解析的行，跳过
                continue
            mapping.append((key.strip(), reply.strip()))
    return mapping


KEYWORDS_MAPPING = load_keywords()


# 认证相关模型
class LoginRequest(BaseModel):
    username: Optional[str] = None
    password: Optional[str] = None
    email: Optional[str] = None
    verification_code: Optional[str] = None


class LoginResponse(BaseModel):
    success: bool
    token: Optional[str] = None
    message: str
    user_id: Optional[int] = None
    username: Optional[str] = None
    is_admin: Optional[bool] = None


class ChangePasswordRequest(BaseModel):
    current_password: str
    new_password: str


class RegisterRequest(BaseModel):
    username: str
    email: str
    password: str
    # 关闭邮箱验证时前端不会带这个字段，设为可选避免直接 422
    verification_code: Optional[str] = None


class RegisterResponse(BaseModel):
    success: bool
    message: str


class SendCodeRequest(BaseModel):
    email: str
    session_id: Optional[str] = None
    type: Optional[str] = 'register'  # 'register' 或 'login'


class SendCodeResponse(BaseModel):
    success: bool
    message: str


class CaptchaRequest(BaseModel):
    session_id: str


class CaptchaResponse(BaseModel):
    success: bool
    captcha_image: str
    session_id: str
    message: str


class VerifyCaptchaRequest(BaseModel):
    session_id: str
    captcha_code: str


class VerifyCaptchaResponse(BaseModel):
    success: bool
    message: str


def generate_token() -> str:
    """生成随机token"""
    return secrets.token_urlsafe(32)


def verify_token(credentials: Optional[HTTPAuthorizationCredentials] = Depends(security)) -> Optional[Dict[str, Any]]:
    """验证token并返回用户信息"""
    if not credentials:
        return None

    token = credentials.credentials
    if token not in SESSION_TOKENS:
        return None

    token_data = SESSION_TOKENS[token]

    # 检查token是否过期
    if time.time() - token_data['timestamp'] > TOKEN_EXPIRE_TIME:
        del SESSION_TOKENS[token]
        return None

    return token_data


def verify_admin_token(credentials: Optional[HTTPAuthorizationCredentials] = Depends(security)) -> Dict[str, Any]:
    """验证管理员token"""
    user_info = verify_token(credentials)
    if not user_info:
        raise HTTPException(status_code=401, detail="未授权访问")

    # 检查是否是管理员
    if user_info['username'] != ADMIN_USERNAME:
        raise HTTPException(status_code=403, detail="需要管理员权限")

    return user_info


def require_auth(user_info: Optional[Dict[str, Any]] = Depends(verify_token)):
    """需要认证的依赖，返回用户信息"""
    if not user_info:
        raise HTTPException(status_code=401, detail="未授权访问")
    return user_info


def get_current_user(user_info: Dict[str, Any] = Depends(require_auth)) -> Dict[str, Any]:
    """获取当前登录用户信息"""
    return user_info


def get_current_user_optional(user_info: Optional[Dict[str, Any]] = Depends(verify_token)) -> Optional[Dict[str, Any]]:
    """获取当前用户信息（可选，不强制要求登录）"""
    return user_info


def get_user_log_prefix(user_info: Dict[str, Any] = None) -> str:
    """获取用户日志前缀"""
    if user_info:
        return f"【{user_info['username']}#{user_info['user_id']}】"
    return "【系统】"


def require_admin(current_user: Dict[str, Any] = Depends(get_current_user)) -> Dict[str, Any]:
    """要求管理员权限"""
    if current_user['username'] != 'admin':
        raise HTTPException(status_code=403, detail="需要管理员权限")
    return current_user


def log_with_user(level: str, message: str, user_info: Dict[str, Any] = None):
    """带用户信息的日志记录"""
    prefix = get_user_log_prefix(user_info)
    full_message = f"{prefix} {message}"

    if level.lower() == 'info':
        logger.info(full_message)
    elif level.lower() == 'error':
        logger.error(full_message)
    elif level.lower() == 'warning':
        logger.warning(full_message)
    elif level.lower() == 'debug':
        logger.debug(full_message)
    else:
        logger.info(full_message)


def match_reply(cookie_id: str, message: str) -> Optional[str]:
    """根据 cookie_id 及消息内容匹配回复
    只有启用的账号才会匹配关键字回复
    """
    mgr = cookie_manager.manager
    if mgr is None:
        return None

    # 检查账号是否启用
    if not mgr.get_cookie_status(cookie_id):
        return None  # 禁用的账号不参与自动回复

    # 优先账号级关键字
    if mgr.get_keywords(cookie_id):
        for k, r in mgr.get_keywords(cookie_id):
            if k in message:
                return r

    # 全局关键字
    for k, r in KEYWORDS_MAPPING:
        if k in message:
            return r
    return None


class RequestModel(BaseModel):
    cookie_id: str
    msg_time: str
    user_url: str
    send_user_id: str
    send_user_name: str
    item_id: str
    send_message: str
    chat_id: str


class ResponseData(BaseModel):
    send_msg: str


class ResponseModel(BaseModel):
    code: int
    data: ResponseData


app = FastAPI(
    title="Xianyu Auto Reply API",
    version="2.5.0",
    description="闲鱼自动回复系统API",
    docs_url="/docs",
    redoc_url="/redoc"
)

# 注册刮刮乐远程控制路由
if CAPTCHA_ROUTER_AVAILABLE:
    app.include_router(captcha_router)
    logger.info("✅ 已注册刮刮乐远程控制路由: /api/captcha")
else:
    logger.warning("⚠️ 刮刮乐远程控制路由未注册")

app.include_router(create_delivery_block_router(get_current_user, db_manager))
logger.info("已注册发货拦截规则路由")

# 初始化文件日志收集器
setup_file_logging()

# 添加一条测试日志
from loguru import logger
logger.info("Web服务器启动，文件日志收集器已初始化")

# 添加请求日志中间件
@app.middleware("http")
async def log_requests(request, call_next):
    start_time = time.time()

    logger.info(f"🌐 API请求: {request.method} {request.url.path}")

    response = await call_next(request)

    process_time = time.time() - start_time
    logger.info(f"✅ API响应: {request.method} {request.url.path} - {response.status_code} ({process_time:.3f}s)")

    return response

# 提供前端静态文件
import os
static_dir = str(PROJECT_ROOT / 'static')
if not os.path.exists(static_dir):
    os.makedirs(static_dir, exist_ok=True)

# 挂载静态文件目录
app.mount('/static', StaticFiles(directory=static_dir), name='static')

# 挂载 /assets 路径，指向 static/assets 目录
# 这样访问 /assets/xxx.js 时会正确映射到 static_dir/assets/xxx.js
assets_dir = os.path.join(static_dir, 'assets')
app.mount('/assets', StaticFiles(directory=assets_dir), name='assets')

# 确保图片上传目录存在
uploads_dir = os.path.join(static_dir, 'uploads', 'images')
if not os.path.exists(uploads_dir):
    os.makedirs(uploads_dir, exist_ok=True)
    logger.info(f"创建图片上传目录: {uploads_dir}")

# 健康检查端点
@app.get('/health')
async def health_check():
    """健康检查端点，用于Docker健康检查和负载均衡器"""
    try:
        # 检查Cookie管理器状态
        manager_status = "ok" if cookie_manager.manager is not None else "error"

        # 检查数据库连接
        from app.db_manager import db_manager
        try:
            db_manager.get_all_cookies()
            db_status = "ok"
        except Exception:
            db_status = "error"

        # 获取系统状态
        import psutil
        cpu_percent = psutil.cpu_percent(interval=None)
        memory_info = psutil.virtual_memory()

        status = {
            "status": "healthy" if manager_status == "ok" and db_status == "ok" else "unhealthy",
            "timestamp": time.time(),
            "services": {
                "cookie_manager": manager_status,
                "database": db_status
            },
            "system": {
                "cpu_percent": cpu_percent,
                "memory_percent": memory_info.percent,
                "memory_available": memory_info.available
            }
        }

        if status["status"] == "unhealthy":
            return JSONResponse(status_code=503, content=status)

        return status

    except Exception as e:
        return JSONResponse(
            status_code=503,
            content={
                "status": "unhealthy",
                "timestamp": time.time(),
                "error": str(e)
            }
        )


# 服务 React 前端 SPA - 所有前端路由都返回 index.html
async def serve_frontend():
    """服务 React 前端 SPA"""
    index_path = os.path.join(static_dir, 'index.html')
    if os.path.exists(index_path):
        with open(index_path, 'r', encoding='utf-8') as f:
            return HTMLResponse(f.read())
    else:
        return HTMLResponse('<h3>Frontend not found. Please build the frontend first.</h3>')

@app.get('/', response_class=HTMLResponse)
async def root():
    return await serve_frontend()


# 登录页面路由 - 重定向到 React 前端
@app.get('/login.html', response_class=HTMLResponse)
async def login_page():
    return await serve_frontend()

@app.get('/login', response_class=HTMLResponse)
async def login_route():
    return await serve_frontend()


# 注册页面路由
@app.get('/register.html', response_class=HTMLResponse)
async def register_page():
    # 检查注册是否开启
    from app.db_manager import db_manager
    registration_enabled = db_manager.get_system_setting('registration_enabled')
    if registration_enabled != 'true':
        return HTMLResponse('''
        <!DOCTYPE html>
        <html>
        <head>
            <title>注册已关闭</title>
            <meta charset="utf-8">
            <style>
                body { font-family: Arial, sans-serif; text-align: center; padding: 50px; }
                .message { color: #666; font-size: 18px; }
                .back-link { margin-top: 20px; }
                .back-link a { color: #007bff; text-decoration: none; }
            </style>
        </head>
        <body>
            <h2>🚫 注册功能已关闭</h2>
            <p class="message">系统管理员已关闭用户注册功能</p>
            <div class="back-link">
                <a href="/">← 返回首页</a>
            </div>
        </body>
        </html>
        ''', status_code=403)

    return await serve_frontend()

@app.get('/register', response_class=HTMLResponse)
async def register_route():
    return await serve_frontend()


# 注意：不要在这里定义 /admin 或 /admin/{path} 路由
# 因为后端有 /admin/users, /admin/logs 等 API 路由
# 前端 SPA 通过根路由 / 加载，由 React Router 处理客户端路由
# 文件末尾的 catch-all 路由会处理前端页面的直接访问



# 登录接口
@app.post('/login')
async def login(request: LoginRequest):
    from app.db_manager import db_manager

    # 判断登录方式
    if request.username and request.password:
        # 用户名/密码登录
        logger.info(f"【{request.username}】尝试用户名登录")

        # 统一使用用户表验证（包括admin用户）
        if db_manager.verify_user_password(request.username, request.password):
            user = db_manager.get_user_by_username(request.username)
            if user:
                # 生成token
                token = generate_token()
                SESSION_TOKENS[token] = {
                    'user_id': user['id'],
                    'username': user['username'],
                    'is_admin': user.get('is_admin', False) or user['username'] == ADMIN_USERNAME,
                    'timestamp': time.time()
                }

                # 区分管理员和普通用户的日志
                if user['username'] == ADMIN_USERNAME:
                    logger.info(f"【{user['username']}#{user['id']}】登录成功（管理员）")
                else:
                    logger.info(f"【{user['username']}#{user['id']}】登录成功")

                return LoginResponse(
                    success=True,
                    token=token,
                    message="登录成功",
                    user_id=user['id'],
                    username=user['username'],
                    is_admin=(user['username'] == ADMIN_USERNAME)
                )

        logger.warning(f"【{request.username}】登录失败：用户名或密码错误")
        return LoginResponse(
            success=False,
            message="用户名或密码错误"
        )

    elif request.email and request.password:
        # 邮箱/密码登录
        logger.info(f"【{request.email}】尝试邮箱密码登录")

        user = db_manager.get_user_by_email(request.email)
        if user and db_manager.verify_user_password(user['username'], request.password):
            # 生成token
            token = generate_token()
            SESSION_TOKENS[token] = {
                'user_id': user['id'],
                'username': user['username'],
                'is_admin': user.get('is_admin', False) or user['username'] == ADMIN_USERNAME,
                'timestamp': time.time()
            }

            logger.info(f"【{user['username']}#{user['id']}】邮箱登录成功")

            return LoginResponse(
                success=True,
                token=token,
                message="登录成功",
                user_id=user['id'],
                username=user['username'],
                is_admin=(user['username'] == ADMIN_USERNAME)
            )

        logger.warning(f"【{request.email}】邮箱登录失败：邮箱或密码错误")
        return LoginResponse(
            success=False,
            message="邮箱或密码错误"
        )

    elif request.email and request.verification_code:
        # 邮箱/验证码登录
        logger.info(f"【{request.email}】尝试邮箱验证码登录")

        # 验证邮箱验证码
        if not db_manager.verify_email_code(request.email, request.verification_code, 'login'):
            logger.warning(f"【{request.email}】验证码登录失败：验证码错误或已过期")
            return LoginResponse(
                success=False,
                message="验证码错误或已过期"
            )

        # 获取用户信息
        user = db_manager.get_user_by_email(request.email)
        if not user:
            logger.warning(f"【{request.email}】验证码登录失败：用户不存在")
            return LoginResponse(
                success=False,
                message="用户不存在"
            )

        # 生成token
        token = generate_token()
        SESSION_TOKENS[token] = {
            'user_id': user['id'],
            'username': user['username'],
            'is_admin': user.get('is_admin', False) or user['username'] == ADMIN_USERNAME,
            'timestamp': time.time()
        }

        logger.info(f"【{user['username']}#{user['id']}】验证码登录成功")

        return LoginResponse(
            success=True,
            token=token,
            message="登录成功",
            user_id=user['id'],
            username=user['username'],
            is_admin=(user['username'] == ADMIN_USERNAME)
        )

    else:
        return LoginResponse(
            success=False,
            message="请提供有效的登录信息"
        )


# 验证token接口
@app.get('/verify')
async def verify(user_info: Optional[Dict[str, Any]] = Depends(verify_token)):
    if user_info:
        return {
            "authenticated": True,
            "user_id": user_info['user_id'],
            "username": user_info['username'],
            "is_admin": user_info['username'] == ADMIN_USERNAME
        }
    return {"authenticated": False}


# 登出接口
@app.post('/logout')
async def logout(credentials: Optional[HTTPAuthorizationCredentials] = Depends(security)):
    if credentials and credentials.credentials in SESSION_TOKENS:
        del SESSION_TOKENS[credentials.credentials]
    return {"message": "已登出"}


# 修改管理员密码接口
@app.post('/change-admin-password')
async def change_admin_password(request: ChangePasswordRequest, admin_user: Dict[str, Any] = Depends(verify_admin_token)):
    from app.db_manager import db_manager

    try:
        # 验证当前密码（使用用户表验证）
        if not db_manager.verify_user_password('admin', request.current_password):
            return {"success": False, "message": "当前密码错误"}

        # 更新密码（使用用户表更新）
        success = db_manager.update_user_password('admin', request.new_password)

        if success:
            logger.info(f"【admin#{admin_user['user_id']}】管理员密码修改成功")
            return {"success": True, "message": "密码修改成功"}
        else:
            return {"success": False, "message": "密码修改失败"}

    except Exception as e:
        logger.error(f"修改管理员密码异常: {e}")
        return {"success": False, "message": "系统错误"}


# 普通用户修改密码接口
@app.post('/change-password')
async def change_user_password(request: ChangePasswordRequest, current_user: Dict[str, Any] = Depends(get_current_user)):
    from app.db_manager import db_manager

    try:
        username = current_user.get('username')
        user_id = current_user.get('user_id')
        
        if not username:
            return {"success": False, "message": "无法获取用户信息"}

        # 验证当前密码
        if not db_manager.verify_user_password(username, request.current_password):
            return {"success": False, "message": "当前密码错误"}

        # 更新密码
        success = db_manager.update_user_password(username, request.new_password)

        if success:
            logger.info(f"【{username}#{user_id}】用户密码修改成功")
            return {"success": True, "message": "密码修改成功"}
        else:
            return {"success": False, "message": "密码修改失败"}

    except Exception as e:
        logger.error(f"修改用户密码异常: {e}")
        return {"success": False, "message": "系统错误"}


# 检查是否使用默认密码
@app.get('/api/check-default-password')
async def check_default_password(current_user: Dict[str, Any] = Depends(get_current_user)):
    from app.db_manager import db_manager

    try:
        username = current_user.get('username')
        is_admin = current_user.get('is_admin', False)
        
        logger.info(f"检查默认密码: username={username}, is_admin={is_admin}")
        
        # 只检查admin用户
        if not is_admin or username != 'admin':
            logger.info(f"非admin用户，跳过检查")
            return {"using_default": False}

        # 检查是否使用默认密码
        using_default = db_manager.verify_user_password('admin', DEFAULT_ADMIN_PASSWORD)
        logger.info(f"默认密码检查结果: {using_default}, DEFAULT_ADMIN_PASSWORD={DEFAULT_ADMIN_PASSWORD}")
        
        return {"using_default": using_default}

    except Exception as e:
        logger.error(f"检查默认密码异常: {e}")
        return {"using_default": False}


# 生成图形验证码接口
@app.post('/generate-captcha')
async def generate_captcha(request: CaptchaRequest):
    from app.db_manager import db_manager

    try:
        # 生成图形验证码
        captcha_text, captcha_image = db_manager.generate_captcha()

        if not captcha_image:
            return CaptchaResponse(
                success=False,
                captcha_image="",
                session_id=request.session_id,
                message="图形验证码生成失败"
            )

        # 保存验证码到数据库
        if db_manager.save_captcha(request.session_id, captcha_text):
            return CaptchaResponse(
                success=True,
                captcha_image=captcha_image,
                session_id=request.session_id,
                message="图形验证码生成成功"
            )
        else:
            return CaptchaResponse(
                success=False,
                captcha_image="",
                session_id=request.session_id,
                message="图形验证码保存失败"
            )

    except Exception as e:
        logger.error(f"生成图形验证码失败: {e}")
        return CaptchaResponse(
            success=False,
            captcha_image="",
            session_id=request.session_id,
            message="图形验证码生成失败"
        )


# 验证图形验证码接口
@app.post('/verify-captcha')
async def verify_captcha(request: VerifyCaptchaRequest):
    from app.db_manager import db_manager

    try:
        if db_manager.verify_captcha(request.session_id, request.captcha_code):
            return VerifyCaptchaResponse(
                success=True,
                message="图形验证码验证成功"
            )
        else:
            return VerifyCaptchaResponse(
                success=False,
                message="图形验证码错误或已过期"
            )

    except Exception as e:
        logger.error(f"验证图形验证码失败: {e}")
        return VerifyCaptchaResponse(
            success=False,
            message="图形验证码验证失败"
        )


# ==================== 极验滑动验证码 ====================

# 极验验证状态存储: {challenge: {"status": int, "expires_at": float}}
geetest_status_store: dict = {}


def cleanup_expired_geetest_status():
    """清理过期的极验验证状态"""
    current_time = time.time()
    expired_keys = [k for k, v in geetest_status_store.items() if v["expires_at"] < current_time]
    for k in expired_keys:
        del geetest_status_store[k]


def set_geetest_status(challenge: str, status: int):
    """设置极验验证状态"""
    cleanup_expired_geetest_status()
    geetest_status_store[challenge] = {
        "status": status,
        "expires_at": time.time() + 300  # 5分钟有效
    }


def get_geetest_status(challenge: str) -> int:
    """获取极验验证状态，返回0表示未验证或已过期"""
    cleanup_expired_geetest_status()
    stored = geetest_status_store.get(challenge)
    if stored and stored["expires_at"] > time.time():
        return stored["status"]
    return 0


class GeetestRegisterResponse(BaseModel):
    """极验验证码初始化响应"""
    success: bool
    code: int = 200
    message: str = ""
    data: Optional[dict] = None


class GeetestValidateRequest(BaseModel):
    """极验二次验证请求"""
    challenge: str
    validate_str: str = Field(..., alias='validate')
    seccode: str

    model_config = {'populate_by_name': True}


class GeetestValidateResponse(BaseModel):
    """极验二次验证响应"""
    success: bool
    code: int = 200
    message: str = ""


@app.get('/geetest/register', response_model=GeetestRegisterResponse)
async def geetest_register():
    """
    获取极验验证码初始化参数
    
    前端调用此接口获取gt、challenge等参数，用于初始化验证码组件
    """
    try:
        from utils.geetest import GeetestLib
        
        gt_lib = GeetestLib()
        result = await gt_lib.register()
        
        data = result.to_dict()
        logger.info(f"极验初始化结果: status={result.status}, data={data}")
        
        # 记录初始状态
        challenge = data.get("challenge", "")
        if challenge:
            set_geetest_status(challenge, 0)
        
        return GeetestRegisterResponse(
            success=True,
            code=200,
            message="获取成功" if result.status == 1 else "宕机模式",
            data=data
        )
            
    except Exception as e:
        logger.error(f"极验初始化失败: {e}")
        # 返回本地初始化结果
        try:
            from utils.geetest import GeetestLib
            gt_lib = GeetestLib()
            result = gt_lib.local_init()
            data = result.to_dict()
            
            # 记录初始状态
            challenge = data.get("challenge", "")
            if challenge:
                set_geetest_status(challenge, 0)
            
            return GeetestRegisterResponse(
                success=True,
                code=200,
                message="本地初始化",
                data=data
            )
        except Exception as e2:
            logger.error(f"极验本地初始化也失败: {e2}")
            return GeetestRegisterResponse(
                success=False,
                code=500,
                message="验证码服务异常"
            )


@app.post('/geetest/validate', response_model=GeetestValidateResponse)
async def geetest_validate(request: GeetestValidateRequest):
    """
    极验二次验证
    
    用户完成滑动验证后，前端调用此接口进行二次验证
    """
    try:
        # 检查是否已经验证过
        if get_geetest_status(request.challenge) == 1:
            return GeetestValidateResponse(
                success=True,
                code=200,
                message="验证通过"
            )
        
        from utils.geetest import GeetestLib
        
        gt_lib = GeetestLib()
        
        # 判断是正常模式还是宕机模式
        # 通过challenge长度判断：正常模式challenge是32位MD5，宕机模式是UUID
        is_normal_mode = len(request.challenge) == 32
        
        if is_normal_mode:
            result = await gt_lib.success_validate(
                request.challenge,
                request.validate_str,
                request.seccode
            )
        else:
            result = gt_lib.fail_validate(
                request.challenge,
                request.validate_str,
                request.seccode
            )
        
        if result.status == 1:
            # 记录验证通过状态
            set_geetest_status(request.challenge, 1)
            
            return GeetestValidateResponse(
                success=True,
                code=200,
                message="验证通过"
            )
        else:
            return GeetestValidateResponse(
                success=False,
                code=400,
                message=result.msg or "验证失败"
            )
            
    except Exception as e:
        logger.error(f"极验二次验证失败: {e}")
        return GeetestValidateResponse(
            success=False,
            code=500,
            message="验证服务异常"
        )


# 发送验证码接口（需要先验证图形验证码）
@app.post('/send-verification-code')
async def send_verification_code(request: SendCodeRequest):
    from app.db_manager import db_manager

    try:
        # 检查是否已验证图形验证码
        # 通过检查数据库中是否存在已验证的图形验证码记录
        with db_manager.lock:
            cursor = db_manager.conn.cursor()
            current_time = time.time()

            # 查找最近5分钟内该session_id的验证记录
            # 由于验证成功后验证码会被删除，我们需要另一种方式来跟踪验证状态
            # 这里我们检查该session_id是否在最近验证过（通过检查是否有已删除的记录）

            # 为了简化，我们要求前端在验证图形验证码成功后立即发送邮件验证码
            # 或者我们可以在验证成功后设置一个临时标记

        # 根据验证码类型进行不同的检查
        if request.type == 'register':
            # 注册验证码：检查邮箱是否已注册
            existing_user = db_manager.get_user_by_email(request.email)
            if existing_user:
                return SendCodeResponse(
                    success=False,
                    message="该邮箱已被注册"
                )
        elif request.type == 'login':
            # 登录验证码：检查邮箱是否存在
            existing_user = db_manager.get_user_by_email(request.email)
            if not existing_user:
                return SendCodeResponse(
                    success=False,
                    message="该邮箱未注册"
                )

        # 生成验证码
        code = db_manager.generate_verification_code()

        # 保存验证码到数据库
        if not db_manager.save_verification_code(request.email, code, request.type):
            return SendCodeResponse(
                success=False,
                message="验证码保存失败，请稍后重试"
            )

        # 发送验证码邮件
        if await db_manager.send_verification_email(request.email, code):
            return SendCodeResponse(
                success=True,
                message="验证码已发送到您的邮箱，请查收"
            )
        else:
            return SendCodeResponse(
                success=False,
                message="验证码发送失败，请检查邮箱地址或稍后重试"
            )

    except Exception as e:
        logger.error(f"发送验证码失败: {e}")
        return SendCodeResponse(
            success=False,
            message="发送验证码失败，请稍后重试"
        )


# 用户注册接口
@app.post('/register')
async def register(request: RegisterRequest):
    from app.db_manager import db_manager

    # 检查注册是否开启
    registration_enabled = db_manager.get_system_setting('registration_enabled')
    if registration_enabled != 'true':
        logger.warning(f"【{request.username}】注册失败: 注册功能已关闭")
        return RegisterResponse(
            success=False,
            message="注册功能已关闭，请联系管理员"
        )

    try:
        logger.info(f"【{request.username}】尝试注册，邮箱: {request.email}")

        # 邮箱验证码是否必填由管理员在系统设置里控制。
        # 没配 SMTP 的部署发不出验证码，强制校验会让注册完全不可用；
        # 关掉这项就能先用起来，配好邮件服务后再开回去。
        email_verification = db_manager.get_system_setting('email_verification_enabled')
        # 老库没有这一项，按开启处理，避免升级后安全性被悄悄降低
        require_email_code = str(email_verification or 'true').strip().lower() not in ('0', 'false', 'no')

        if require_email_code:
            # 验证邮箱验证码
            if not db_manager.verify_email_code(request.email, request.verification_code):
                logger.warning(f"【{request.username}】注册失败: 验证码错误或已过期")
                return RegisterResponse(
                    success=False,
                    message="验证码错误或已过期"
                )

        # 检查用户名是否已存在
        existing_user = db_manager.get_user_by_username(request.username)
        if existing_user:
            logger.warning(f"【{request.username}】注册失败: 用户名已存在")
            return RegisterResponse(
                success=False,
                message="用户名已存在"
            )

        # 检查邮箱是否已注册
        existing_email = db_manager.get_user_by_email(request.email)
        if existing_email:
            logger.warning(f"【{request.username}】注册失败: 邮箱已被注册")
            return RegisterResponse(
                success=False,
                message="该邮箱已被注册"
            )

        # 创建用户
        if db_manager.create_user(request.username, request.email, request.password):
            logger.info(f"【{request.username}】注册成功")
            return RegisterResponse(
                success=True,
                message="注册成功，请登录"
            )
        else:
            logger.error(f"【{request.username}】注册失败: 数据库操作失败")
            return RegisterResponse(
                success=False,
                message="注册失败，请稍后重试"
            )

    except Exception as e:
        logger.error(f"【{request.username}】注册异常: {e}")
        return RegisterResponse(
            success=False,
            message="注册失败，请稍后重试"
        )


# ------------------------- 发送消息接口 -------------------------

# 固定的API秘钥（生产环境中应该从配置文件或环境变量读取）
# 注意：现在从系统设置中读取QQ回复消息秘钥
API_SECRET_KEY = "xianyu_api_secret_2024"  # 保留作为后备

class SendMessageRequest(BaseModel):
    api_key: str
    cookie_id: str
    chat_id: str
    to_user_id: str
    message: str


class SendMessageResponse(BaseModel):
    success: bool
    message: str


class ChatSendMessageRequest(BaseModel):
    cid: str = Field(..., min_length=1, max_length=200)
    to_user_id: str = Field(..., min_length=1, max_length=100)
    text: str = Field(..., min_length=1, max_length=2000)


def verify_api_key(api_key: str) -> bool:
    """验证API秘钥"""
    try:
        # 从系统设置中获取QQ回复消息秘钥
        from app.db_manager import db_manager
        qq_secret_key = db_manager.get_system_setting('qq_reply_secret_key')

        # 如果系统设置中没有配置，使用默认值
        if not qq_secret_key:
            qq_secret_key = API_SECRET_KEY

        return api_key == qq_secret_key
    except Exception as e:
        logger.error(f"验证API秘钥时发生异常: {e}")
        # 异常情况下使用默认秘钥验证
        return api_key == API_SECRET_KEY


def _get_owned_chat_account(cookie_id: str, current_user: Dict[str, Any]) -> str:
    owned_cookies = db_manager.get_all_cookies(current_user["user_id"])
    if cookie_id not in owned_cookies:
        raise HTTPException(status_code=403, detail="无权访问该闲鱼账号")
    return cookie_id


class _AccountRequestDedup:
    """把短时间内重复的同一账号请求合并成一次。

    消息与会话列表在前端是轮询的，多开标签页、或前后端各自重试时，同一份数据
    会被反复经 WebSocket 透传到闲鱼，很容易把账号打到限流（429 flow controled）。
    这里做两件事：
    1. 正在执行的相同请求直接复用同一个 Future，不再发第二次；
    2. 结果按 TTL 短暂缓存；命中限流时用更长的 TTL，避免雪上加霜。
    """

    NORMAL_TTL = 3.0
    THROTTLED_TTL = 15.0
    MAX_ENTRIES = 512

    def __init__(self):
        self._pending: Dict[str, asyncio.Future] = {}
        self._cache: "OrderedDict[str, Tuple[float, Any]]" = OrderedDict()

    @staticmethod
    def _looks_throttled(result) -> bool:
        text = str(result)
        return 'flow controled' in text or 'FAIL_SYS_FLOW_LIMIT' in text or '429' in text

    def _prune(self) -> None:
        now = time.monotonic()
        for key in [k for k, (expire_at, _) in self._cache.items() if expire_at <= now]:
            self._cache.pop(key, None)
        while len(self._cache) > self.MAX_ENTRIES:
            self._cache.popitem(last=False)

    async def run(self, key: str, factory):
        cached = self._cache.get(key)
        if cached and cached[0] > time.monotonic():
            return cached[1]

        pending = self._pending.get(key)
        if pending is not None and not pending.done():
            return await asyncio.shield(pending)

        task = asyncio.ensure_future(factory())
        self._pending[key] = task
        try:
            result = await task
        finally:
            self._pending.pop(key, None)

        ttl = self.THROTTLED_TTL if self._looks_throttled(result) else self.NORMAL_TTL
        self._cache[key] = (time.monotonic() + ttl, result)
        self._prune()
        return result


account_request_dedup = _AccountRequestDedup()


def _is_account_connection_alive(instance) -> bool:
    """判断账号的 WebSocket 是否真的可用。

    只看 connection_state 不够：状态是 CONNECTED 但连接对象已被对端关闭时
    （浏览器标签页长时间挂后台会出现这种"半开"），请求仍会一路等到超时。
    """
    state = getattr(instance, 'connection_state', None)
    state_value = getattr(state, 'value', state)
    if state_value != 'connected':
        return False

    ws = getattr(instance, 'ws', None)
    if ws is None:
        return False
    # websockets 的连接对象用 closed 标记；不同版本属性不一致，缺失时按可用处理
    closed = getattr(ws, 'closed', False)
    return not closed


async def _run_on_account_loop(cookie_id: str, operation):
    manager = cookie_manager.manager
    if manager is None or not manager.loop.is_running():
        raise HTTPException(status_code=503, detail="闲鱼账号服务尚未启动")

    instance = manager.instances.get(cookie_id)
    if instance is None:
        # 区分风控和未登录：两者的处理方式完全不同，笼统提示"未在线"会误导用户
        from utils import risk_control

        guard = risk_control.registry.get(cookie_id)
        if guard.is_blocked:
            minutes = max(1, guard.remaining_seconds // 60)
            raise HTTPException(
                status_code=409,
                detail=(
                    f"闲鱼要求人机验证，账号暂停请求中（约 {minutes} 分钟后自动重试）。"
                    f"请勿频繁操作，等待自动恢复或稍后重新扫码登录。"
                ),
            )
        raise HTTPException(status_code=409, detail="账号未在线，请先启用账号并等待连接成功")

    # 连接不可用时立刻返回，不要走到下面等 25 秒。浏览器标签页切到后台后
    # WebSocket 会进入"半开"状态（TCP 还连着、应用层已不通），此时发请求会一路
    # 等到超时：内层 15 秒 + 外层 25 秒，用户要等约 40 秒才看到 504。
    if not _is_account_connection_alive(instance):
        # 登录态过期是终态，别让用户以为等一会儿就好
        if getattr(instance, 'needs_relogin', False):
            raise HTTPException(
                status_code=409,
                detail=(
                    getattr(instance, 'relogin_reason', '')
                    or '闲鱼登录态已过期，请重新扫码登录该账号'
                ),
            )
        await asyncio.sleep(3)  # 给自动重连一点时间
        if not _is_account_connection_alive(instance):
            if getattr(instance, 'needs_relogin', False):
                raise HTTPException(
                    status_code=409,
                    detail='闲鱼登录态已过期，请重新扫码登录该账号',
                )
            raise HTTPException(
                status_code=503,
                detail="账号连接已断开，正在自动重连，请稍后刷新重试",
            )

    try:
        current_loop = asyncio.get_running_loop()
        if current_loop is manager.loop:
            return await operation(instance)
        future = asyncio.run_coroutine_threadsafe(operation(instance), manager.loop)
        return await asyncio.wait_for(asyncio.wrap_future(future), timeout=25)
    except asyncio.TimeoutError as exc:
        raise HTTPException(status_code=504, detail="闲鱼消息服务响应超时") from exc
    except HTTPException:
        raise
    except (ConnectionError, TimeoutError) as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        logger.warning(f"【{cookie_id}】闲鱼消息请求失败: {exc}")
        raise HTTPException(status_code=502, detail=f"闲鱼消息服务请求失败: {exc}") from exc


@app.get("/chat/accounts")
async def get_chat_accounts(current_user: Dict[str, Any] = Depends(get_current_user)):
    owned_cookies = db_manager.get_all_cookies(current_user["user_id"])
    manager = cookie_manager.manager
    instances = manager.instances if manager else {}
    result = []
    for cookie_id in owned_cookies:
        details = db_manager.get_cookie_details(cookie_id) or {}
        instance = instances.get(cookie_id)
        websocket = getattr(instance, "ws", None) if instance else None
        result.append({
            "accountId": cookie_id,
            "displayName": details.get("nickname") or details.get("remark") or cookie_id,
            "avatarUrl": details.get("avatar_url") or "",
            "connected": websocket is not None and not bool(getattr(websocket, "closed", False)),
            "xianyuUserId": str(getattr(instance, "myid", "") or ""),
        })
    return {"success": True, "data": result}


@app.get("/chat/conversations/{cookie_id}")
async def get_chat_conversations(
    cookie_id: str,
    cursor: Optional[int] = Query(None),
    limit: int = Query(30, ge=1, le=100),
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    _get_owned_chat_account(cookie_id, current_user)
    # 走去重：会话列表是前端轮询的，多标签页并发时同一份数据会被重复
    # 透传到闲鱼，容易触发账号限流
    body = await account_request_dedup.run(
        f"conversations|{cookie_id}|{cursor}|{limit}",
        lambda: _run_on_account_loop(
            cookie_id,
            lambda instance: instance.get_im_conversations(cursor, limit),
        ),
    )
    if isinstance(body, dict) and (body.get("reason") or body.get("code") == "400600001"):
        reason = body.get("developerMessage") or body.get("reason") or body.get("code")
        raise HTTPException(status_code=429, detail=f"闲鱼会话请求受限: {reason}")

    from app.xianyu_im import parse_conversation

    manager = cookie_manager.manager
    instance = manager.instances.get(cookie_id) if manager else None
    my_id = str(getattr(instance, "myid", "") or "")
    conversations = []
    for item in body.get("userConvs", []) if isinstance(body, dict) else []:
        raw = item.get("singleChatUserConversation", item) if isinstance(item, dict) else {}
        parsed = parse_conversation(raw, my_id)
        if parsed:
            conversations.append(parsed)
    return {
        "success": True,
        "data": {
            "conversations": conversations,
            "hasMore": bool(body.get("hasMore", False)) if isinstance(body, dict) else False,
            "nextCursor": body.get("nextCursor") if isinstance(body, dict) else None,
        },
    }


@app.get("/chat/messages/{cookie_id}/{cid}")
async def get_chat_messages(
    cookie_id: str,
    cid: str,
    cursor: Optional[int] = Query(None),
    limit: int = Query(50, ge=1, le=100),
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    _get_owned_chat_account(cookie_id, current_user)
    body = await account_request_dedup.run(
        f"messages|{cookie_id}|{cid}|{cursor}|{limit}",
        lambda: _run_on_account_loop(
            cookie_id,
            lambda instance: instance.get_im_messages(cid, cursor, limit),
        ),
    )
    if isinstance(body, dict) and body.get("reason"):
        reason = body.get("developerMessage") or body.get("reason")
        raise HTTPException(status_code=502, detail=f"闲鱼消息记录获取失败: {reason}")

    from app.xianyu_im import parse_message

    manager = cookie_manager.manager
    instance = manager.instances.get(cookie_id) if manager else None
    my_id = str(getattr(instance, "myid", "") or "")

    # 这个接口的响应结构有两种：消息可能直接挂在 body 下，也可能包在 body.data 里。
    # 只读一层时另一种结构会永远拿到空列表，表现为「点开会话看不到任何消息」。
    payload = body if isinstance(body, dict) else {}
    inner = payload.get("data") if isinstance(payload.get("data"), dict) else {}

    def pick(field, default=None):
        value = inner.get(field)
        if value is None:
            value = payload.get(field)
        return default if value is None else value

    messages = []
    for model in pick("userMessageModels", []) or []:
        parsed = parse_message(model, my_id)
        if parsed:
            messages.append(parsed)
    messages.reverse()
    return {
        "success": True,
        "data": {
            "messages": messages,
            "hasMore": bool(pick("hasMore", False)),
            "nextCursor": pick("nextCursor"),
        },
    }


@app.post("/chat/send/{cookie_id}")
async def send_chat_message(
    cookie_id: str,
    request: ChatSendMessageRequest,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    _get_owned_chat_account(cookie_id, current_user)
    response = await _run_on_account_loop(
        cookie_id,
        lambda instance: instance.send_im_text(
            request.cid.strip(),
            request.to_user_id.strip(),
            request.text,
        ),
    )
    body = response.get("body", {}) if isinstance(response, dict) else {}
    message_id = ""
    if isinstance(body, dict):
        message_id = str(body.get("messageId") or body.get("msgId") or "")
    logger.info(
        f"【{cookie_id}】后台用户 {current_user.get('username')} 人工发送闲鱼消息，"
        f"会话={request.cid}, 对方={request.to_user_id}, 长度={len(request.text)}"
    )
    return {"success": True, "message": "发送成功", "data": {"messageId": message_id}}


@app.post('/send-message', response_model=SendMessageResponse)
async def send_message_api(request: SendMessageRequest):
    """发送消息API接口（使用秘钥验证）"""
    try:
        # 清理所有参数中的换行符
        def clean_param(param_str):
            """清理参数中的换行符"""
            if isinstance(param_str, str):
                return param_str.replace('\\n', '').replace('\n', '')
            return param_str

        # 清理所有参数
        cleaned_api_key = clean_param(request.api_key)
        cleaned_cookie_id = clean_param(request.cookie_id)
        cleaned_chat_id = clean_param(request.chat_id)
        cleaned_to_user_id = clean_param(request.to_user_id)
        cleaned_message = clean_param(request.message)

        # 验证API秘钥不能为空
        if not cleaned_api_key:
            logger.warning("API秘钥为空")
            return SendMessageResponse(
                success=False,
                message="API秘钥不能为空"
            )

        # 验证API秘钥
        if not verify_api_key(cleaned_api_key):
            logger.warning(f"API秘钥验证失败: {cleaned_api_key}")
            return SendMessageResponse(
                success=False,
                message="API秘钥验证失败"
            )

        # 验证必需参数不能为空
        required_params = {
            'cookie_id': cleaned_cookie_id,
            'chat_id': cleaned_chat_id,
            'to_user_id': cleaned_to_user_id,
            'message': cleaned_message
        }

        for param_name, param_value in required_params.items():
            if not param_value:
                logger.warning(f"必需参数 {param_name} 为空")
                return SendMessageResponse(
                    success=False,
                    message=f"参数 {param_name} 不能为空"
                )

        # 直接获取XianyuLive实例，跳过cookie_manager检查
        from XianyuAutoAsync import XianyuLive
        live_instance = XianyuLive.get_instance(cleaned_cookie_id)

        if not live_instance:
            logger.warning(f"账号实例不存在或未连接: {cleaned_cookie_id}")
            return SendMessageResponse(
                success=False,
                message="账号实例不存在或未连接，请检查账号状态"
            )

        # 检查WebSocket连接状态
        if not live_instance.ws or live_instance.ws.closed:
            logger.warning(f"账号WebSocket连接已断开: {cleaned_cookie_id}")
            return SendMessageResponse(
                success=False,
                message="账号WebSocket连接已断开，请等待重连"
            )

        # 发送消息（使用清理后的所有参数）
        await live_instance.send_msg(
            live_instance.ws,
            cleaned_chat_id,
            cleaned_to_user_id,
            cleaned_message
        )

        logger.info(f"API成功发送消息: {cleaned_cookie_id} -> {cleaned_to_user_id}, 内容: {cleaned_message[:50]}{'...' if len(cleaned_message) > 50 else ''}")

        return SendMessageResponse(
            success=True,
            message="消息发送成功"
        )

    except Exception as e:
        # 使用清理后的参数记录日志
        cookie_id_for_log = clean_param(request.cookie_id) if 'clean_param' in locals() else request.cookie_id
        to_user_id_for_log = clean_param(request.to_user_id) if 'clean_param' in locals() else request.to_user_id
        logger.error(f"API发送消息异常: {cookie_id_for_log} -> {to_user_id_for_log}, 错误: {str(e)}")
        return SendMessageResponse(
            success=False,
            message=f"发送消息失败: {str(e)}"
        )


@app.post("/xianyu/reply", response_model=ResponseModel)
async def xianyu_reply(req: RequestModel):
    msg_template = match_reply(req.cookie_id, req.send_message)
    is_default_reply = False

    if not msg_template:
        # 从数据库获取默认回复
        from app.db_manager import db_manager
        default_reply_settings = db_manager.get_default_reply(req.cookie_id)

        if default_reply_settings and default_reply_settings.get('enabled', False):
            # 检查是否开启了"只回复一次"功能
            if default_reply_settings.get('reply_once', False):
                # 检查是否已经回复过这个chat_id
                if db_manager.has_default_reply_record(req.cookie_id, req.chat_id):
                    raise HTTPException(status_code=404, detail="该对话已使用默认回复，不再重复回复")

            msg_template = default_reply_settings.get('reply_content', '')
            is_default_reply = True

        # 如果数据库中没有设置或为空，返回错误
        if not msg_template:
            raise HTTPException(status_code=404, detail="未找到匹配的回复规则且未设置默认回复")

    # 按占位符格式化
    try:
        send_msg = msg_template.format(
            send_user_id=req.send_user_id,
            send_user_name=req.send_user_name,
            send_message=req.send_message,
        )
    except Exception:
        # 如果格式化失败，返回原始内容
        send_msg = msg_template

    # 如果是默认回复且开启了"只回复一次"，记录回复记录
    if is_default_reply:
        from app.db_manager import db_manager
        default_reply_settings = db_manager.get_default_reply(req.cookie_id)
        if default_reply_settings and default_reply_settings.get('reply_once', False):
            db_manager.add_default_reply_record(req.cookie_id, req.chat_id)

    return {"code": 200, "data": {"send_msg": send_msg}}

# ------------------------- 账号 / 关键字管理接口 -------------------------


class CookieIn(BaseModel):
    id: str
    value: str


class CookieStatusIn(BaseModel):
    enabled: bool


class DefaultReplyIn(BaseModel):
    enabled: bool
    reply_content: Optional[str] = None
    reply_image_url: Optional[str] = None
    reply_once: bool = False


class NotificationChannelIn(BaseModel):
    name: str
    type: str
    config: str


class NotificationChannelUpdate(BaseModel):
    name: Optional[str] = None
    config: Optional[str] = None
    enabled: Optional[bool] = None


class MessageNotificationIn(BaseModel):
    channel_id: int
    enabled: bool = True


class MessageFilterIn(BaseModel):
    cookie_id: str
    keyword: str
    filter_type: str
    enabled: bool = True


class MessageFilterUpdate(BaseModel):
    keyword: Optional[str] = None
    filter_type: Optional[str] = None
    enabled: Optional[bool] = None


class MessageFilterBatchIn(BaseModel):
    cookie_id: str
    keywords: List[str]
    filter_type: str
    enabled: bool = True


class MessageFilterBatchDelete(BaseModel):
    ids: List[int]


MESSAGE_FILTER_TYPES = {"skip_reply", "skip_notify"}


def validate_message_filter(cookie_id: str, keyword: str, filter_type: str) -> Tuple[str, str, str]:
    normalized_cookie_id = (cookie_id or "").strip()
    normalized_keyword = (keyword or "").strip()
    normalized_type = (filter_type or "").strip().lower()
    if not normalized_cookie_id:
        raise ValueError("账号不能为空")
    if not normalized_keyword:
        raise ValueError("过滤关键词不能为空")
    if len(normalized_keyword) > 200:
        raise ValueError("过滤关键词不能超过 200 个字符")
    if normalized_type not in MESSAGE_FILTER_TYPES:
        raise ValueError("无效的过滤类型")
    return normalized_cookie_id, normalized_keyword, normalized_type


NOTIFICATION_CHANNEL_REQUIRED_FIELDS = {
    "dingtalk": ("webhook_url",),
    "feishu": ("webhook_url",),
    "bark": ("device_key",),
    "email": ("smtp_server", "smtp_port", "email_user", "email_password", "recipient_email"),
    "webhook": ("webhook_url",),
    "wechat": ("webhook_url",),
    "telegram": ("bot_token", "chat_id"),
}
NOTIFICATION_CHANNEL_TYPE_ALIASES = {
    "ding_talk": "dingtalk",
    "lark": "feishu",
}


def validate_notification_channel(
    name: str,
    channel_type: str,
    config: str,
) -> Tuple[str, str, str]:
    """Validate and normalize notification channel data before persistence."""
    normalized_name = (name or "").strip()
    if not normalized_name:
        raise ValueError("通知渠道名称不能为空")
    if len(normalized_name) > 80:
        raise ValueError("通知渠道名称不能超过 80 个字符")

    normalized_type = (channel_type or "").strip().lower()
    normalized_type = NOTIFICATION_CHANNEL_TYPE_ALIASES.get(normalized_type, normalized_type)
    if normalized_type not in NOTIFICATION_CHANNEL_REQUIRED_FIELDS:
        raise ValueError("不支持的通知渠道类型")

    try:
        config_data = json.loads(config)
    except (TypeError, json.JSONDecodeError) as exc:
        raise ValueError("通知渠道配置必须是有效的 JSON") from exc
    if not isinstance(config_data, dict):
        raise ValueError("通知渠道配置必须是 JSON 对象")

    missing_fields = [
        field
        for field in NOTIFICATION_CHANNEL_REQUIRED_FIELDS[normalized_type]
        if config_data.get(field) in (None, "")
    ]
    if missing_fields:
        raise ValueError(f"通知渠道配置缺少字段: {', '.join(missing_fields)}")

    if normalized_type == "email":
        try:
            smtp_port = int(config_data["smtp_port"])
        except (TypeError, ValueError) as exc:
            raise ValueError("SMTP 端口必须是数字") from exc
        if not 1 <= smtp_port <= 65535:
            raise ValueError("SMTP 端口必须在 1-65535 之间")
        config_data["smtp_port"] = smtp_port

    if normalized_type == "webhook":
        http_method = str(config_data.get("http_method", "POST")).upper()
        if http_method not in {"POST", "PUT"}:
            raise ValueError("Webhook 请求方法仅支持 POST 或 PUT")
        config_data["http_method"] = http_method

        headers = config_data.get("headers")
        if isinstance(headers, str) and headers.strip():
            try:
                parsed_headers = json.loads(headers)
            except json.JSONDecodeError as exc:
                raise ValueError("Webhook 请求头必须是有效的 JSON 对象") from exc
            if not isinstance(parsed_headers, dict):
                raise ValueError("Webhook 请求头必须是 JSON 对象")
        elif headers is not None and not isinstance(headers, dict):
            raise ValueError("Webhook 请求头必须是 JSON 对象")

    return (
        normalized_name,
        normalized_type,
        json.dumps(config_data, ensure_ascii=False, separators=(",", ":")),
    )


class SystemSettingIn(BaseModel):
    value: str
    description: Optional[str] = None


class SystemSettingCreateIn(BaseModel):
    key: str
    value: str
    description: Optional[str] = None





@app.get("/cookies")
def list_cookies(current_user: Dict[str, Any] = Depends(get_current_user)):
    if cookie_manager.manager is None:
        return []

    # 获取当前用户的cookies
    user_id = current_user['user_id']
    from app.db_manager import db_manager
    user_cookies = db_manager.get_all_cookies(user_id)
    return list(user_cookies.keys())


@app.get("/cookies/details")
def get_cookies_details(current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取所有Cookie的详细信息（包括值和状态）"""
    if cookie_manager.manager is None:
        return []

    # 获取当前用户的cookies
    user_id = current_user['user_id']
    from app.db_manager import db_manager
    user_cookies = db_manager.get_all_cookies(user_id)

    result = []
    for cookie_id, cookie_value in user_cookies.items():
        cookie_enabled = cookie_manager.manager.get_cookie_status(cookie_id)
        auto_confirm = db_manager.get_auto_confirm(cookie_id)
        # 获取备注信息
        cookie_details = db_manager.get_cookie_details(cookie_id) or {}
        remark = cookie_details.get('remark', '')

        result.append({
            'id': cookie_id,
            'value': cookie_value,
            'enabled': cookie_enabled,
            'auto_confirm': auto_confirm,
            'remark': remark,
            'pause_duration': cookie_details.get('pause_duration', 10),
            'nickname': cookie_details.get('nickname', ''),
            'avatar_url': cookie_details.get('avatar_url', ''),
            'location': cookie_details.get('location', ''),
            'bio': cookie_details.get('bio', ''),
            'followers': cookie_details.get('followers'),
            'following': cookie_details.get('following'),
            'profile_updated_at': cookie_details.get('profile_updated_at'),
            'runtime_state': cookie_manager.manager.get_task_state(cookie_id),
            # 登录信息一并返回。这几个字段本来要前端再逐个请求
            # /cookie/{id}/details 才拿得到，等于对同一份数据做 N+1 查询：
            # 页面上有 7 处组件各自调用，账号一多首屏就被这些重复请求拖慢。
            'username': cookie_details.get('username', ''),
            'login_password': cookie_details.get('password', '') or cookie_details.get('login_password', ''),
            'show_browser': bool(cookie_details.get('show_browser', False)),
        })
    return result


@app.post("/cookies")
def add_cookie(item: CookieIn, current_user: Dict[str, Any] = Depends(get_current_user)):
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")
    try:
        # 添加cookie时绑定到当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager

        log_with_user('info', f"尝试添加Cookie: {item.id}, 当前用户ID: {user_id}, 用户名: {current_user.get('username', 'unknown')}", current_user)

        # 检查cookie是否已存在且属于其他用户
        existing_cookies = db_manager.get_all_cookies()
        if item.id in existing_cookies:
            # 检查是否属于当前用户
            user_cookies = db_manager.get_all_cookies(user_id)
            if item.id not in user_cookies:
                log_with_user('warning', f"Cookie ID冲突: {item.id} 已被其他用户使用", current_user)
                raise HTTPException(status_code=400, detail="该Cookie ID已被其他用户使用")

        # 保存到数据库时指定用户ID
        db_manager.save_cookie(item.id, item.value, user_id)

        # 添加到CookieManager，同时指定用户ID
        cookie_manager.manager.add_cookie(item.id, item.value, user_id=user_id)
        log_with_user('info', f"Cookie添加成功: {item.id}", current_user)
        return {"msg": "success"}
    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"添加Cookie失败: {item.id} - {str(e)}", current_user)
        raise HTTPException(status_code=400, detail=str(e))


# ============ 带子路径的 /cookies/{cid}/xxx 路由必须在 /cookies/{cid} 之前定义 ============

class AccountLoginInfoUpdate(BaseModel):
    username: Optional[str] = None
    login_password: Optional[str] = None
    show_browser: Optional[bool] = None


async def _fetch_and_store_account_profile(
    cookie_id: str,
    cookies_str: str,
    user_id: int,
) -> Dict[str, Any]:
    """抓取并保存账号公开资料，不返回或记录Cookie内容。"""
    from XianyuAutoAsync import XianyuLive

    instance = XianyuLive(
        cookies_str=cookies_str,
        cookie_id=cookie_id,
        user_id=user_id,
    )
    result = await instance.fetch_account_profile()
    if not result.get('success'):
        return result

    profile = result.get('profile') or {}
    if not db_manager.update_cookie_profile(cookie_id, profile):
        return {
            'success': False,
            'error': '账号资料已获取，但保存失败',
            'error_type': 'ProfilePersistenceError',
        }

    details = db_manager.get_cookie_details(cookie_id) or {}
    return {
        'success': True,
        'profile': {
            key: details.get(key)
            for key in (
                'nickname',
                'avatar_url',
                'location',
                'bio',
                'followers',
                'following',
                'profile_updated_at',
            )
        },
    }


@app.put("/cookies/{cid}/login-info")
def update_cookie_login_info(cid: str, update_data: AccountLoginInfoUpdate, current_user: Dict[str, Any] = Depends(get_current_user)):
    """更新账号登录信息（用户名、密码、是否显示浏览器）"""
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        # 使用现有的update_cookie_account_info方法更新登录信息
        success = db_manager.update_cookie_account_info(
            cid,
            username=update_data.username,
            password=update_data.login_password,
            show_browser=update_data.show_browser
        )

        if success:
            return {"success": True, "message": "登录信息已更新"}
        else:
            raise HTTPException(status_code=500, detail="更新登录信息失败")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/cookies/{cid}/refresh-profile")
async def refresh_cookie_profile(
    cid: str,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    """使用现有Cookie刷新闲鱼账号的公开资料。"""
    user_id = current_user['user_id']
    user_cookies = db_manager.get_all_cookies(user_id)
    if cid not in user_cookies:
        raise HTTPException(status_code=403, detail="无权限操作该Cookie")

    log_with_user('info', f"开始刷新账号公开资料: 账号={cid}", current_user)
    result = await _fetch_and_store_account_profile(
        cookie_id=cid,
        cookies_str=user_cookies[cid],
        user_id=user_id,
    )
    if not result.get('success'):
        log_with_user(
            'warning',
            f"账号公开资料刷新失败: 账号={cid}, "
            f"异常类型={result.get('error_type', 'unknown')}",
            current_user,
        )
        raise HTTPException(
            status_code=502,
            detail=result.get('error', '账号资料刷新失败'),
        )

    log_with_user('info', f"账号公开资料刷新成功: 账号={cid}", current_user)
    return result


# ============ 通用的 /cookies/{cid} 路由 ============

@app.put('/cookies/{cid}')
def update_cookie(cid: str, item: CookieIn, current_user: Dict[str, Any] = Depends(get_current_user)):
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail='CookieManager 未就绪')
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        # 获取旧的 cookie 值，用于判断是否需要重启任务
        old_cookie_details = db_manager.get_cookie_details(cid)
        old_cookie_value = old_cookie_details.get('value') if old_cookie_details else None

        # 使用 update_cookie_account_info 更新（只更新cookie值，不覆盖其他字段）
        success = db_manager.update_cookie_account_info(cid, cookie_value=item.value)
        
        if not success:
            raise HTTPException(status_code=400, detail="更新Cookie失败")
        
        # 只有当 cookie 值真的发生变化时才重启任务
        if item.value != old_cookie_value:
            logger.info(f"Cookie值已变化，重启任务: {cid}")
            cookie_manager.manager.update_cookie(cid, item.value, save_to_db=False)
        else:
            logger.info(f"Cookie值未变化，无需重启任务: {cid}")
        
        return {'msg': 'updated'}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


class CookieAccountInfo(BaseModel):
    """账号信息更新模型"""
    value: Optional[str] = None
    username: Optional[str] = None
    password: Optional[str] = None
    show_browser: Optional[bool] = None


@app.post("/cookie/{cid}/account-info")
def update_cookie_account_info(cid: str, info: CookieAccountInfo, current_user: Dict[str, Any] = Depends(get_current_user)):
    """更新账号信息（Cookie、用户名、密码、显示浏览器设置）"""
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail='CookieManager 未就绪')
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        # 获取旧的 cookie 值，用于判断是否需要重启任务
        old_cookie_details = db_manager.get_cookie_details(cid)
        old_cookie_value = old_cookie_details.get('value') if old_cookie_details else None
        
        # 更新数据库
        success = db_manager.update_cookie_account_info(
            cid, 
            cookie_value=info.value,
            username=info.username,
            password=info.password,
            show_browser=info.show_browser
        )
        
        if not success:
            raise HTTPException(status_code=400, detail="更新账号信息失败")
        
        # 只有当 cookie 值真的发生变化时才重启任务
        if info.value is not None and info.value != old_cookie_value:
            logger.info(f"Cookie值已变化，重启任务: {cid}")
            cookie_manager.manager.update_cookie(cid, info.value, save_to_db=False)
        else:
            logger.info(f"Cookie值未变化，无需重启任务: {cid}")
        
        return {'msg': 'updated', 'success': True}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"更新账号信息失败: {e}")
        raise HTTPException(status_code=400, detail=str(e))


@app.get("/cookie/{cid}/details")
def get_cookie_account_details(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取账号详细信息（包括用户名、密码、显示浏览器设置）"""
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        # 获取详细信息
        details = db_manager.get_cookie_details(cid)
        
        if not details:
            raise HTTPException(status_code=404, detail="账号不存在")
        
        return details
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"获取账号详情失败: {e}")
        raise HTTPException(status_code=400, detail=str(e))


# ========================= 账号密码登录相关接口 =========================

async def _execute_password_login(session_id: str, account_id: str, account: str, password: str, show_browser: bool, user_id: int, current_user: Dict[str, Any]):
    """后台执行账号密码登录任务"""
    try:
        log_with_user('info', f"开始执行账号密码登录任务: {session_id}, 账号: {account_id}", current_user)
        
        # 导入 XianyuSliderStealth
        from utils.xianyu_slider_stealth import XianyuSliderStealth
        
        # 创建 XianyuSliderStealth 实例
        slider_instance = XianyuSliderStealth(
            user_id=account_id,
            enable_learning=True,
            headless=not show_browser
        )
        
        # 更新会话信息
        password_login_sessions[session_id]['slider_instance'] = slider_instance
        
        # 定义通知回调函数，用于检测到人脸认证时返回验证链接或截图（同步函数）
        def notification_callback(message: str, screenshot_path: str = None, verification_url: str = None, screenshot_path_new: str = None):
            """人脸认证通知回调（同步）
            
            Args:
                message: 通知消息
                screenshot_path: 旧版截图路径（兼容参数）
                verification_url: 验证链接
                screenshot_path_new: 新版截图路径（新参数，优先使用）
            """
            try:
                # 优先使用新的截图路径参数
                actual_screenshot_path = screenshot_path_new if screenshot_path_new else screenshot_path
                
                # 优先使用截图路径，如果没有截图则使用验证链接
                if actual_screenshot_path and os.path.exists(actual_screenshot_path):
                    # 更新会话状态，保存截图路径
                    password_login_sessions[session_id]['status'] = 'verification_required'
                    password_login_sessions[session_id]['screenshot_path'] = actual_screenshot_path
                    password_login_sessions[session_id]['verification_url'] = None
                    password_login_sessions[session_id]['qr_code_url'] = None
                    log_with_user('info', f"人脸认证截图已保存: {session_id}, 路径: {actual_screenshot_path}", current_user)
                    
                    # 发送通知到用户配置的渠道
                    def send_face_verification_notification():
                        """在后台线程中发送人脸验证通知"""
                        try:
                            from XianyuAutoAsync import XianyuLive
                            log_with_user('info', f"开始尝试发送人脸验证通知: {account_id}", current_user)
                            
                            # 尝试获取XianyuLive实例（如果账号已经存在）
                            live_instance = XianyuLive.get_instance(account_id)
                            
                            if live_instance:
                                log_with_user('info', f"找到账号实例，准备发送通知: {account_id}", current_user)
                                # 创建新的事件循环来运行异步通知
                                new_loop = asyncio.new_event_loop()
                                asyncio.set_event_loop(new_loop)
                                try:
                                    new_loop.run_until_complete(
                                        live_instance.send_token_refresh_notification(
                                            error_message=message,
                                            notification_type="face_verification",
                                            verification_url=None,
                                            attachment_path=actual_screenshot_path
                                        )
                                    )
                                    log_with_user('info', f"✅ 已发送人脸验证通知: {account_id}", current_user)
                                except Exception as notify_err:
                                    log_with_user('error', f"发送人脸验证通知失败: {str(notify_err)}", current_user)
                                    import traceback
                                    log_with_user('error', f"通知错误详情: {traceback.format_exc()}", current_user)
                                finally:
                                    new_loop.close()
                            else:
                                # 如果账号实例不存在，记录警告并尝试从数据库获取通知配置
                                log_with_user('warning', f"账号实例不存在: {account_id}，尝试从数据库获取通知配置", current_user)
                                try:
                                    # 尝试从数据库获取通知配置
                                    notifications = db_manager.get_account_notifications(account_id)
                                    if notifications:
                                        log_with_user('info', f"找到 {len(notifications)} 个通知配置，但需要账号实例才能发送", current_user)
                                        log_with_user('warning', f"账号实例不存在，无法发送通知: {account_id}。请确保账号已登录并运行中。", current_user)
                                    else:
                                        log_with_user('warning', f"账号 {account_id} 未配置通知渠道", current_user)
                                except Exception as db_err:
                                    log_with_user('error', f"获取通知配置失败: {str(db_err)}", current_user)
                        except Exception as notify_err:
                            log_with_user('error', f"发送人脸验证通知时出错: {str(notify_err)}", current_user)
                            import traceback
                            log_with_user('error', f"通知错误详情: {traceback.format_exc()}", current_user)
                    
                    # 在后台线程中发送通知，避免阻塞登录流程
                    import threading
                    notification_thread = threading.Thread(target=send_face_verification_notification)
                    notification_thread.daemon = True
                    notification_thread.start()
                    log_with_user('info', f"已启动人脸验证通知发送线程: {account_id}", current_user)
                elif verification_url:
                    # 如果没有截图，使用验证链接（兼容旧版本）
                    password_login_sessions[session_id]['status'] = 'verification_required'
                    password_login_sessions[session_id]['verification_url'] = verification_url
                    password_login_sessions[session_id]['screenshot_path'] = None
                    password_login_sessions[session_id]['qr_code_url'] = None
                    log_with_user('info', f"人脸认证验证链接已保存: {session_id}, URL: {verification_url}", current_user)
                    
                    # 发送通知到用户配置的渠道
                    def send_face_verification_notification():
                        """在后台线程中发送人脸验证通知"""
                        try:
                            from XianyuAutoAsync import XianyuLive
                            log_with_user('info', f"开始尝试发送人脸验证通知: {account_id}", current_user)
                            
                            # 尝试获取XianyuLive实例（如果账号已经存在）
                            live_instance = XianyuLive.get_instance(account_id)
                            
                            if live_instance:
                                log_with_user('info', f"找到账号实例，准备发送通知: {account_id}", current_user)
                                # 创建新的事件循环来运行异步通知
                                new_loop = asyncio.new_event_loop()
                                asyncio.set_event_loop(new_loop)
                                try:
                                    new_loop.run_until_complete(
                                        live_instance.send_token_refresh_notification(
                                            error_message=message,
                                            notification_type="face_verification",
                                            verification_url=verification_url
                                        )
                                    )
                                    log_with_user('info', f"✅ 已发送人脸验证通知: {account_id}", current_user)
                                except Exception as notify_err:
                                    log_with_user('error', f"发送人脸验证通知失败: {str(notify_err)}", current_user)
                                    import traceback
                                    log_with_user('error', f"通知错误详情: {traceback.format_exc()}", current_user)
                                finally:
                                    new_loop.close()
                            else:
                                # 如果账号实例不存在，记录警告并尝试从数据库获取通知配置
                                log_with_user('warning', f"账号实例不存在: {account_id}，尝试从数据库获取通知配置", current_user)
                                try:
                                    # 尝试从数据库获取通知配置
                                    notifications = db_manager.get_account_notifications(account_id)
                                    if notifications:
                                        log_with_user('info', f"找到 {len(notifications)} 个通知配置，但需要账号实例才能发送", current_user)
                                        log_with_user('warning', f"账号实例不存在，无法发送通知: {account_id}。请确保账号已登录并运行中。", current_user)
                                    else:
                                        log_with_user('warning', f"账号 {account_id} 未配置通知渠道", current_user)
                                except Exception as db_err:
                                    log_with_user('error', f"获取通知配置失败: {str(db_err)}", current_user)
                        except Exception as notify_err:
                            log_with_user('error', f"发送人脸验证通知时出错: {str(notify_err)}", current_user)
                            import traceback
                            log_with_user('error', f"通知错误详情: {traceback.format_exc()}", current_user)
                    
                    # 在后台线程中发送通知，避免阻塞登录流程
                    import threading
                    notification_thread = threading.Thread(target=send_face_verification_notification)
                    notification_thread.daemon = True
                    notification_thread.start()
                    log_with_user('info', f"已启动人脸验证通知发送线程: {account_id}", current_user)
            except Exception as e:
                log_with_user('error', f"处理人脸认证通知失败: {str(e)}", current_user)
        
        # 调用登录方法（同步方法，需要在后台线程中执行）
        import threading
        
        def run_login():
            try:
                cookies_dict = slider_instance.login_with_password_playwright(
                    account=account,
                    password=password,
                    show_browser=show_browser,
                    notification_callback=notification_callback
                )
                
                if cookies_dict is None:
                    password_login_sessions[session_id]['status'] = 'failed'
                    password_login_sessions[session_id]['error'] = '登录失败，请检查账号密码是否正确'
                    log_with_user('error', f"账号密码登录失败: {account_id}", current_user)
                    return
                
                # 将cookie字典转换为字符串格式
                cookies_str = '; '.join([f"{k}={v}" for k, v in cookies_dict.items()])
                
                log_with_user('info', f"账号密码登录成功，获取到 {len(cookies_dict)} 个Cookie字段: {account_id}", current_user)
                
                # 检查是否已存在相同账号ID的Cookie
                existing_cookies = db_manager.get_all_cookies(user_id)
                is_new_account = account_id not in existing_cookies
                
                # 保存账号密码和Cookie到数据库
                # 使用 update_cookie_account_info 来保存，它会自动处理新账号和现有账号的情况
                update_success = db_manager.update_cookie_account_info(
                    account_id,
                    cookie_value=cookies_str,
                    username=account,
                    password=password,
                    show_browser=show_browser,
                    user_id=user_id  # 新账号时需要提供user_id
                )
                
                if update_success:
                    if is_new_account:
                        log_with_user('info', f"新账号Cookie和账号密码已保存: {account_id}", current_user)
                    else:
                        log_with_user('info', f"现有账号Cookie和账号密码已更新: {account_id}", current_user)
                else:
                    log_with_user('error', f"保存账号信息失败: {account_id}", current_user)
                
                # 账号信息已写入数据库；统一重启监听任务，确保自动回复和订单监听立即生效。
                if cookie_manager.manager:
                    try:
                        cookie_manager.manager.ensure_cookie_task(
                            account_id,
                            cookies_str,
                            user_id=user_id,
                            restart=True,
                            save_to_db=False,
                        )
                        runtime_state = cookie_manager.manager.get_task_state(account_id)
                        log_with_user(
                            'info',
                            f"账号登录后监听任务状态: {account_id} -> {runtime_state}",
                            current_user,
                        )
                    except Exception as task_err:
                        log_with_user(
                            'error',
                            f"账号登录成功但监听任务启动失败: {account_id}, 错误: {str(task_err)}",
                            current_user,
                        )
                        raise RuntimeError("登录凭证已保存，但自动回复和订单监听启动失败") from task_err
                
                # 登录成功后，调用_refresh_cookies_via_browser刷新Cookie
                try:
                    log_with_user('info', f"开始调用_refresh_cookies_via_browser刷新Cookie: {account_id}", current_user)
                    from XianyuAutoAsync import XianyuLive
                    
                    # 创建临时的XianyuLive实例来刷新Cookie
                    temp_xianyu = XianyuLive(
                        cookies_str=cookies_str,
                        cookie_id=account_id,
                        user_id=user_id
                    )
                    
                    # 重置扫码登录Cookie刷新标志，确保账号密码登录后能立即刷新
                    try:
                        temp_xianyu.reset_qr_cookie_refresh_flag()
                        log_with_user('info', f"已重置扫码登录Cookie刷新标志: {account_id}", current_user)
                    except Exception as reset_err:
                        log_with_user('debug', f"重置扫码登录Cookie刷新标志失败（不影响刷新）: {str(reset_err)}", current_user)
                    
                    # 在后台异步执行刷新（不阻塞主流程）
                    async def refresh_cookies_task():
                        try:
                            refresh_success = await temp_xianyu._refresh_cookies_via_browser(triggered_by_refresh_token=False)
                            if refresh_success:
                                log_with_user('info', f"Cookie刷新成功: {account_id}", current_user)
                                # 刷新成功后，从数据库获取更新后的Cookie
                                updated_cookie_info = db_manager.get_cookie_details(account_id)
                                if updated_cookie_info:
                                    refreshed_cookies = updated_cookie_info.get('value', '')
                                    if refreshed_cookies:
                                        # 更新cookie_manager中的Cookie
                                        if cookie_manager.manager:
                                            cookie_manager.manager.update_cookie(account_id, refreshed_cookies, save_to_db=False)
                                        log_with_user('info', f"已更新刷新后的Cookie到cookie_manager: {account_id}", current_user)
                            else:
                                log_with_user('warning', f"Cookie刷新失败或跳过: {account_id}", current_user)
                        except Exception as refresh_e:
                            log_with_user('error', f"刷新Cookie时出错: {account_id}, 错误: {str(refresh_e)}", current_user)
                            import traceback
                            logger.error(traceback.format_exc())
                    
                    # 在后台线程中运行异步任务
                    # 由于run_login是在线程中运行的，需要创建新的事件循环
                    def run_async_refresh():
                        try:
                            import asyncio
                            # 创建新的事件循环
                            new_loop = asyncio.new_event_loop()
                            asyncio.set_event_loop(new_loop)
                            try:
                                new_loop.run_until_complete(refresh_cookies_task())
                            finally:
                                new_loop.close()
                        except Exception as e:
                            log_with_user('error', f"运行异步刷新任务失败: {account_id}, 错误: {str(e)}", current_user)
                    
                    # 在后台线程中执行刷新任务
                    refresh_thread = threading.Thread(target=run_async_refresh, daemon=True)
                    refresh_thread.start()
                    
                except Exception as refresh_err:
                    log_with_user('warning', f"调用_refresh_cookies_via_browser失败: {account_id}, 错误: {str(refresh_err)}", current_user)
                    # 刷新失败不影响登录成功
                
                # 更新会话状态
                password_login_sessions[session_id]['status'] = 'success'
                password_login_sessions[session_id]['account_id'] = account_id
                password_login_sessions[session_id]['is_new_account'] = is_new_account
                password_login_sessions[session_id]['cookie_count'] = len(cookies_dict)
                
            except Exception as e:
                error_msg = str(e)
                password_login_sessions[session_id]['status'] = 'failed'
                password_login_sessions[session_id]['error'] = error_msg
                log_with_user('error', f"账号密码登录失败: {account_id}, 错误: {error_msg}", current_user)
                logger.info(f"会话 {session_id} 状态已更新为 failed，错误消息: {error_msg}")  # 添加日志确认状态更新
                import traceback
                logger.error(traceback.format_exc())
            finally:
                # 清理实例（释放并发槽位）
                try:
                    from utils.xianyu_slider_stealth import concurrency_manager
                    concurrency_manager.unregister_instance(account_id)
                    log_with_user('debug', f"已释放并发槽位: {account_id}", current_user)
                except Exception as cleanup_e:
                    log_with_user('warning', f"清理实例时出错: {str(cleanup_e)}", current_user)
        
        # 在后台线程中执行登录
        login_thread = threading.Thread(target=run_login, daemon=True)
        login_thread.start()
        
    except Exception as e:
        password_login_sessions[session_id]['status'] = 'failed'
        password_login_sessions[session_id]['error'] = str(e)
        log_with_user('error', f"执行账号密码登录任务异常: {str(e)}", current_user)
        import traceback
        logger.error(traceback.format_exc())


@app.post("/password-login")
async def password_login(
    request: Dict[str, Any],
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """账号密码登录接口（异步，支持人脸认证）"""
    try:
        account_id = request.get('account_id')
        account = request.get('account')
        password = request.get('password')
        show_browser = request.get('show_browser', False)
        
        if not account_id or not account or not password:
            return {'success': False, 'message': '账号ID、登录账号和密码不能为空'}
        
        log_with_user('info', f"开始账号密码登录: {account_id}, 账号: {account}", current_user)
        
        # 生成会话ID
        import secrets
        session_id = secrets.token_urlsafe(16)
        
        user_id = current_user['user_id']
        
        # 创建登录会话
        password_login_sessions[session_id] = {
            'account_id': account_id,
            'account': account,
            'password': password,
            'show_browser': show_browser,
            'status': 'processing',
            'verification_url': None,
            'screenshot_path': None,
            'qr_code_url': None,
            'slider_instance': None,
            'task': None,
            'timestamp': time.time(),
            'user_id': user_id
        }
        
        # 启动后台登录任务
        task = asyncio.create_task(_execute_password_login(
            session_id, account_id, account, password, show_browser, user_id, current_user
        ))
        password_login_sessions[session_id]['task'] = task
        
        return {
            'success': True,
            'session_id': session_id,
            'status': 'processing',
            'message': '登录任务已启动，请等待...'
        }
        
    except Exception as e:
        log_with_user('error', f"账号密码登录异常: {str(e)}", current_user)
        import traceback
        logger.error(traceback.format_exc())
        return {'success': False, 'message': f'登录失败: {str(e)}'}


@app.get("/password-login/check/{session_id}")
async def check_password_login_status(
    session_id: str,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """检查账号密码登录状态"""
    try:
        # 清理过期会话（超过1小时）
        current_time = time.time()
        expired_sessions = [
            sid for sid, session in password_login_sessions.items()
            if current_time - session['timestamp'] > 3600
        ]
        for sid in expired_sessions:
            if sid in password_login_sessions:
                del password_login_sessions[sid]
        
        if session_id not in password_login_sessions:
            return {'status': 'not_found', 'message': '会话不存在或已过期'}
        
        session = password_login_sessions[session_id]
        
        # 检查用户权限
        if session['user_id'] != current_user['user_id']:
            return {'status': 'forbidden', 'message': '无权限访问该会话'}
        
        status = session['status']
        
        if status == 'verification_required':
            # 需要人脸认证
            screenshot_path = session.get('screenshot_path')
            verification_url = session.get('verification_url')
            return {
                'status': 'verification_required',
                'verification_url': verification_url,
                'screenshot_path': screenshot_path,
                'qr_code_url': session.get('qr_code_url'),  # 保留兼容性
                'message': '需要人脸验证，请查看验证截图' if screenshot_path else '需要人脸验证，请点击验证链接'
            }
        elif status == 'success':
            # 登录成功
            # 删除截图（如果存在）
            screenshot_path = session.get('screenshot_path')
            if screenshot_path:
                try:
                    from utils.image_utils import image_manager
                    if image_manager.delete_image(screenshot_path):
                        log_with_user('info', f"验证成功后已删除截图: {screenshot_path}", current_user)
                    else:
                        log_with_user('warning', f"删除截图失败: {screenshot_path}", current_user)
                except Exception as e:
                    log_with_user('error', f"删除截图时出错: {str(e)}", current_user)
            
            result = {
                'status': 'success',
                'message': f'账号 {session["account_id"]} 登录成功',
                'account_id': session['account_id'],
                'is_new_account': session.get('is_new_account', False),
                'cookie_count': session.get('cookie_count', 0)
            }
            # 清理会话
            del password_login_sessions[session_id]
            return result
        elif status == 'failed':
            # 登录失败
            # 删除截图（如果存在）
            screenshot_path = session.get('screenshot_path')
            if screenshot_path:
                try:
                    from utils.image_utils import image_manager
                    if image_manager.delete_image(screenshot_path):
                        log_with_user('info', f"验证失败后已删除截图: {screenshot_path}", current_user)
                    else:
                        log_with_user('warning', f"删除截图失败: {screenshot_path}", current_user)
                except Exception as e:
                    log_with_user('error', f"删除截图时出错: {str(e)}", current_user)
            
            error_msg = session.get('error', '登录失败')
            log_with_user('info', f"返回登录失败状态: {session_id}, 错误消息: {error_msg}", current_user)  # 添加日志
            result = {
                'status': 'failed',
                'message': error_msg,
                'error': error_msg  # 也包含error字段，确保前端能获取到
            }
            # 清理会话
            del password_login_sessions[session_id]
            return result
        else:
            # 处理中
            return {
                'status': 'processing',
                'message': '登录处理中，请稍候...'
            }
        
    except Exception as e:
        log_with_user('error', f"检查账号密码登录状态异常: {str(e)}", current_user)
        return {'status': 'error', 'message': str(e)}


# ========================= 人脸验证截图相关接口 =========================

@app.get("/face-verification/screenshot/{account_id}")
async def get_account_face_verification_screenshot(
    account_id: str,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """获取指定账号的人脸验证截图"""
    try:
        import glob
        from datetime import datetime
        
        # 检查账号是否属于当前用户
        user_id = current_user['user_id']
        username = current_user['username']
        
        # 如果是管理员，允许访问所有账号
        is_admin = username == 'admin'
        
        if not is_admin:
            cookie_info = db_manager.get_cookie_details(account_id)
            if not cookie_info:
                log_with_user('warning', f"账号 {account_id} 不存在", current_user)
                return {
                    'success': False,
                    'message': '账号不存在'
                }
            
            cookie_user_id = cookie_info.get('user_id')
            if cookie_user_id != user_id:
                log_with_user('warning', f"用户 {user_id} 尝试访问账号 {account_id}（归属用户: {cookie_user_id}）", current_user)
                return {
                    'success': False,
                    'message': '无权访问该账号'
                }
        
        # 获取该账号的验证截图
        screenshots_dir = os.path.join(static_dir, 'uploads', 'images')
        pattern = os.path.join(screenshots_dir, f'face_verify_{account_id}_*.jpg')
        screenshot_files = glob.glob(pattern)
        
        log_with_user('debug', f"查找截图: {pattern}, 找到 {len(screenshot_files)} 个文件", current_user)
        
        if not screenshot_files:
            log_with_user('warning', f"账号 {account_id} 没有找到验证截图", current_user)
            return {
                'success': False,
                'message': '未找到验证截图'
            }
        
        # 获取最新的截图
        latest_file = max(screenshot_files, key=os.path.getmtime)
        filename = os.path.basename(latest_file)
        stat = os.stat(latest_file)
        
        screenshot_info = {
            'filename': filename,
            'account_id': account_id,
            'path': f'/static/uploads/images/{filename}',
            'size': stat.st_size,
            'created_time': stat.st_ctime,
            'created_time_str': datetime.fromtimestamp(stat.st_ctime).strftime('%Y-%m-%d %H:%M:%S')
        }
        
        log_with_user('info', f"获取账号 {account_id} 的验证截图", current_user)
        
        return {
            'success': True,
            'screenshot': screenshot_info
        }
        
    except Exception as e:
        log_with_user('error', f"获取验证截图失败: {str(e)}", current_user)
        return {
            'success': False,
            'message': str(e)
        }


@app.delete("/face-verification/screenshot/{account_id}")
async def delete_account_face_verification_screenshot(
    account_id: str,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """删除指定账号的人脸验证截图"""
    try:
        import glob
        
        # 检查账号是否属于当前用户
        user_id = current_user['user_id']
        cookie_info = db_manager.get_cookie_details(account_id)
        if not cookie_info or cookie_info.get('user_id') != user_id:
            return {
                'success': False,
                'message': '无权访问该账号'
            }
        
        # 删除该账号的所有验证截图
        screenshots_dir = os.path.join(static_dir, 'uploads', 'images')
        pattern = os.path.join(screenshots_dir, f'face_verify_{account_id}_*.jpg')
        screenshot_files = glob.glob(pattern)
        
        deleted_count = 0
        for file_path in screenshot_files:
            try:
                if os.path.exists(file_path):
                    os.remove(file_path)
                    deleted_count += 1
                    log_with_user('info', f"删除账号 {account_id} 的验证截图: {os.path.basename(file_path)}", current_user)
            except Exception as e:
                log_with_user('error', f"删除截图失败 {file_path}: {str(e)}", current_user)
        
        return {
            'success': True,
            'message': f'已删除 {deleted_count} 个验证截图',
            'deleted_count': deleted_count
        }
        
    except Exception as e:
        log_with_user('error', f"删除验证截图失败: {str(e)}", current_user)
        return {
            'success': False,
            'message': str(e)
        }


# ========================= 扫码登录相关接口 =========================

@app.post("/qr-login/generate")
async def generate_qr_code(current_user: Dict[str, Any] = Depends(get_current_user)):
    """生成扫码登录二维码"""
    try:
        log_with_user('info', "请求生成扫码登录二维码", current_user)

        result = await qr_login_manager.generate_qr_code()

        if result['success']:
            log_with_user('info', f"扫码登录二维码生成成功: {result['session_id']}", current_user)
        else:
            log_with_user('warning', f"扫码登录二维码生成失败: {result.get('message', '未知错误')}", current_user)

        return result

    except Exception as e:
        log_with_user('error', f"生成扫码登录二维码异常: {str(e)}", current_user)
        return {'success': False, 'message': f'生成二维码失败: {str(e)}'}


@app.get("/qr-login/check/{session_id}")
async def check_qr_code_status(session_id: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """检查扫码登录状态"""
    try:
        cleanup_qr_check_records()

        record = qr_check_processed.get(session_id)
        if record:
            return record.get('result', {
                'status': 'processing',
                'message': '已确认，正在准备账号...'
            })

        session_lock = qr_check_locks[session_id]
        if session_lock.locked():
            return {'status': 'processing', 'message': '已确认，正在准备账号...'}

        async with session_lock:
            record = qr_check_processed.get(session_id)
            if record:
                return record.get('result', {
                    'status': 'processing',
                    'message': '已确认，正在准备账号...'
                })

            qr_login_manager.cleanup_expired_sessions()
            status_info = qr_login_manager.get_session_status(session_id)
            log_with_user(
                'debug',
                f"扫码登录会话状态: session={session_id}, status={status_info['status']}",
                current_user
            )
            if status_info['status'] == 'success':
                cookies_info = qr_login_manager.get_session_cookies(session_id)
                if not cookies_info or not cookies_info.get('cookies') or not cookies_info.get('unb'):
                    return {
                        'status': 'error',
                        'message': '扫码已确认，但登录响应缺少账号Cookie，请重新扫码'
                    }

                cookie_fields = sorted(trans_cookies(cookies_info['cookies']).keys())
                log_with_user(
                    'info',
                    f"扫码确认已获取Cookie: 账号={cookies_info.get('unb')}, "
                    f"字段数={len(cookie_fields)}",
                    current_user
                )
                processing_result = {
                    'status': 'processing',
                    'message': '已确认，正在准备账号...'
                }
                qr_check_processed[session_id] = {
                    'processed': False,
                    'timestamp': time.time(),
                    'result': processing_result
                }
                task = asyncio.create_task(
                    _process_qr_login_session(
                        session_id,
                        cookies_info,
                        dict(current_user)
                    )
                )
                qr_check_tasks[session_id] = task
                task.add_done_callback(
                    lambda _, completed_session_id=session_id:
                    qr_check_tasks.pop(completed_session_id, None)
                )
                return processing_result

            return status_info

    except Exception as e:
        log_with_user('error', f"检查扫码登录状态异常: {str(e)}", current_user)
        return {'status': 'error', 'message': str(e)}


async def _process_qr_login_session(
    session_id: str,
    cookies_info: Dict[str, Any],
    current_user: Dict[str, Any]
) -> None:
    """快速保存扫码Cookie，再在后台增强Cookie。"""
    started_at = time.perf_counter()
    try:
        account_info = await process_qr_login_cookies(
            cookies_info['cookies'],
            cookies_info['unb'],
            current_user
        )
        manager_operation = account_info.pop('_manager_operation', None)
        result = {
            'status': 'success',
            'account_ready': True,
            'cookie_refresh_status': 'processing',
            'account_info': account_info,
            'message': '账号已保存，正在后台增强Cookie'
        }
        qr_check_processed[session_id] = {
            'processed': True,
            'timestamp': time.time(),
            'result': result
        }
        log_with_user(
            'info',
            f"扫码登录账号已保存: 账号={account_info.get('account_id', 'unknown')}, "
            f"字段数={account_info.get('cookie_field_count', 0)}, "
            f"耗时={time.perf_counter() - started_at:.2f}s",
            current_user
        )
        await _enhance_qr_login_cookies(
            session_id=session_id,
            account_info=account_info,
            cookies=cookies_info['cookies'],
            current_user=current_user,
            manager_operation=manager_operation,
        )
    except Exception as exc:
        result = {
            'status': 'error',
            'account_ready': False,
            'message': '账号保存失败，请重新扫码'
        }
        log_with_user(
            'error',
            f"扫码登录账号保存失败: 异常类型={type(exc).__name__}, "
            f"耗时={time.perf_counter() - started_at:.2f}s",
            current_user
        )
        qr_check_processed[session_id] = {
            'processed': True,
            'timestamp': time.time(),
            'result': result
        }


async def process_qr_login_cookies(cookies: str, unb: str, current_user: Dict[str, Any]) -> Dict[str, Any]:
    """验证并快速保存扫码Cookie，不等待浏览器增强刷新。"""
    user_id = current_user['user_id']
    cookie_fields = trans_cookies(cookies)
    if not cookie_fields:
        raise ValueError("扫码Cookie为空")

    response_unb = str(unb or "").strip()
    cookie_unb = str(cookie_fields.get('unb') or "").strip()
    if not response_unb and not cookie_unb:
        raise ValueError("扫码响应缺少账号标识")
    if response_unb and cookie_unb and response_unb != cookie_unb:
        raise ValueError("扫码响应账号标识不一致")
    normalized_unb = cookie_unb or response_unb

    existing_cookies = db_manager.get_all_cookies(user_id)
    existing_account_id = None
    for existing_id, cookie_value in existing_cookies.items():
        try:
            if str(trans_cookies(cookie_value).get('unb') or "").strip() == normalized_unb:
                existing_account_id = existing_id
                break
        except Exception:
            continue

    account_id = existing_account_id or normalized_unb
    is_new_account = existing_account_id is None
    if is_new_account:
        counter = 1
        original_account_id = account_id
        while account_id in existing_cookies:
            account_id = f"{original_account_id}_{counter}"
            counter += 1

    if is_new_account:
        saved = db_manager.save_cookie(account_id, cookies, user_id)
    else:
        saved = db_manager.update_cookie_account_info(
            account_id,
            cookie_value=cookies,
            user_id=user_id,
        )
    if not saved:
        raise RuntimeError("扫码Cookie保存失败")

    manager_operation = None
    if cookie_manager.manager:
        if is_new_account:
            manager_operation = cookie_manager.manager.add_cookie(
                account_id,
                cookies,
                user_id=user_id,
            )
        else:
            manager_operation = cookie_manager.manager.update_cookie(
                account_id,
                cookies,
                save_to_db=False,
            )

    log_with_user(
        'info',
        f"扫码Cookie快速持久化完成: 账号={account_id}, "
        f"新账号={is_new_account}, 字段数={len(cookie_fields)}",
        current_user,
    )
    return {
        'account_id': account_id,
        'is_new_account': is_new_account,
        'real_cookie_refreshed': False,
        'cookie_field_count': len(cookie_fields),
        '_manager_operation': manager_operation,
    }


async def _await_cookie_manager_operation(
    operation: Any,
    account_id: str,
    current_user: Dict[str, Any],
) -> None:
    """等待同事件循环中的Manager任务，异常日志不包含敏感值。"""
    if not asyncio.isfuture(operation):
        return
    try:
        await operation
    except asyncio.CancelledError:
        raise
    except Exception as exc:
        log_with_user(
            'warning',
            f"账号任务同步失败: 账号={account_id}, 异常类型={type(exc).__name__}",
            current_user,
        )


async def _enhance_qr_login_cookies(
    session_id: str,
    account_info: Dict[str, Any],
    cookies: str,
    current_user: Dict[str, Any],
    manager_operation: Any = None,
) -> None:
    """后台补全真实Cookie；失败不回滚已经保存的扫码账号。"""
    account_id = account_info['account_id']
    user_id = current_user['user_id']
    started_at = time.perf_counter()
    await _await_cookie_manager_operation(manager_operation, account_id, current_user)

    refresh_success = False
    refresh_error_type = None
    profile_success = False
    profile_error_type = None
    try:
        from XianyuAutoAsync import XianyuLive

        log_with_user(
            'info',
            f"扫码Cookie后台增强开始: 账号={account_id}, "
            f"字段数={account_info.get('cookie_field_count', 0)}",
            current_user,
        )
        temp_instance = XianyuLive(
            cookies_str=cookies,
            cookie_id=account_id,
            user_id=user_id,
        )
        refresh_success = await temp_instance.refresh_cookies_from_qr_login(
            qr_cookies_str=cookies,
            cookie_id=account_id,
            user_id=user_id,
        )
        if refresh_success:
            updated_cookie_info = db_manager.get_cookie_by_id(account_id)
            refreshed_cookies = (
                updated_cookie_info.get('cookies_str')
                if updated_cookie_info
                else None
            )
            if not refreshed_cookies:
                refresh_success = False
            elif cookie_manager.manager:
                operation = cookie_manager.manager.update_cookie(
                    account_id,
                    refreshed_cookies,
                    save_to_db=False,
                )
                await _await_cookie_manager_operation(
                    operation,
                    account_id,
                    current_user,
                )
    except asyncio.CancelledError:
        raise
    except Exception as exc:
        refresh_error_type = type(exc).__name__

    try:
        latest_cookie_info = db_manager.get_cookie_by_id(account_id)
        profile_cookies = (
            latest_cookie_info.get('cookies_str')
            if latest_cookie_info
            else cookies
        )
        if profile_cookies:
            profile_result = await _fetch_and_store_account_profile(
                cookie_id=account_id,
                cookies_str=profile_cookies,
                user_id=user_id,
            )
            profile_success = bool(profile_result.get('success'))
            profile_error_type = profile_result.get('error_type')
    except asyncio.CancelledError:
        raise
    except Exception as exc:
        profile_error_type = type(exc).__name__

    record = qr_check_processed.get(session_id)
    if record:
        result = record.get('result', {})
        result['cookie_refresh_status'] = 'success' if refresh_success else 'warning'
        result['profile_refresh_status'] = 'success' if profile_success else 'warning'
        result['message'] = (
            '账号已保存，Cookie和账号资料补全完成'
            if refresh_success and profile_success
            else '账号已保存，后台资料仍可在账号管理中手动刷新'
        )
        result.setdefault('account_info', {})['real_cookie_refreshed'] = refresh_success
        result.setdefault('account_info', {})['profile_refreshed'] = profile_success
        record['timestamp'] = time.time()

    log_with_user(
        'info' if refresh_success and profile_success else 'warning',
        f"扫码Cookie后台增强结束: 账号={account_id}, 成功={refresh_success}, "
        f"异常类型={refresh_error_type or 'none'}, "
        f"资料成功={profile_success}, "
        f"资料异常类型={profile_error_type or 'none'}, "
        f"耗时={time.perf_counter() - started_at:.2f}s",
        current_user,
    )


@app.post("/qr-login/refresh-cookies")
async def refresh_cookies_from_qr_login(
    request: Dict[str, Any],
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """使用扫码登录获取的cookie访问指定界面获取真实cookie并存入数据库"""
    try:
        qr_cookies = request.get('qr_cookies')
        cookie_id = request.get('cookie_id')

        if not qr_cookies:
            return {'success': False, 'message': '缺少扫码登录cookie'}

        if not cookie_id:
            return {'success': False, 'message': '缺少cookie_id'}

        log_with_user('info', f"开始使用扫码cookie刷新真实cookie: {cookie_id}", current_user)

        # 创建一个临时的XianyuLive实例来执行cookie刷新
        from XianyuAutoAsync import XianyuLive

        # 使用扫码登录的cookie创建临时实例
        temp_instance = XianyuLive(
            cookies_str=qr_cookies,
            cookie_id=cookie_id,
            user_id=current_user['user_id']
        )

        # 执行cookie刷新
        success = await temp_instance.refresh_cookies_from_qr_login(
            qr_cookies_str=qr_cookies,
            cookie_id=cookie_id,
            user_id=current_user['user_id']
        )

        if success:
            log_with_user('info', f"扫码cookie刷新成功: {cookie_id}", current_user)

            # 如果cookie_manager存在，更新其中的cookie
            if cookie_manager.manager:
                # 从数据库获取更新后的cookie
                updated_cookie_info = db_manager.get_cookie_by_id(cookie_id)
                if updated_cookie_info:
                    # refresh_cookies_from_qr_login 已经保存到数据库了，这里不需要再保存
                    cookie_manager.manager.update_cookie(cookie_id, updated_cookie_info['cookies_str'], save_to_db=False)
                    log_with_user('info', f"已更新cookie_manager中的cookie: {cookie_id}", current_user)

            return {
                'success': True,
                'message': '真实cookie获取并保存成功',
                'cookie_id': cookie_id
            }
        else:
            log_with_user('error', f"扫码cookie刷新失败: {cookie_id}", current_user)
            return {'success': False, 'message': '获取真实cookie失败'}

    except Exception as e:
        log_with_user('error', f"扫码cookie刷新异常: {str(e)}", current_user)
        return {'success': False, 'message': f'刷新cookie失败: {str(e)}'}


@app.post("/qr-login/reset-cooldown/{cookie_id}")
async def reset_qr_cookie_refresh_cooldown(
    cookie_id: str,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """重置指定账号的扫码登录Cookie刷新冷却时间"""
    try:
        log_with_user('info', f"重置扫码登录Cookie刷新冷却时间: {cookie_id}", current_user)

        # 检查cookie是否存在
        cookie_info = db_manager.get_cookie_by_id(cookie_id)
        if not cookie_info:
            return {'success': False, 'message': '账号不存在'}

        # 如果cookie_manager中有对应的实例，直接重置
        if cookie_manager.manager and cookie_id in cookie_manager.manager.instances:
            instance = cookie_manager.manager.instances[cookie_id]
            remaining_time_before = instance.get_qr_cookie_refresh_remaining_time()
            instance.reset_qr_cookie_refresh_flag()

            log_with_user('info', f"已重置账号 {cookie_id} 的扫码登录冷却时间，原剩余时间: {remaining_time_before}秒", current_user)

            return {
                'success': True,
                'message': '扫码登录Cookie刷新冷却时间已重置',
                'cookie_id': cookie_id,
                'previous_remaining_time': remaining_time_before
            }
        else:
            # 如果没有活跃实例，返回成功（因为没有冷却时间需要重置）
            log_with_user('info', f"账号 {cookie_id} 没有活跃实例，无需重置冷却时间", current_user)
            return {
                'success': True,
                'message': '账号没有活跃实例，无需重置冷却时间',
                'cookie_id': cookie_id
            }

    except Exception as e:
        log_with_user('error', f"重置扫码登录冷却时间异常: {str(e)}", current_user)
        return {'success': False, 'message': f'重置冷却时间失败: {str(e)}'}


@app.get("/qr-login/cooldown-status/{cookie_id}")
async def get_qr_cookie_refresh_cooldown_status(
    cookie_id: str,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """获取指定账号的扫码登录Cookie刷新冷却状态"""
    try:
        # 检查cookie是否存在
        cookie_info = db_manager.get_cookie_by_id(cookie_id)
        if not cookie_info:
            return {'success': False, 'message': '账号不存在'}

        # 如果cookie_manager中有对应的实例，获取冷却状态
        if cookie_manager.manager and cookie_id in cookie_manager.manager.instances:
            instance = cookie_manager.manager.instances[cookie_id]
            remaining_time = instance.get_qr_cookie_refresh_remaining_time()
            cooldown_duration = instance.qr_cookie_refresh_cooldown
            last_refresh_time = instance.last_qr_cookie_refresh_time

            return {
                'success': True,
                'cookie_id': cookie_id,
                'remaining_time': remaining_time,
                'cooldown_duration': cooldown_duration,
                'last_refresh_time': last_refresh_time,
                'is_in_cooldown': remaining_time > 0,
                'remaining_minutes': remaining_time // 60,
                'remaining_seconds': remaining_time % 60
            }
        else:
            return {
                'success': True,
                'cookie_id': cookie_id,
                'remaining_time': 0,
                'cooldown_duration': 600,  # 默认10分钟
                'last_refresh_time': 0,
                'is_in_cooldown': False,
                'message': '账号没有活跃实例'
            }

    except Exception as e:
        log_with_user('error', f"获取扫码登录冷却状态异常: {str(e)}", current_user)
        return {'success': False, 'message': f'获取冷却状态失败: {str(e)}'}


@app.put('/cookies/{cid}/status')
def update_cookie_status(cid: str, status_data: CookieStatusIn, current_user: Dict[str, Any] = Depends(get_current_user)):
    """更新账号的启用/禁用状态"""
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail='CookieManager 未就绪')
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        cookie_manager.manager.update_cookie_status(cid, status_data.enabled)
        return {'msg': 'status updated', 'enabled': status_data.enabled}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


# ------------------------- 默认回复管理接口 -------------------------

@app.get('/default-replies/{cid}')
def get_default_reply(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取指定账号的默认回复设置"""
    from app.db_manager import db_manager
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限访问该Cookie")

        result = db_manager.get_default_reply(cid)
        if result is None:
            # 如果没有设置，返回默认值
            return {'enabled': False, 'reply_content': '', 'reply_once': False}
        return result
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.put('/default-replies/{cid}')
def update_default_reply(cid: str, reply_data: DefaultReplyIn, current_user: Dict[str, Any] = Depends(get_current_user)):
    """更新指定账号的默认回复设置"""
    from app.db_manager import db_manager
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        db_manager.save_default_reply(cid, reply_data.enabled, reply_data.reply_content, reply_data.reply_once, reply_data.reply_image_url)
        return {'msg': 'default reply updated', 'enabled': reply_data.enabled, 'reply_once': reply_data.reply_once, 'reply_image_url': reply_data.reply_image_url}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get('/default-replies')
def get_all_default_replies(current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取当前用户所有账号的默认回复设置"""
    from app.db_manager import db_manager
    try:
        # 只返回当前用户的默认回复设置
        user_id = current_user['user_id']
        user_cookies = db_manager.get_all_cookies(user_id)

        all_replies = db_manager.get_all_default_replies()
        # 过滤只属于当前用户的回复设置
        user_replies = {cid: reply for cid, reply in all_replies.items() if cid in user_cookies}
        return user_replies
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.delete('/default-replies/{cid}')
def delete_default_reply(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """删除指定账号的默认回复设置"""
    from app.db_manager import db_manager
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        success = db_manager.delete_default_reply(cid)
        if success:
            return {'msg': 'default reply deleted'}
        else:
            raise HTTPException(status_code=400, detail='删除失败')
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post('/default-replies/{cid}/clear-records')
def clear_default_reply_records(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """清空指定账号的默认回复记录"""
    from app.db_manager import db_manager
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        db_manager.clear_default_reply_records(cid)
        return {'msg': 'default reply records cleared'}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ------------------------- 默认回复管理接口（单数形式兼容路由） -------------------------
# 兼容前端使用 /api/default-reply/ 的请求

@app.get('/api/default-reply/{cid}')
def get_default_reply_compat(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取指定账号的默认回复设置（兼容路由）"""
    return get_default_reply(cid, current_user)


@app.put('/api/default-reply/{cid}')
def update_default_reply_compat(cid: str, reply_data: DefaultReplyIn, current_user: Dict[str, Any] = Depends(get_current_user)):
    """更新指定账号的默认回复设置（兼容路由）"""
    return update_default_reply(cid, reply_data, current_user)


@app.delete('/api/default-reply/{cid}')
def delete_default_reply_compat(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """删除指定账号的默认回复设置（兼容路由）"""
    return delete_default_reply(cid, current_user)


@app.post('/api/default-reply/{cid}/clear-records')
def clear_default_reply_records_compat(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """清空指定账号的默认回复记录（兼容路由）"""
    return clear_default_reply_records(cid, current_user)


# ------------------------- 通知渠道管理接口 -------------------------

@app.get('/notification-channels')
def get_notification_channels(current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取所有通知渠道"""
    from app.db_manager import db_manager
    try:
        user_id = current_user['user_id']
        return db_manager.get_notification_channels(user_id)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post('/notification-channels')
def create_notification_channel(channel_data: NotificationChannelIn, current_user: Dict[str, Any] = Depends(get_current_user)):
    """创建通知渠道"""
    from app.db_manager import db_manager
    try:
        user_id = current_user['user_id']
        name, channel_type, config = validate_notification_channel(
            channel_data.name,
            channel_data.type,
            channel_data.config,
        )
        channel_id = db_manager.create_notification_channel(
            name,
            channel_type,
            config,
            user_id
        )
        return {'msg': 'notification channel created', 'id': channel_id}
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.get('/notification-channels/{channel_id}')
def get_notification_channel(channel_id: int, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取指定通知渠道"""
    from app.db_manager import db_manager
    try:
        channel = db_manager.get_notification_channel(channel_id, current_user['user_id'])
        if not channel:
            raise HTTPException(status_code=404, detail='通知渠道不存在')
        return channel
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.put('/notification-channels/{channel_id}')
def update_notification_channel(channel_id: int, channel_data: NotificationChannelUpdate,
                                current_user: Dict[str, Any] = Depends(get_current_user)):
    """更新通知渠道"""
    from app.db_manager import db_manager
    try:
        existing_channel = db_manager.get_notification_channel(channel_id, current_user['user_id'])
        if not existing_channel:
            raise HTTPException(status_code=404, detail='通知渠道不存在')

        name, _, config = validate_notification_channel(
            channel_data.name if channel_data.name is not None else existing_channel["name"],
            existing_channel["type"],
            channel_data.config if channel_data.config is not None else existing_channel["config"],
        )
        success = db_manager.update_notification_channel(
            channel_id,
            name,
            config,
            channel_data.enabled if channel_data.enabled is not None else existing_channel["enabled"],
            current_user['user_id']
        )
        if success:
            return {'msg': 'notification channel updated'}
        else:
            raise HTTPException(status_code=404, detail='通知渠道不存在')
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.delete('/notification-channels/{channel_id}')
def delete_notification_channel(channel_id: int, current_user: Dict[str, Any] = Depends(get_current_user)):
    """删除通知渠道"""
    from app.db_manager import db_manager
    try:
        success = db_manager.delete_notification_channel(channel_id, current_user['user_id'])
        if success:
            return {'msg': 'notification channel deleted'}
        else:
            raise HTTPException(status_code=404, detail='通知渠道不存在')
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ------------------------- 消息通知配置接口 -------------------------

@app.get('/message-notifications')
def get_all_message_notifications(current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取当前用户所有账号的消息通知配置"""
    from app.db_manager import db_manager
    try:
        # 只返回当前用户的消息通知配置
        user_id = current_user['user_id']
        return db_manager.get_all_message_notifications(user_id)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get('/message-notifications/{cid}')
def get_account_notifications(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取指定账号的消息通知配置"""
    from app.db_manager import db_manager
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限访问该Cookie")

        return db_manager.get_account_notifications(cid, user_id)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post('/message-notifications/{cid}')
def set_message_notification(cid: str, notification_data: MessageNotificationIn, current_user: Dict[str, Any] = Depends(get_current_user)):
    """设置账号的消息通知"""
    from app.db_manager import db_manager
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        # 检查通知渠道是否存在
        channel = db_manager.get_notification_channel(notification_data.channel_id, user_id)
        if not channel:
            raise HTTPException(status_code=404, detail='通知渠道不存在')

        success = db_manager.set_message_notification(cid, notification_data.channel_id, notification_data.enabled)
        if success:
            return {'msg': 'message notification set'}
        else:
            raise HTTPException(status_code=400, detail='设置失败')
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.delete('/message-notifications/account/{cid}')
def delete_account_notifications(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """删除账号的所有消息通知配置"""
    from app.db_manager import db_manager
    try:
        success = db_manager.delete_account_notifications(cid, current_user['user_id'])
        if success:
            return {'msg': 'account notifications deleted'}
        else:
            raise HTTPException(status_code=404, detail='账号通知配置不存在')
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.delete('/message-notifications/{notification_id}')
def delete_message_notification(notification_id: int, current_user: Dict[str, Any] = Depends(get_current_user)):
    """删除消息通知配置"""
    from app.db_manager import db_manager
    try:
        success = db_manager.delete_message_notification(notification_id, current_user['user_id'])
        if success:
            return {'msg': 'message notification deleted'}
        else:
            raise HTTPException(status_code=404, detail='通知配置不存在')
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ------------------------- 消息过滤与回复决策日志 -------------------------

@app.get('/message-filters')
def list_message_filters(
    cookie_id: str = None,
    filter_type: str = None,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    normalized_type = (filter_type or "").strip().lower() or None
    if normalized_type and normalized_type not in MESSAGE_FILTER_TYPES:
        raise HTTPException(status_code=400, detail="无效的过滤类型")
    if cookie_id and cookie_id not in db_manager.get_all_cookies(current_user["user_id"]):
        raise HTTPException(status_code=403, detail="无权限访问该账号")
    return {
        "success": True,
        "data": db_manager.get_message_filters(
            current_user["user_id"],
            cookie_id=(cookie_id or "").strip() or None,
            filter_type=normalized_type,
        ),
    }


@app.post('/message-filters')
def create_message_filter(
    data: MessageFilterIn,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    try:
        cookie_id, keyword, filter_type = validate_message_filter(
            data.cookie_id, data.keyword, data.filter_type
        )
        filter_id = db_manager.create_message_filter(
            cookie_id, keyword, filter_type, current_user["user_id"], data.enabled
        )
        return {
            "success": True,
            "data": db_manager.get_message_filter(filter_id, current_user["user_id"]),
        }
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail=str(exc))
    except sqlite3.IntegrityError:
        raise HTTPException(status_code=409, detail="相同账号、关键词和类型的规则已存在")
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))


@app.post('/message-filters/batch-create')
def batch_create_message_filters(
    data: MessageFilterBatchIn,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    unique_keywords = []
    seen = set()
    for raw_keyword in data.keywords:
        cookie_id, keyword, filter_type = validate_message_filter(
            data.cookie_id, raw_keyword, data.filter_type
        )
        normalized_key = keyword.casefold()
        if normalized_key not in seen:
            seen.add(normalized_key)
            unique_keywords.append(keyword)
    if not unique_keywords:
        raise HTTPException(status_code=400, detail="至少提供一个过滤关键词")

    created = 0
    skipped = 0
    for keyword in unique_keywords:
        try:
            db_manager.create_message_filter(
                cookie_id, keyword, filter_type, current_user["user_id"], data.enabled
            )
            created += 1
        except PermissionError as exc:
            raise HTTPException(status_code=403, detail=str(exc))
        except sqlite3.IntegrityError:
            skipped += 1
    return {"success": True, "created": created, "skipped": skipped}


@app.put('/message-filters/{filter_id}')
def update_message_filter(
    filter_id: int,
    data: MessageFilterUpdate,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    existing = db_manager.get_message_filter(filter_id, current_user["user_id"])
    if not existing:
        raise HTTPException(status_code=404, detail="过滤规则不存在")
    try:
        _, keyword, filter_type = validate_message_filter(
            existing["cookie_id"],
            data.keyword if data.keyword is not None else existing["keyword"],
            data.filter_type if data.filter_type is not None else existing["filter_type"],
        )
        db_manager.update_message_filter(
            filter_id,
            current_user["user_id"],
            keyword=keyword,
            filter_type=filter_type,
            enabled=data.enabled,
        )
        return {
            "success": True,
            "data": db_manager.get_message_filter(filter_id, current_user["user_id"]),
        }
    except sqlite3.IntegrityError:
        raise HTTPException(status_code=409, detail="相同账号、关键词和类型的规则已存在")
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))


@app.put('/message-filters/{filter_id}/toggle')
def toggle_message_filter(
    filter_id: int,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    existing = db_manager.get_message_filter(filter_id, current_user["user_id"])
    if not existing:
        raise HTTPException(status_code=404, detail="过滤规则不存在")
    db_manager.update_message_filter(
        filter_id,
        current_user["user_id"],
        enabled=not existing["enabled"],
    )
    return {
        "success": True,
        "data": db_manager.get_message_filter(filter_id, current_user["user_id"]),
    }


@app.delete('/message-filters/{filter_id}')
def delete_message_filter(
    filter_id: int,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    if not db_manager.delete_message_filter(filter_id, current_user["user_id"]):
        raise HTTPException(status_code=404, detail="过滤规则不存在")
    return {"success": True}


@app.post('/message-filters/batch-delete')
def batch_delete_message_filters(
    data: MessageFilterBatchDelete,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    ids = list(dict.fromkeys(filter_id for filter_id in data.ids if filter_id > 0))
    if not ids:
        raise HTTPException(status_code=400, detail="请选择要删除的过滤规则")
    return {
        "success": True,
        "deleted": db_manager.delete_message_filters(ids, current_user["user_id"]),
    }


@app.get('/auto-reply-logs')
def list_auto_reply_logs(
    cookie_id: str = None,
    process_status: str = None,
    reply_strategy: str = None,
    send_status: str = None,
    keyword: str = None,
    page: int = 1,
    page_size: int = 20,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    process_status = (process_status or "").strip() or None
    reply_strategy = (reply_strategy or "").strip() or None
    send_status = (send_status or "").strip() or None
    if process_status not in (None, "success", "skipped", "failed"):
        raise HTTPException(status_code=400, detail="无效的处理状态")
    if reply_strategy not in (None, "keyword", "ai", "default", "api", "none"):
        raise HTTPException(status_code=400, detail="无效的回复策略")
    if send_status not in (None, "success", "failed", "unknown"):
        raise HTTPException(status_code=400, detail="无效的发送状态")
    if cookie_id and cookie_id not in db_manager.get_all_cookies(current_user["user_id"]):
        raise HTTPException(status_code=403, detail="无权限访问该账号")

    page = max(1, page)
    page_size = max(1, min(page_size, 100))
    filters = {
        "cookie_id": (cookie_id or "").strip() or None,
        "process_status": process_status,
        "reply_strategy": reply_strategy,
        "send_status": send_status,
        "keyword": (keyword or "").strip() or None,
    }
    total = db_manager.get_auto_reply_logs_count(current_user["user_id"], **filters)
    data = db_manager.get_auto_reply_logs(
        current_user["user_id"],
        limit=page_size,
        offset=(page - 1) * page_size,
        **filters,
    )
    return {
        "success": True,
        "data": data,
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_pages": max(1, (total + page_size - 1) // page_size),
    }


# ------------------------- 系统设置接口 -------------------------

@app.get('/system-settings/public')
def get_public_system_settings():
    """获取公开的系统设置（无需认证）"""
    from app.db_manager import db_manager
    try:
        all_settings = db_manager.get_all_system_settings()
        # 只返回公开的配置项
        public_keys = {
            "registration_enabled",
            "show_default_login_info",
            "login_captcha_enabled",
            # 注册表单要据此决定是否显示验证码输入框
            "email_verification_enabled",
        }
        result = {k: v for k, v in all_settings.items() if k in public_keys}
        # 没写过这项的老库按开启处理，与后端校验逻辑保持一致
        result.setdefault("email_verification_enabled", "true")
        return result
    except Exception as e:
        logger.error(f"获取公开系统设置失败: {e}")
        # 返回默认值
        return {
            "registration_enabled": "true",
            "show_default_login_info": "true",
            "login_captcha_enabled": "true",
            "email_verification_enabled": "true"
        }


@app.get('/system-settings')
def get_system_settings(_: Dict[str, Any] = Depends(require_admin)):
    """获取系统设置（排除敏感信息）"""
    from app.db_manager import db_manager
    try:
        settings = db_manager.get_all_system_settings()
        # 移除敏感信息
        if 'admin_password_hash' in settings:
            del settings['admin_password_hash']
        return settings
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.put('/system-settings/{key}')
def update_system_setting(key: str, setting_data: SystemSettingIn,
                          _: Dict[str, Any] = Depends(require_admin)):
    """更新系统设置"""
    from app.db_manager import db_manager
    try:
        # 禁止直接修改密码哈希
        if key == 'admin_password_hash':
            raise HTTPException(status_code=400, detail='请使用密码修改接口')

        success = db_manager.set_system_setting(key, setting_data.value, setting_data.description)
        if success:
            return {'msg': 'system setting updated'}
        else:
            raise HTTPException(status_code=400, detail='更新失败')
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ------------------------- 注册设置接口 -------------------------

@app.get('/registration-status')
def get_registration_status():
    """获取注册开关状态（公开接口，无需认证）"""
    from app.db_manager import db_manager
    try:
        enabled_str = db_manager.get_system_setting('registration_enabled')
        logger.info(f"从数据库获取的注册设置值: '{enabled_str}'")  # 调试信息

        # 如果设置不存在，默认为开启
        if enabled_str is None:
            enabled_bool = True
            message = '注册功能已开启'
        else:
            enabled_bool = enabled_str == 'true'
            message = '注册功能已开启' if enabled_bool else '注册功能已关闭'

        logger.info(f"解析后的注册状态: enabled={enabled_bool}, message='{message}'")  # 调试信息

        return {
            'enabled': enabled_bool,
            'message': message
        }
    except Exception as e:
        logger.error(f"获取注册状态失败: {e}")
        return {'enabled': True, 'message': '注册功能已开启'}  # 出错时默认开启


@app.get('/login-info-status')
def get_login_info_status():
    """获取默认登录信息显示状态（公开接口，无需认证）"""
    from app.db_manager import db_manager
    try:
        enabled_str = db_manager.get_system_setting('show_default_login_info')
        logger.debug(f"从数据库获取的登录信息显示设置值: '{enabled_str}'")

        # 如果设置不存在，默认为开启
        if enabled_str is None:
            enabled_bool = True
        else:
            enabled_bool = enabled_str == 'true'

        return {"enabled": enabled_bool}
    except Exception as e:
        logger.error(f"获取登录信息显示状态失败: {e}")
        # 出错时默认为开启
        return {"enabled": True}


class RegistrationSettingUpdate(BaseModel):
    enabled: bool


class LoginInfoSettingUpdate(BaseModel):
    enabled: bool


@app.put('/registration-settings')
def update_registration_settings(setting_data: RegistrationSettingUpdate, admin_user: Dict[str, Any] = Depends(require_admin)):
    """更新注册开关设置（仅管理员）"""
    from app.db_manager import db_manager
    try:
        enabled = setting_data.enabled
        success = db_manager.set_system_setting(
            'registration_enabled',
            'true' if enabled else 'false',
            '是否开启用户注册'
        )
        if success:
            log_with_user('info', f"更新注册设置: {'开启' if enabled else '关闭'}", admin_user)
            return {
                'success': True,
                'enabled': enabled,
                'message': f"注册功能已{'开启' if enabled else '关闭'}"
            }
        else:
            raise HTTPException(status_code=500, detail='更新注册设置失败')
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"更新注册设置失败: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@app.put('/login-info-settings')
def update_login_info_settings(setting_data: LoginInfoSettingUpdate, admin_user: Dict[str, Any] = Depends(require_admin)):
    """更新默认登录信息显示设置（仅管理员）"""
    from app.db_manager import db_manager
    try:
        enabled = setting_data.enabled
        success = db_manager.set_system_setting(
            'show_default_login_info',
            'true' if enabled else 'false',
            '是否显示默认登录信息'
        )
        if success:
            log_with_user('info', f"更新登录信息显示设置: {'开启' if enabled else '关闭'}", admin_user)
            return {
                'success': True,
                'enabled': enabled,
                'message': f"默认登录信息显示已{'开启' if enabled else '关闭'}"
            }
        else:
            raise HTTPException(status_code=500, detail='更新登录信息显示设置失败')
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"更新登录信息显示设置失败: {e}")
        raise HTTPException(status_code=500, detail=str(e))




@app.delete("/cookies/{cid}")
def remove_cookie(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        cookie_manager.manager.remove_cookie(cid)
        return {"msg": "removed"}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


class AutoConfirmUpdate(BaseModel):
    auto_confirm: bool


class RemarkUpdate(BaseModel):
    remark: str


class PauseDurationUpdate(BaseModel):
    pause_duration: int


@app.put("/cookies/{cid}/auto-confirm")
def update_auto_confirm(cid: str, update_data: AutoConfirmUpdate, current_user: Dict[str, Any] = Depends(get_current_user)):
    """更新账号的自动确认发货设置"""
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        # 更新数据库中的auto_confirm设置
        success = db_manager.update_auto_confirm(cid, update_data.auto_confirm)
        if not success:
            raise HTTPException(status_code=500, detail="更新自动确认发货设置失败")

        # 通知CookieManager更新设置（如果账号正在运行）
        if hasattr(cookie_manager.manager, 'update_auto_confirm_setting'):
            cookie_manager.manager.update_auto_confirm_setting(cid, update_data.auto_confirm)

        return {
            "msg": "success",
            "auto_confirm": update_data.auto_confirm,
            "message": f"自动确认发货已{'开启' if update_data.auto_confirm else '关闭'}"
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/cookies/{cid}/auto-confirm")
def get_auto_confirm(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取账号的自动确认发货设置"""
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        # 获取auto_confirm设置
        auto_confirm = db_manager.get_auto_confirm(cid)
        return {
            "auto_confirm": auto_confirm,
            "message": f"自动确认发货当前{'开启' if auto_confirm else '关闭'}"
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.put("/cookies/{cid}/remark")
def update_cookie_remark(cid: str, update_data: RemarkUpdate, current_user: Dict[str, Any] = Depends(get_current_user)):
    """更新账号备注"""
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        # 更新备注
        success = db_manager.update_cookie_remark(cid, update_data.remark)
        if success:
            log_with_user('info', f"更新账号备注: {cid} -> {update_data.remark}", current_user)
            return {
                "message": "备注更新成功",
                "remark": update_data.remark
            }
        else:
            raise HTTPException(status_code=500, detail="备注更新失败")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/cookies/{cid}/remark")
def get_cookie_remark(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取账号备注"""
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        # 获取Cookie详细信息（包含备注）
        cookie_details = db_manager.get_cookie_details(cid)
        if cookie_details:
            return {
                "remark": cookie_details.get('remark', ''),
                "message": "获取备注成功"
            }
        else:
            raise HTTPException(status_code=404, detail="账号不存在")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.put("/cookies/{cid}/pause-duration")
def update_cookie_pause_duration(cid: str, update_data: PauseDurationUpdate, current_user: Dict[str, Any] = Depends(get_current_user)):
    """更新账号自动回复暂停时间"""
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        # 验证暂停时间范围（0-120分钟，0表示不暂停）
        if not (0 <= update_data.pause_duration <= 120):
            raise HTTPException(status_code=400, detail="暂停时间必须在0-120分钟之间（0表示不暂停）")

        # 更新暂停时间
        success = db_manager.update_cookie_pause_duration(cid, update_data.pause_duration)
        if success:
            log_with_user('info', f"更新账号自动回复暂停时间: {cid} -> {update_data.pause_duration}分钟", current_user)
            return {
                "message": "暂停时间更新成功",
                "pause_duration": update_data.pause_duration
            }
        else:
            raise HTTPException(status_code=500, detail="暂停时间更新失败")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/cookies/{cid}/pause-duration")
def get_cookie_pause_duration(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取账号自动回复暂停时间"""
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cid not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        # 获取暂停时间
        pause_duration = db_manager.get_cookie_pause_duration(cid)
        return {
            "pause_duration": pause_duration,
            "message": "获取暂停时间成功"
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


class KeywordIn(BaseModel):
    keywords: Dict[str, str]  # key -> reply

class KeywordWithItemIdIn(BaseModel):
    keywords: List[Dict[str, Any]]  # [{"keyword": str, "reply": str, "item_id": str}]


@app.get("/keywords/{cid}")
def get_keywords(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")

    # 检查cookie是否属于当前用户
    user_id = current_user['user_id']
    from app.db_manager import db_manager
    user_cookies = db_manager.get_all_cookies(user_id)

    if cid not in user_cookies:
        raise HTTPException(status_code=403, detail="无权限访问该Cookie")

    # 直接从数据库获取所有关键词（避免重复计算）
    item_keywords = db_manager.get_keywords_with_item_id(cid)

    # 转换为统一格式
    all_keywords = []
    for keyword, reply, item_id in item_keywords:
        all_keywords.append({
            "keyword": keyword,
            "reply": reply,
            "item_id": item_id,
            "type": "item" if item_id else "normal"
        })

    return all_keywords


@app.get("/keywords-with-item-id/{cid}")
def get_keywords_with_item_id(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取包含商品ID的关键词列表"""
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")

    # 检查cookie是否属于当前用户
    user_id = current_user['user_id']
    from app.db_manager import db_manager
    user_cookies = db_manager.get_all_cookies(user_id)

    if cid not in user_cookies:
        raise HTTPException(status_code=403, detail="无权限访问该Cookie")

    # 获取包含类型信息的关键词
    keywords = db_manager.get_keywords_with_type(cid)

    # 转换为前端需要的格式
    result = []
    for keyword_data in keywords:
        result.append({
            "keyword": keyword_data['keyword'],
            "reply": keyword_data['reply'],
            "item_id": keyword_data['item_id'] or "",
            "type": keyword_data['type'],
            "image_url": keyword_data['image_url']
        })

    return result


@app.post("/keywords/{cid}")
def update_keywords(cid: str, body: KeywordIn, current_user: Dict[str, Any] = Depends(get_current_user)):
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")

    # 检查cookie是否属于当前用户
    user_id = current_user['user_id']
    from app.db_manager import db_manager
    user_cookies = db_manager.get_all_cookies(user_id)

    if cid not in user_cookies:
        log_with_user('warning', f"尝试操作其他用户的Cookie关键字: {cid}", current_user)
        raise HTTPException(status_code=403, detail="无权限操作该Cookie")

    kw_list = [(k, v) for k, v in body.keywords.items()]
    log_with_user('info', f"更新Cookie关键字: {cid}, 数量: {len(kw_list)}", current_user)

    cookie_manager.manager.update_keywords(cid, kw_list)
    log_with_user('info', f"Cookie关键字更新成功: {cid}", current_user)
    return {"msg": "updated", "count": len(kw_list)}


@app.post("/keywords-with-item-id/{cid}")
def update_keywords_with_item_id(cid: str, body: KeywordWithItemIdIn, current_user: Dict[str, Any] = Depends(get_current_user)):
    """更新包含商品ID的关键词列表"""
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")

    # 检查cookie是否属于当前用户
    user_id = current_user['user_id']
    from app.db_manager import db_manager
    user_cookies = db_manager.get_all_cookies(user_id)

    if cid not in user_cookies:
        log_with_user('warning', f"尝试操作其他用户的Cookie关键字: {cid}", current_user)
        raise HTTPException(status_code=403, detail="无权限操作该Cookie")

    # 验证数据格式
    keywords_to_save = []
    keyword_set = set()  # 用于检查当前提交的关键词中是否有重复

    for kw_data in body.keywords:
        keyword = kw_data.get('keyword', '').strip()
        reply = kw_data.get('reply', '').strip()
        item_id = kw_data.get('item_id', '').strip() or None

        if not keyword:
            raise HTTPException(status_code=400, detail="关键词不能为空")

        # 检查当前提交的关键词中是否有重复
        keyword_key = f"{keyword}|{item_id or ''}"
        if keyword_key in keyword_set:
            item_id_text = f"（商品ID: {item_id}）" if item_id else "（通用关键词）"
            raise HTTPException(status_code=400, detail=f"关键词 '{keyword}' {item_id_text} 在当前提交中重复")
        keyword_set.add(keyword_key)

        keywords_to_save.append((keyword, reply, item_id))

    # 保存关键词（只保存文本关键词，保留图片关键词）
    try:
        success = db_manager.save_text_keywords_only(cid, keywords_to_save)
        if not success:
            raise HTTPException(status_code=500, detail="保存关键词失败")
    except Exception as e:
        error_msg = str(e)

        # 检查是否是图片关键词冲突
        if "已存在（图片关键词）" in error_msg:
            # 直接使用数据库管理器提供的友好错误信息
            raise HTTPException(status_code=400, detail=error_msg)
        elif "UNIQUE constraint failed" in error_msg or "唯一约束冲突" in error_msg:
            # 尝试从错误信息中提取具体的冲突关键词
            conflict_keyword = None
            conflict_type = None

            # 检查是否是数据库管理器抛出的详细错误
            if "关键词唯一约束冲突" in error_msg:
                # 解析详细错误信息：关键词唯一约束冲突: Cookie=xxx, 关键词='xxx', 通用关键词/商品ID: xxx
                import re
                keyword_match = re.search(r"关键词='([^']+)'", error_msg)
                if keyword_match:
                    conflict_keyword = keyword_match.group(1)

                if "通用关键词" in error_msg:
                    conflict_type = "通用关键词"
                elif "商品ID:" in error_msg:
                    item_match = re.search(r"商品ID: ([^\s,]+)", error_msg)
                    if item_match:
                        conflict_type = f"商品关键词（商品ID: {item_match.group(1)}）"

            # 构造用户友好的错误信息
            if conflict_keyword and conflict_type:
                detail_msg = f'关键词 "{conflict_keyword}" （{conflict_type}） 已存在，请使用其他关键词或商品ID'
            elif "keywords.cookie_id, keywords.keyword" in error_msg:
                detail_msg = "关键词重复！该关键词已存在（可能是图片关键词或文本关键词），请使用其他关键词"
            else:
                detail_msg = "关键词重复！请使用不同的关键词或商品ID组合"

            raise HTTPException(status_code=400, detail=detail_msg)
        else:
            log_with_user('error', f"保存关键词时发生未知错误: {error_msg}", current_user)
            raise HTTPException(status_code=500, detail="保存关键词失败")

    log_with_user('info', f"更新Cookie关键字(含商品ID): {cid}, 数量: {len(keywords_to_save)}", current_user)
    return {"msg": "updated", "count": len(keywords_to_save)}


@app.get("/items/{cid}")
def get_items_list(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取指定账号的商品列表"""
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")

    # 检查cookie是否属于当前用户
    user_id = current_user['user_id']
    from app.db_manager import db_manager
    user_cookies = db_manager.get_all_cookies(user_id)

    if cid not in user_cookies:
        raise HTTPException(status_code=403, detail="无权限访问该Cookie")

    try:
        # 获取该账号的所有商品
        with db_manager.lock:
            cursor = db_manager.conn.cursor()
            cursor.execute('''
            SELECT item_id, item_title, item_price, created_at
            FROM item_info
            WHERE cookie_id = ?
            ORDER BY created_at DESC
            ''', (cid,))

            items = []
            for row in cursor.fetchall():
                items.append({
                    'item_id': row[0],
                    'item_title': row[1] or '未知商品',
                    'item_price': row[2] or '价格未知',
                    'created_at': row[3]
                })

            return {"items": items, "count": len(items)}

    except Exception as e:
        logger.error(f"获取商品列表失败: {e}")
        raise HTTPException(status_code=500, detail="获取商品列表失败")


@app.get("/keywords-export/{cid}")
def export_keywords(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """导出指定账号的关键词为Excel文件"""
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")

    # 检查cookie是否属于当前用户
    user_id = current_user['user_id']
    from app.db_manager import db_manager
    user_cookies = db_manager.get_all_cookies(user_id)

    if cid not in user_cookies:
        raise HTTPException(status_code=403, detail="无权限访问该Cookie")

    try:
        # 获取关键词数据（包含类型信息）
        keywords = db_manager.get_keywords_with_type(cid)

        # 创建DataFrame，只导出文本类型的关键词
        data = []
        for keyword_data in keywords:
            # 只导出文本类型的关键词
            if keyword_data.get('type', 'text') == 'text':
                data.append({
                    '关键词': keyword_data['keyword'],
                    '商品ID': keyword_data['item_id'] or '',
                    '关键词内容': keyword_data['reply']
                })

        # 如果没有数据，创建空的DataFrame但保留列名（作为模板）
        if not data:
            df = pd.DataFrame(columns=['关键词', '商品ID', '关键词内容'])
        else:
            df = pd.DataFrame(data)

        # 创建Excel文件
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='关键词数据', index=False)

            # 如果是空模板，添加一些示例说明
            if data == []:
                worksheet = writer.sheets['关键词数据']
                # 添加示例数据作为注释（从第2行开始）
                worksheet['A2'] = '你好'
                worksheet['B2'] = ''
                worksheet['C2'] = '您好！欢迎咨询，有什么可以帮助您的吗？'

                worksheet['A3'] = '价格'
                worksheet['B3'] = '123456'
                worksheet['C3'] = '这个商品的价格是99元，现在有优惠活动哦！'

                worksheet['A4'] = '发货'
                worksheet['B4'] = ''
                worksheet['C4'] = '我们会在24小时内发货，请耐心等待。'

                # 设置示例行的样式（浅灰色背景）
                from openpyxl.styles import PatternFill
                gray_fill = PatternFill(start_color='F0F0F0', end_color='F0F0F0', fill_type='solid')
                for row in range(2, 5):
                    for col in range(1, 4):
                        worksheet.cell(row=row, column=col).fill = gray_fill

        output.seek(0)

        # 生成文件名（使用URL编码处理中文）
        from urllib.parse import quote
        if not data:
            filename = f"keywords_template_{cid}_{int(time.time())}.xlsx"
        else:
            filename = f"keywords_{cid}_{int(time.time())}.xlsx"
        encoded_filename = quote(filename.encode('utf-8'))

        # 返回文件
        return StreamingResponse(
            io.BytesIO(output.read()),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"
            }
        )

    except Exception as e:
        logger.error(f"导出关键词失败: {e}")
        raise HTTPException(status_code=500, detail=f"导出关键词失败: {str(e)}")


@app.post("/keywords-import/{cid}")
async def import_keywords(cid: str, file: UploadFile = File(...), current_user: Dict[str, Any] = Depends(get_current_user)):
    """导入Excel文件中的关键词到指定账号"""
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")

    # 检查cookie是否属于当前用户
    user_id = current_user['user_id']
    from app.db_manager import db_manager
    user_cookies = db_manager.get_all_cookies(user_id)

    if cid not in user_cookies:
        raise HTTPException(status_code=403, detail="无权限访问该Cookie")

    # 检查文件类型
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="请上传Excel文件(.xlsx或.xls)")

    try:
        # 读取Excel文件
        contents = await file.read()
        df = pd.read_excel(io.BytesIO(contents))

        # 检查必要的列
        required_columns = ['关键词', '商品ID', '关键词内容']
        missing_columns = [col for col in required_columns if col not in df.columns]
        if missing_columns:
            raise HTTPException(status_code=400, detail=f"Excel文件缺少必要的列: {', '.join(missing_columns)}")

        # 获取现有的文本类型关键词（用于比较更新/新增）
        existing_keywords = db_manager.get_keywords_with_type(cid)
        existing_dict = {}
        for keyword_data in existing_keywords:
            # 只考虑文本类型的关键词
            if keyword_data.get('type', 'text') == 'text':
                keyword = keyword_data['keyword']
                reply = keyword_data['reply']
                item_id = keyword_data['item_id']
                key = f"{keyword}|{item_id or ''}"
                existing_dict[key] = (keyword, reply, item_id)

        # 处理导入数据
        import_data = []
        update_count = 0
        add_count = 0

        def clean_cell_value(value):
            """清理单元格值，处理数字转字符串时的 .0 后缀问题"""
            if pd.isna(value):
                return ''
            # 如果是数字类型，先转为整数（如果是整数值）再转字符串
            if isinstance(value, float) and value == int(value):
                return str(int(value)).strip()
            return str(value).strip()

        for index, row in df.iterrows():
            keyword = clean_cell_value(row['关键词'])
            item_id = clean_cell_value(row['商品ID']) or None
            reply = clean_cell_value(row['关键词内容'])

            if not keyword:
                continue  # 跳过没有关键词的行

            # 检查是否重复
            key = f"{keyword}|{item_id or ''}"
            if key in existing_dict:
                # 更新现有关键词
                update_count += 1
            else:
                # 新增关键词
                add_count += 1

            import_data.append((keyword, reply, item_id))

        if not import_data:
            raise HTTPException(status_code=400, detail="Excel文件中没有有效的关键词数据")

        # 保存到数据库（只影响文本关键词，保留图片关键词）
        success = db_manager.save_text_keywords_only(cid, import_data)
        if not success:
            raise HTTPException(status_code=500, detail="保存关键词到数据库失败")

        log_with_user('info', f"导入关键词成功: {cid}, 新增: {add_count}, 更新: {update_count}", current_user)

        return {
            "msg": "导入成功",
            "total": len(import_data),
            "added": add_count,
            "updated": update_count
        }

    except pd.errors.EmptyDataError:
        raise HTTPException(status_code=400, detail="Excel文件为空")
    except pd.errors.ParserError:
        raise HTTPException(status_code=400, detail="Excel文件格式错误")
    except Exception as e:
        logger.error(f"导入关键词失败: {e}")
        raise HTTPException(status_code=500, detail=f"导入关键词失败: {str(e)}")


@app.post("/keywords/{cid}/image")
async def add_image_keyword(
    cid: str,
    keyword: str = Form(...),
    item_id: str = Form(default=""),
    image: UploadFile = File(...),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """添加图片关键词"""
    logger.info(f"接收到图片关键词添加请求: cid={cid}, keyword={keyword}, item_id={item_id}")

    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")

    # 检查参数
    if not keyword or not keyword.strip():
        raise HTTPException(status_code=400, detail="关键词不能为空")

    if not image or not image.filename:
        raise HTTPException(status_code=400, detail="请选择图片文件")

    # 检查cookie是否属于当前用户
    cookie_details = db_manager.get_cookie_details(cid)
    if not cookie_details or cookie_details['user_id'] != current_user['user_id']:
        raise HTTPException(status_code=404, detail="账号不存在或无权限")

    try:
        logger.info(f"接收到图片关键词添加请求: cid={cid}, keyword={keyword}, item_id={item_id}, filename={image.filename}")

        # 验证图片文件
        if not image.content_type or not image.content_type.startswith('image/'):
            logger.warning(f"无效的图片文件类型: {image.content_type}")
            raise HTTPException(status_code=400, detail="请上传图片文件")

        # 读取图片数据
        image_data = await image.read()
        logger.info(f"读取图片数据成功，大小: {len(image_data)} bytes")

        # 保存图片
        image_url = image_manager.save_image(image_data, image.filename)
        if not image_url:
            logger.error("图片保存失败")
            raise HTTPException(status_code=400, detail="图片保存失败")

        logger.info(f"图片保存成功: {image_url}")

        # 先检查关键词是否已存在
        normalized_item_id = item_id if item_id and item_id.strip() else None
        if db_manager.check_keyword_duplicate(cid, keyword, normalized_item_id):
            # 删除已保存的图片
            image_manager.delete_image(image_url)
            if normalized_item_id:
                raise HTTPException(status_code=400, detail=f"关键词 '{keyword}' 在商品 '{normalized_item_id}' 中已存在")
            else:
                raise HTTPException(status_code=400, detail=f"通用关键词 '{keyword}' 已存在")

        # 保存图片关键词到数据库
        success = db_manager.save_image_keyword(cid, keyword, image_url, item_id or None)
        if not success:
            # 如果数据库保存失败，删除已保存的图片
            logger.error("数据库保存失败，删除已保存的图片")
            image_manager.delete_image(image_url)
            raise HTTPException(status_code=400, detail="图片关键词保存失败，请稍后重试")

        log_with_user('info', f"添加图片关键词成功: {cid}, 关键词: {keyword}", current_user)

        return {
            "msg": "图片关键词添加成功",
            "keyword": keyword,
            "image_url": image_url,
            "item_id": item_id or None
        }

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"添加图片关键词失败: {e}")
        raise HTTPException(status_code=500, detail=f"添加图片关键词失败: {str(e)}")


@app.post("/upload-image")
async def upload_image(
    image: UploadFile = File(...),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """上传图片（用于卡券等功能）"""
    try:
        logger.info(f"接收到图片上传请求: filename={image.filename}")

        # 验证图片文件
        if not image.content_type or not image.content_type.startswith('image/'):
            logger.warning(f"无效的图片文件类型: {image.content_type}")
            raise HTTPException(status_code=400, detail="请上传图片文件")

        # 读取图片数据
        image_data = await image.read()
        logger.info(f"读取图片数据成功，大小: {len(image_data)} bytes")

        # 保存图片
        image_url = image_manager.save_image(image_data, image.filename)
        if not image_url:
            logger.error("图片保存失败")
            raise HTTPException(status_code=400, detail="图片保存失败")

        logger.info(f"图片上传成功: {image_url}")

        return {
            "message": "图片上传成功",
            "image_url": image_url
        }

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"图片上传失败: {e}")
        raise HTTPException(status_code=500, detail=f"图片上传失败: {str(e)}")


@app.get("/keywords-with-type/{cid}")
def get_keywords_with_type(cid: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取包含类型信息的关键词列表"""
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")

    # 检查cookie是否属于当前用户
    cookie_details = db_manager.get_cookie_details(cid)
    if not cookie_details or cookie_details['user_id'] != current_user['user_id']:
        raise HTTPException(status_code=404, detail="账号不存在或无权限")

    try:
        keywords = db_manager.get_keywords_with_type(cid)
        return keywords
    except Exception as e:
        logger.error(f"获取关键词列表失败: {e}")
        raise HTTPException(status_code=500, detail=f"获取关键词列表失败: {str(e)}")


@app.delete("/keywords/{cid}/{index}")
def delete_keyword_by_index(cid: str, index: int, current_user: Dict[str, Any] = Depends(get_current_user)):
    """根据索引删除关键词"""
    if cookie_manager.manager is None:
        raise HTTPException(status_code=500, detail="CookieManager 未就绪")

    # 检查cookie是否属于当前用户
    cookie_details = db_manager.get_cookie_details(cid)
    if not cookie_details or cookie_details['user_id'] != current_user['user_id']:
        raise HTTPException(status_code=404, detail="账号不存在或无权限")

    try:
        # 先获取要删除的关键词信息（用于删除图片文件）
        keywords = db_manager.get_keywords_with_type(cid)
        if 0 <= index < len(keywords):
            keyword_data = keywords[index]

            # 删除关键词
            success = db_manager.delete_keyword_by_index(cid, index)
            if not success:
                raise HTTPException(status_code=400, detail="删除关键词失败")

            # 如果是图片关键词，删除对应的图片文件
            if keyword_data.get('type') == 'image' and keyword_data.get('image_url'):
                image_manager.delete_image(keyword_data['image_url'])

            log_with_user('info', f"删除关键词成功: {cid}, 索引: {index}, 关键词: {keyword_data.get('keyword')}", current_user)

            return {"msg": "删除成功"}
        else:
            raise HTTPException(status_code=400, detail="关键词索引无效")

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"删除关键词失败: {e}")
        raise HTTPException(status_code=500, detail=f"删除关键词失败: {str(e)}")


@app.get("/debug/keywords-table-info")
def debug_keywords_table_info(current_user: Dict[str, Any] = Depends(get_current_user)):
    """调试：检查keywords表结构"""
    try:
        import sqlite3
        conn = sqlite3.connect(db_manager.db_path)
        cursor = conn.cursor()

        # 获取表结构信息
        cursor.execute("PRAGMA table_info(keywords)")
        columns = cursor.fetchall()

        # 获取数据库版本
        cursor.execute("SELECT value FROM system_settings WHERE key = 'db_version'")
        version_result = cursor.fetchone()
        db_version = version_result[0] if version_result else "未知"

        conn.close()

        return {
            "db_version": db_version,
            "table_columns": [{"name": col[1], "type": col[2], "default": col[4]} for col in columns]
        }
    except Exception as e:
        logger.error(f"检查表结构失败: {e}")
        raise HTTPException(status_code=500, detail=f"检查表结构失败: {str(e)}")


# 卡券管理API
@app.get("/cards")
def get_cards(current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取当前用户的卡券列表"""
    try:
        from app.db_manager import db_manager
        user_id = current_user['user_id']
        cards = db_manager.get_all_cards(user_id)
        return cards
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/cards")
def create_card(card_data: dict, current_user: Dict[str, Any] = Depends(get_current_user)):
    """创建新卡券"""
    try:
        from app.db_manager import db_manager
        user_id = current_user['user_id']
        card_name = card_data.get('name', '未命名卡券')

        log_with_user('info', f"创建卡券: {card_name}", current_user)

        # 验证多规格字段
        is_multi_spec = card_data.get('is_multi_spec', False)
        if is_multi_spec:
            if not card_data.get('spec_name') or not card_data.get('spec_value'):
                raise HTTPException(status_code=400, detail="多规格卡券必须提供规格名称和规格值")

        card_id = db_manager.create_card(
            name=card_data.get('name'),
            card_type=card_data.get('type'),
            api_config=card_data.get('api_config'),
            text_content=card_data.get('text_content'),
            data_content=card_data.get('data_content'),
            image_url=card_data.get('image_url'),
            description=card_data.get('description'),
            enabled=card_data.get('enabled', True),
            delay_seconds=card_data.get('delay_seconds', 0),
            is_multi_spec=is_multi_spec,
            spec_name=card_data.get('spec_name') if is_multi_spec else None,
            spec_value=card_data.get('spec_value') if is_multi_spec else None,
            user_id=user_id
        )

        log_with_user('info', f"卡券创建成功: {card_name} (ID: {card_id})", current_user)
        return {"id": card_id, "message": "卡券创建成功"}
    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"创建卡券失败: {card_data.get('name', '未知')} - {str(e)}", current_user)
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/cards/{card_id}")
def get_card(card_id: int, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取单个卡券详情"""
    try:
        from app.db_manager import db_manager
        user_id = current_user['user_id']
        card = db_manager.get_card_by_id(card_id, user_id)
        if card:
            return card
        else:
            raise HTTPException(status_code=404, detail="卡券不存在")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.put("/cards/{card_id}")
def update_card(card_id: int, card_data: dict, current_user: Dict[str, Any] = Depends(get_current_user)):
    """更新卡券"""
    try:
        from app.db_manager import db_manager
        # 验证多规格字段
        is_multi_spec = card_data.get('is_multi_spec')
        if is_multi_spec:
            if not card_data.get('spec_name') or not card_data.get('spec_value'):
                raise HTTPException(status_code=400, detail="多规格卡券必须提供规格名称和规格值")

        success = db_manager.update_card(
            card_id=card_id,
            name=card_data.get('name'),
            card_type=card_data.get('type'),
            api_config=card_data.get('api_config'),
            text_content=card_data.get('text_content'),
            data_content=card_data.get('data_content'),
            image_url=card_data.get('image_url'),
            description=card_data.get('description'),
            enabled=card_data.get('enabled', True),
            delay_seconds=card_data.get('delay_seconds'),
            is_multi_spec=is_multi_spec,
            spec_name=card_data.get('spec_name'),
            spec_value=card_data.get('spec_value'),
            user_id=current_user['user_id']
        )
        if success:
            return {"message": "卡券更新成功"}
        else:
            raise HTTPException(status_code=404, detail="卡券不存在")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.put("/cards/{card_id}/image")
async def update_card_with_image(
    card_id: int,
    image: UploadFile = File(...),
    name: str = Form(...),
    type: str = Form(...),
    description: str = Form(default=""),
    delay_seconds: int = Form(default=0),
    enabled: bool = Form(default=True),
    is_multi_spec: bool = Form(default=False),
    spec_name: str = Form(default=""),
    spec_value: str = Form(default=""),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """更新带图片的卡券"""
    try:
        logger.info(f"接收到带图片的卡券更新请求: card_id={card_id}, name={name}, type={type}")

        # 验证图片文件
        if not image.content_type or not image.content_type.startswith('image/'):
            logger.warning(f"无效的图片文件类型: {image.content_type}")
            raise HTTPException(status_code=400, detail="请上传图片文件")

        # 验证多规格字段
        if is_multi_spec:
            if not spec_name or not spec_value:
                raise HTTPException(status_code=400, detail="多规格卡券必须提供规格名称和规格值")

        # 读取图片数据
        image_data = await image.read()
        logger.info(f"读取图片数据成功，大小: {len(image_data)} bytes")

        # 保存图片
        image_url = image_manager.save_image(image_data, image.filename)
        if not image_url:
            logger.error("图片保存失败")
            raise HTTPException(status_code=400, detail="图片保存失败")

        logger.info(f"图片保存成功: {image_url}")

        # 更新卡券
        from app.db_manager import db_manager
        success = db_manager.update_card(
            card_id=card_id,
            name=name,
            card_type=type,
            image_url=image_url,
            description=description,
            enabled=enabled,
            delay_seconds=delay_seconds,
            is_multi_spec=is_multi_spec,
            spec_name=spec_name if is_multi_spec else None,
            spec_value=spec_value if is_multi_spec else None,
            user_id=current_user['user_id']
        )

        if success:
            logger.info(f"卡券更新成功: {name} (ID: {card_id})")
            return {"message": "卡券更新成功", "image_url": image_url}
        else:
            # 如果数据库更新失败，删除已保存的图片
            image_manager.delete_image(image_url)
            raise HTTPException(status_code=404, detail="卡券不存在")

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"更新带图片的卡券失败: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# 自动发货规则API
def _validate_delivery_rule_scope(rule_data: dict, user_id: int):
    keyword = str(rule_data.get("keyword") or "").strip()
    cookie_id = str(rule_data.get("cookie_id") or "").strip() or None
    item_id = str(rule_data.get("item_id") or "").strip() or None
    card_id = rule_data.get("card_id")
    delivery_count_raw = rule_data.get("delivery_count", 1)

    if not card_id or not db_manager.get_card_by_id(int(card_id), user_id):
        raise HTTPException(status_code=400, detail="请选择当前用户可用的卡券")
    if isinstance(delivery_count_raw, bool):
        raise HTTPException(status_code=400, detail="每单发货数量必须是大于等于 1 的整数")
    try:
        delivery_count = int(delivery_count_raw)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="每单发货数量必须是大于等于 1 的整数")
    if str(delivery_count_raw).strip() != str(delivery_count) or delivery_count < 1:
        raise HTTPException(status_code=400, detail="每单发货数量必须是大于等于 1 的整数")
    if item_id and not cookie_id:
        raise HTTPException(status_code=400, detail="指定商品时必须同时选择账号")
    if cookie_id:
        owned_cookies = db_manager.get_all_cookies(user_id)
        if cookie_id not in owned_cookies:
            raise HTTPException(status_code=403, detail="无权使用该闲鱼账号")
    if item_id and not db_manager.get_item_info(cookie_id, item_id):
        raise HTTPException(status_code=400, detail="所选商品不存在，请先同步商品")
    if not item_id and not keyword:
        raise HTTPException(status_code=400, detail="通用规则必须填写触发关键词")

    return keyword, int(card_id), cookie_id, item_id, delivery_count


@app.get("/delivery-rules")
def get_delivery_rules(current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取发货规则列表"""
    try:
        from app.db_manager import db_manager
        user_id = current_user['user_id']
        rules = db_manager.get_all_delivery_rules(user_id)
        return rules
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/delivery-rules")
def create_delivery_rule(rule_data: dict, current_user: Dict[str, Any] = Depends(get_current_user)):
    """创建新发货规则"""
    try:
        from app.db_manager import db_manager
        user_id = current_user['user_id']
        keyword, card_id, cookie_id, item_id, delivery_count = _validate_delivery_rule_scope(
            rule_data, user_id
        )
        if item_id:
            duplicate = next(
                (
                    rule for rule in db_manager.get_all_delivery_rules(user_id)
                    if rule.get("cookie_id") == cookie_id
                    and rule.get("item_id") == item_id
                ),
                None,
            )
            if duplicate:
                raise HTTPException(
                    status_code=409,
                    detail="该商品已配置自动发货，请编辑现有策略",
                )
        rule_id = db_manager.create_delivery_rule(
            keyword=keyword,
            card_id=card_id,
            delivery_count=delivery_count,
            enabled=rule_data.get('enabled', True),
            description=rule_data.get('description'),
            user_id=user_id,
            cookie_id=cookie_id,
            item_id=item_id,
        )
        return {"id": rule_id, "message": "发货规则创建成功"}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/delivery-rules/{rule_id}")
def get_delivery_rule(rule_id: int, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取单个发货规则详情"""
    try:
        from app.db_manager import db_manager
        user_id = current_user['user_id']
        rule = db_manager.get_delivery_rule_by_id(rule_id, user_id)
        if rule:
            return rule
        else:
            raise HTTPException(status_code=404, detail="发货规则不存在")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.put("/delivery-rules/{rule_id}")
def update_delivery_rule(rule_id: int, rule_data: dict, current_user: Dict[str, Any] = Depends(get_current_user)):
    """更新发货规则"""
    try:
        from app.db_manager import db_manager
        user_id = current_user['user_id']
        keyword, card_id, cookie_id, item_id, delivery_count = _validate_delivery_rule_scope(
            rule_data, user_id
        )
        success = db_manager.update_delivery_rule(
            rule_id=rule_id,
            keyword=keyword,
            card_id=card_id,
            delivery_count=delivery_count,
            enabled=rule_data.get('enabled', True),
            description=rule_data.get('description'),
            user_id=user_id,
            cookie_id=cookie_id,
            item_id=item_id,
            scope_updated=True,
        )
        if success:
            return {"message": "发货规则更新成功"}
        else:
            raise HTTPException(status_code=404, detail="发货规则不存在")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.delete("/cards/{card_id}")
def delete_card(card_id: int, current_user: Dict[str, Any] = Depends(get_current_user)):
    """删除卡券"""
    try:
        from app.db_manager import db_manager
        references = db_manager.get_card_variant_references(
            card_id,
            current_user['user_id'],
        )
        if references:
            reference = references[0]
            item_name = reference.get("item_title") or reference["item_id"]
            raise HTTPException(
                status_code=409,
                detail=(
                    f"该卡密正被商品“{item_name}”的规格"
                    f"“{reference['variant_name']}”使用，请先解除绑定"
                ),
            )
        success = db_manager.delete_card(card_id, current_user['user_id'])
        if success:
            return {"message": "卡券删除成功"}
        else:
            raise HTTPException(status_code=404, detail="卡券不存在")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.delete("/delivery-rules/{rule_id}")
def delete_delivery_rule(rule_id: int, current_user: Dict[str, Any] = Depends(get_current_user)):
    """删除发货规则"""
    try:
        from app.db_manager import db_manager
        user_id = current_user['user_id']
        success = db_manager.delete_delivery_rule(rule_id, user_id)
        if success:
            return {"message": "发货规则删除成功"}
        else:
            raise HTTPException(status_code=404, detail="发货规则不存在")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ==================== 商品自动化 API ====================

def _raise_product_automation_error(error: Exception):
    if isinstance(error, PermissionError):
        raise HTTPException(status_code=403, detail=str(error))
    if isinstance(error, LookupError):
        raise HTTPException(status_code=404, detail=str(error))
    if isinstance(error, ValueError):
        raise HTTPException(status_code=400, detail=str(error))
    logger.exception("商品自动化操作失败")
    raise HTTPException(status_code=500, detail="商品自动化操作失败")


@app.get("/product-automation/materials")
def get_product_materials(
    cookie_id: Optional[str] = Query(default=None),
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    try:
        return {
            "success": True,
            "data": product_automation.list_materials(current_user["user_id"], cookie_id),
        }
    except Exception as error:
        _raise_product_automation_error(error)


@app.put("/product-automation/materials/{material_id}")
def update_product_material(
    material_id: int,
    payload: Dict[str, Any],
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    try:
        return {
            "success": True,
            "data": product_automation.update_material(
                current_user["user_id"], material_id, payload
            ),
        }
    except Exception as error:
        _raise_product_automation_error(error)


@app.delete("/product-automation/materials/{material_id}")
def delete_product_material(
    material_id: int,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    try:
        if not product_automation.delete_material(current_user["user_id"], material_id):
            raise LookupError("素材不存在")
        return {"success": True, "message": "素材已删除"}
    except Exception as error:
        _raise_product_automation_error(error)


@app.get("/product-automation/filter-rules")
def get_product_filter_rules(
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    return {
        "success": True,
        "data": product_automation.list_filter_rules(current_user["user_id"]),
    }


@app.post("/product-automation/filter-rules")
def create_product_filter_rule(
    payload: Dict[str, Any],
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    try:
        return {
            "success": True,
            "data": product_automation.save_filter_rule(
                current_user["user_id"], payload
            ),
        }
    except Exception as error:
        _raise_product_automation_error(error)


@app.put("/product-automation/filter-rules/{rule_id}")
def update_product_filter_rule(
    rule_id: int,
    payload: Dict[str, Any],
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    try:
        return {
            "success": True,
            "data": product_automation.save_filter_rule(
                current_user["user_id"], payload, rule_id
            ),
        }
    except Exception as error:
        _raise_product_automation_error(error)


@app.delete("/product-automation/filter-rules/{rule_id}")
def delete_product_filter_rule(
    rule_id: int,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    try:
        if not product_automation.delete_filter_rule(current_user["user_id"], rule_id):
            raise LookupError("筛选规则不存在")
        return {"success": True, "message": "筛选规则已删除"}
    except Exception as error:
        _raise_product_automation_error(error)


@app.post("/product-automation/filter-rules/{rule_id}/run")
def run_product_filter_rule(
    rule_id: int,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    try:
        return {
            "success": True,
            "data": product_automation.run_filter_rule(
                current_user["user_id"], rule_id
            ),
        }
    except Exception as error:
        _raise_product_automation_error(error)


@app.get("/product-automation/delete-rules")
def get_product_delete_rules(
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    return {
        "success": True,
        "data": product_automation.list_delete_rules(current_user["user_id"]),
    }


@app.post("/product-automation/delete-rules")
def create_product_delete_rule(
    payload: Dict[str, Any],
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    try:
        return {
            "success": True,
            "data": product_automation.save_delete_rule(
                current_user["user_id"], payload
            ),
        }
    except Exception as error:
        _raise_product_automation_error(error)


@app.put("/product-automation/delete-rules/{rule_id}")
def update_product_delete_rule(
    rule_id: int,
    payload: Dict[str, Any],
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    try:
        return {
            "success": True,
            "data": product_automation.save_delete_rule(
                current_user["user_id"], payload, rule_id
            ),
        }
    except Exception as error:
        _raise_product_automation_error(error)


@app.delete("/product-automation/delete-rules/{rule_id}")
def delete_product_delete_rule(
    rule_id: int,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    try:
        if not product_automation.delete_delete_rule(current_user["user_id"], rule_id):
            raise LookupError("删除计划不存在")
        return {"success": True, "message": "删除计划已删除"}
    except Exception as error:
        _raise_product_automation_error(error)


@app.post("/product-automation/delete-rules/{rule_id}/preview")
def preview_product_delete_rule(
    rule_id: int,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    try:
        return {
            "success": True,
            "data": product_automation.preview_delete_rule(
                current_user["user_id"], rule_id
            ),
        }
    except Exception as error:
        _raise_product_automation_error(error)


@app.get("/product-automation/runs")
def get_product_automation_runs(
    limit: int = Query(default=50, ge=1, le=200),
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    return {
        "success": True,
        "data": product_automation.list_runs(current_user["user_id"], limit),
    }


@app.post("/product-automation/repairs/published-ids")
def repair_product_published_ids(
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    return {
        "success": True,
        "data": product_automation.repair_published_ids(current_user["user_id"]),
    }


@app.post("/product-automation/repairs/short-links")
def repair_product_short_links(
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    return {
        "success": True,
        "data": product_automation.repair_short_links(current_user["user_id"]),
    }


@app.post("/product-automation/repairs/cards")
def compensate_product_cards(
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    return {
        "success": True,
        "data": product_automation.compensate_cards(current_user["user_id"]),
    }


# ==================== 备份和恢复 API ====================

@app.get("/backup/export")
def export_backup(current_user: Dict[str, Any] = Depends(get_current_user)):
    """导出用户备份"""
    try:
        from app.db_manager import db_manager
        user_id = current_user['user_id']
        username = current_user['username']

        # 导出当前用户的数据
        backup_data = db_manager.export_backup(user_id)

        # 生成文件名
        import datetime
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"xianyu_backup_{username}_{timestamp}.json"

        # 返回JSON响应，设置下载头
        response = JSONResponse(content=backup_data)
        response.headers["Content-Disposition"] = f"attachment; filename={filename}"
        response.headers["Content-Type"] = "application/json"

        return response
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"导出备份失败: {str(e)}")


@app.post("/backup/import")
def import_backup(file: UploadFile = File(...), current_user: Dict[str, Any] = Depends(get_current_user)):
    """导入用户备份"""
    try:
        # 验证文件类型
        if not file.filename.endswith('.json'):
            raise HTTPException(status_code=400, detail="只支持JSON格式的备份文件")

        # 读取文件内容
        content = file.file.read()
        backup_data = json.loads(content.decode('utf-8'))

        # 导入备份到当前用户
        from app.db_manager import db_manager
        user_id = current_user['user_id']
        success = db_manager.import_backup(backup_data, user_id)

        if success:
            # 备份导入成功后，刷新 CookieManager 的内存缓存
            from app import cookie_manager
            if cookie_manager.manager:
                try:
                    cookie_manager.manager.reload_from_db()
                    logger.info("备份导入后已刷新 CookieManager 缓存")
                except Exception as e:
                    logger.error(f"刷新 CookieManager 缓存失败: {e}")

            return {"message": "备份导入成功"}
        else:
            raise HTTPException(status_code=400, detail="备份导入失败")

    except json.JSONDecodeError:
        raise HTTPException(status_code=400, detail="备份文件格式无效")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"导入备份失败: {str(e)}")


@app.post("/system/reload-cache")
def reload_cache(_: Dict[str, Any] = Depends(require_admin)):
    """重新加载系统缓存（仅管理员）"""
    try:
        from app import cookie_manager
        if cookie_manager.manager:
            success = cookie_manager.manager.reload_from_db()
            if success:
                return {"message": "系统缓存已刷新", "success": True}
            else:
                raise HTTPException(status_code=500, detail="缓存刷新失败")
        else:
            raise HTTPException(status_code=500, detail="CookieManager 未初始化")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"刷新缓存失败: {str(e)}")


# ==================== 商品管理 API ====================

@app.get("/items")
def get_all_items(current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取当前用户的所有商品信息"""
    try:
        # 只返回当前用户的商品信息
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        all_items = []
        for cookie_id in user_cookies.keys():
            items = db_manager.get_items_by_cookie(cookie_id)
            all_items.extend(items)

        return {"items": all_items}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"获取商品信息失败: {str(e)}")


class ManualItemCreate(BaseModel):
    cookie_id: str
    item_id: str
    title: str
    price: Optional[str] = ""
    image_url: Optional[str] = ""
    description: Optional[str] = ""
    detail: Optional[str] = ""


@app.post("/items")
def create_manual_item(
    item: ManualItemCreate,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    """手动添加未从闲鱼同步到本地的商品。"""
    try:
        from app.db_manager import db_manager

        cookie_id = item.cookie_id.strip()
        item_id = item.item_id.strip()
        title = item.title.strip()
        if not cookie_id:
            raise HTTPException(status_code=400, detail="请选择商品所属账号")
        if not item_id:
            raise HTTPException(status_code=400, detail="请填写商品 ID")
        if not title:
            raise HTTPException(status_code=400, detail="请填写商品标题")

        user_cookies = db_manager.get_all_cookies(current_user["user_id"])
        if cookie_id not in user_cookies:
            raise HTTPException(status_code=403, detail="无权使用该闲鱼账号")
        if db_manager.get_item_info(cookie_id, item_id):
            raise HTTPException(status_code=409, detail="该商品已存在，请直接编辑现有商品")

        saved = db_manager.save_item_info(
            cookie_id,
            item_id,
            {
                "title": title,
                "description": item.description.strip(),
                "category": "",
                "price": item.price.strip(),
                "item_image": item.image_url.strip(),
                "item_detail": item.detail.strip(),
                "source": "manual",
            },
        )
        if not saved:
            raise HTTPException(status_code=500, detail="手动商品保存失败")

        return {
            "success": True,
            "message": "商品添加成功",
            "item": db_manager.get_item_info(cookie_id, item_id),
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("手动添加商品失败")
        raise HTTPException(status_code=500, detail=f"手动添加商品失败: {str(e)}")


class ProductVariantBindingIn(BaseModel):
    display_name: str = ""
    spec_text: str = ""
    spec_payload: Optional[Dict[str, Any]] = None
    platform_sku_id: str = ""
    card_id: int
    delivery_count: int = Field(default=1, ge=1)
    enabled: bool = True
    binding_enabled: bool = True
    source: str = "manual"


class ItemDeliveryConfigIn(BaseModel):
    enabled: bool = True
    is_multi_spec: bool = False
    variants: List[ProductVariantBindingIn]


@app.get("/item-delivery-configs")
def get_item_delivery_config_summaries(
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    """获取商品变体发货配置摘要，供商品列表一次性展示。"""
    return {
        "configs": db_manager.get_item_delivery_config_summaries(current_user["user_id"])
    }


@app.get("/items/{cookie_id}/{item_id}/delivery-config")
def get_item_delivery_config(
    cookie_id: str,
    item_id: str,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    """获取一个商品的完整变体发货配置。"""
    if cookie_id not in db_manager.get_all_cookies(current_user["user_id"]):
        raise HTTPException(status_code=403, detail="无权访问该闲鱼账号")
    if not db_manager.get_item_info(cookie_id, item_id):
        raise HTTPException(status_code=404, detail="商品不存在，请先同步或手动添加")

    config = db_manager.get_item_delivery_config(
        cookie_id,
        item_id,
        current_user["user_id"],
    )
    if config:
        return {"configured": True, **config}
    return {
        "configured": False,
        "cookie_id": cookie_id,
        "item_id": item_id,
        "enabled": False,
        "is_multi_spec": False,
        "variant_count": 0,
        "configured_count": 0,
        "complete": False,
        "variants": [],
    }


@app.put("/items/{cookie_id}/{item_id}/delivery-config")
def save_item_delivery_config(
    cookie_id: str,
    item_id: str,
    config_data: ItemDeliveryConfigIn,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    """整体保存商品规格与逐规格发货库存绑定。"""
    user_id = current_user["user_id"]
    if cookie_id not in db_manager.get_all_cookies(user_id):
        raise HTTPException(status_code=403, detail="无权操作该闲鱼账号")
    try:
        config = db_manager.save_item_delivery_config(
            user_id=user_id,
            cookie_id=cookie_id,
            item_id=item_id,
            enabled=config_data.enabled,
            is_multi_spec=config_data.is_multi_spec,
            variants=[
                variant.model_dump()
                if hasattr(variant, "model_dump")
                else variant.dict()
                for variant in config_data.variants
            ],
        )
        return {
            "success": True,
            "message": "商品自动发货配置已保存",
            "config": config,
        }
    except ValueError as e:
        # 记下具体原因。只看访问日志里的 400 无法判断是哪一条校验没过，
        # 用户那边也只会看到「Request failed with status code 400」。
        logger.warning(f"保存商品发货配置被拒绝: {cookie_id}/{item_id} - {e}")
        raise HTTPException(status_code=400, detail=str(e))
    except sqlite3.IntegrityError as e:
        logger.warning(f"保存商品变体配置冲突: {cookie_id}/{item_id} - {e}")
        raise HTTPException(status_code=409, detail="规格组合或平台 SKU ID 重复")
    except HTTPException:
        raise
    except Exception as e:
        logger.exception(f"保存商品变体配置失败: {cookie_id}/{item_id}")
        raise HTTPException(status_code=500, detail=f"保存商品发货配置失败: {str(e)}")


# ==================== 商品搜索 API ====================

class ItemSearchRequest(BaseModel):
    keyword: str
    page: int = 1
    page_size: int = 20

class ItemSearchMultipleRequest(BaseModel):
    keyword: str
    total_pages: int = 1

@app.post("/items/search")
async def search_items(
    search_request: ItemSearchRequest,
    current_user: Optional[Dict[str, Any]] = Depends(get_current_user_optional)
):
    """搜索闲鱼商品"""
    user_info = f"【{current_user.get('username', 'unknown')}#{current_user.get('user_id', 'unknown')}】" if current_user else "【未登录】"

    try:
        logger.info(f"{user_info} 开始单页搜索: 关键词='{search_request.keyword}', 页码={search_request.page}, 每页={search_request.page_size}")

        from utils.item_search import search_xianyu_items

        # 执行搜索
        result = await search_xianyu_items(
            keyword=search_request.keyword,
            page=search_request.page,
            page_size=search_request.page_size
        )

        # 检查是否有错误
        has_error = result.get("error")
        items_count = len(result.get("items", []))

        logger.info(f"{user_info} 单页搜索完成: 获取到 {items_count} 条数据" +
                   (f", 错误: {has_error}" if has_error else ""))

        response_data = {
            "success": True,
            "data": result.get("items", []),
            "total": result.get("total", 0),
            "page": search_request.page,
            "page_size": search_request.page_size,
            "keyword": search_request.keyword,
            "is_real_data": result.get("is_real_data", False),
            "source": result.get("source", "unknown")
        }

        # 如果有错误信息，也包含在响应中
        if has_error:
            response_data["error"] = has_error

        return response_data

    except Exception as e:
        error_msg = str(e)
        logger.error(f"{user_info} 商品搜索失败: {error_msg}")
        raise HTTPException(status_code=500, detail=f"商品搜索失败: {error_msg}")


@app.get("/cookies/check")
async def check_valid_cookies(
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """检查是否有有效的cookies账户（必须是启用状态）"""
    try:
        if cookie_manager.manager is None:
            return {
                "success": True,
                "hasValidCookies": False,
                "validCount": 0,
                "enabledCount": 0,
                "totalCount": 0
            }

        from app.db_manager import db_manager

        all_cookies = db_manager.get_all_cookies(current_user['user_id'])

        # 检查启用状态和有效性
        valid_cookies = []
        enabled_cookies = []

        for cookie_id, cookie_value in all_cookies.items():
            # 检查是否启用
            is_enabled = cookie_manager.manager.get_cookie_status(cookie_id)
            if is_enabled:
                enabled_cookies.append(cookie_id)
                # 检查是否有效（长度大于50）
                if len(cookie_value) > 50:
                    valid_cookies.append(cookie_id)

        return {
            "success": True,
            "hasValidCookies": len(valid_cookies) > 0,
            "validCount": len(valid_cookies),
            "enabledCount": len(enabled_cookies),
            "totalCount": len(all_cookies)
        }

    except Exception as e:
        logger.error(f"检查cookies失败: {str(e)}")
        return {
            "success": False,
            "hasValidCookies": False,
            "error": str(e)
        }

@app.post("/items/search_multiple")
async def search_multiple_pages(
    search_request: ItemSearchMultipleRequest,
    current_user: Optional[Dict[str, Any]] = Depends(get_current_user_optional)
):
    """搜索多页闲鱼商品"""
    user_info = f"【{current_user.get('username', 'unknown')}#{current_user.get('user_id', 'unknown')}】" if current_user else "【未登录】"

    try:
        logger.info(f"{user_info} 开始多页搜索: 关键词='{search_request.keyword}', 页数={search_request.total_pages}")

        from utils.item_search import search_multiple_pages_xianyu

        # 执行多页搜索
        result = await search_multiple_pages_xianyu(
            keyword=search_request.keyword,
            total_pages=search_request.total_pages
        )

        # 检查是否有错误
        has_error = result.get("error")
        items_count = len(result.get("items", []))

        logger.info(f"{user_info} 多页搜索完成: 获取到 {items_count} 条数据" +
                   (f", 错误: {has_error}" if has_error else ""))

        response_data = {
            "success": True,
            "data": result.get("items", []),
            "total": result.get("total", 0),
            "total_pages": search_request.total_pages,
            "keyword": search_request.keyword,
            "is_real_data": result.get("is_real_data", False),
            "is_fallback": result.get("is_fallback", False),
            "source": result.get("source", "unknown")
        }

        # 如果有错误信息，也包含在响应中
        if has_error:
            response_data["error"] = has_error

        return response_data

    except Exception as e:
        error_msg = str(e)
        logger.error(f"{user_info} 多页商品搜索失败: {error_msg}")
        raise HTTPException(status_code=500, detail=f"多页商品搜索失败: {error_msg}")



@app.get("/items/cookie/{cookie_id}")
def get_items_by_cookie(cookie_id: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取指定Cookie的商品信息"""
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cookie_id not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限访问该Cookie")

        items = db_manager.get_items_by_cookie(cookie_id)
        return {"items": items}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"获取商品信息失败: {str(e)}")


@app.get("/items/{cookie_id}/{item_id}")
def get_item_detail(cookie_id: str, item_id: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取商品详情"""
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cookie_id not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限访问该Cookie")

        item = db_manager.get_item_info(cookie_id, item_id)
        if not item:
            raise HTTPException(status_code=404, detail="商品不存在")
        return {"item": item}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"获取商品详情失败: {str(e)}")


class ItemDetailUpdate(BaseModel):
    item_detail: str


@app.put("/items/{cookie_id}/{item_id}")
def update_item_detail(
    cookie_id: str,
    item_id: str,
    update_data: ItemDetailUpdate,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """更新商品详情"""
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cookie_id not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        success = db_manager.update_item_detail(cookie_id, item_id, update_data.item_detail)
        if success:
            return {"message": "商品详情更新成功"}
        else:
            raise HTTPException(status_code=400, detail="更新失败")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"更新商品详情失败: {str(e)}")


@app.delete("/items/{cookie_id}/{item_id}")
def delete_item_info(
    cookie_id: str,
    item_id: str,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """删除商品信息"""
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cookie_id not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        success = db_manager.delete_item_info(cookie_id, item_id)
        if success:
            return {"message": "商品信息删除成功"}
        else:
            raise HTTPException(status_code=404, detail="商品信息不存在")
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"删除商品信息异常: {e}")
        raise HTTPException(status_code=500, detail=f"服务器错误: {str(e)}")


class BatchDeleteRequest(BaseModel):
    items: List[dict]  # [{"cookie_id": "xxx", "item_id": "yyy"}, ...]


class AIReplySettings(BaseModel):
    ai_enabled: bool
    model_name: str = "qwen-plus"
    api_key: str = ""
    base_url: str = "https://ai.corleom.com/v1"
    max_discount_percent: int = 10
    max_discount_amount: int = 100
    max_bargain_rounds: int = 3
    context_enabled: bool = True
    context_message_limit: int = 12
    context_expire_minutes: int = 120
    custom_prompts: str = ""


def _public_ai_reply_settings(settings: dict) -> dict:
    """返回前端可展示的AI配置，不暴露密钥。"""
    public_settings = dict(settings)
    public_settings['api_key_configured'] = bool(public_settings.get('api_key'))
    public_settings['api_key'] = ''
    return public_settings


@app.delete("/items/batch")
def batch_delete_items(
    request: BatchDeleteRequest,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """批量删除商品信息"""
    try:
        if not request.items:
            raise HTTPException(status_code=400, detail="删除列表不能为空")

        user_cookies = db_manager.get_all_cookies(current_user['user_id'])
        unauthorized = [item for item in request.items if item.get('cookie_id') not in user_cookies]
        if unauthorized:
            raise HTTPException(status_code=403, detail="删除列表中包含无权操作的账号")

        success_count = db_manager.batch_delete_item_info(request.items)
        total_count = len(request.items)

        return {
            "message": f"批量删除完成",
            "success_count": success_count,
            "total_count": total_count,
            "failed_count": total_count - success_count
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"批量删除商品信息异常: {e}")
        raise HTTPException(status_code=500, detail=f"服务器错误: {str(e)}")


# ==================== AI回复管理API ====================

@app.get("/ai-reply-settings/{cookie_id}")
def get_ai_reply_settings(cookie_id: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取指定账号的AI回复设置"""
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cookie_id not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限访问该Cookie")

        settings = db_manager.get_ai_reply_settings(cookie_id)
        return _public_ai_reply_settings(settings)
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"获取AI回复设置异常: {e}")
        raise HTTPException(status_code=500, detail=f"服务器错误: {str(e)}")


@app.put("/ai-reply-settings/{cookie_id}")
def update_ai_reply_settings(cookie_id: str, settings: AIReplySettings, current_user: Dict[str, Any] = Depends(get_current_user)):
    """更新指定账号的AI回复设置"""
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cookie_id not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该Cookie")

        # 检查账号是否存在
        if cookie_manager.manager is None:
            raise HTTPException(status_code=500, detail='CookieManager 未就绪')

        # 保存设置
        settings_dict = settings.dict()
        settings_dict['context_message_limit'] = max(
            2, min(30, settings.context_message_limit)
        )
        settings_dict['context_expire_minutes'] = max(
            5, min(1440, settings.context_expire_minutes)
        )
        if not settings_dict.get('api_key'):
            settings_dict['api_key'] = db_manager.get_account_ai_api_key(cookie_id)
        success = db_manager.save_ai_reply_settings(cookie_id, settings_dict)

        if success:

            # 如果启用了AI回复，记录日志
            if settings.ai_enabled:
                logger.info(f"账号 {cookie_id} 启用AI回复")
            else:
                logger.info(f"账号 {cookie_id} 禁用AI回复")

            return {"message": "AI回复设置更新成功"}
        else:
            raise HTTPException(status_code=400, detail="更新失败")
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"更新AI回复设置异常: {e}")
        raise HTTPException(status_code=500, detail=f"服务器错误: {str(e)}")


@app.get("/ai-reply-settings")
def get_all_ai_reply_settings(current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取当前用户所有账号的AI回复设置"""
    try:
        # 只返回当前用户的AI回复设置
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        user_settings = {
            cid: _public_ai_reply_settings(db_manager.get_ai_reply_settings(cid))
            for cid in user_cookies
        }
        return user_settings
    except Exception as e:
        logger.error(f"获取所有AI回复设置异常: {e}")
        raise HTTPException(status_code=500, detail=f"服务器错误: {str(e)}")


@app.post("/ai-reply-test/{cookie_id}")
async def test_ai_reply(cookie_id: str, test_data: dict,
                        current_user: Dict[str, Any] = Depends(get_current_user)):
    """测试AI回复功能"""
    try:
        # 检查账号是否存在
        if cookie_manager.manager is None:
            raise HTTPException(status_code=500, detail='CookieManager 未就绪')

        user_cookies = db_manager.get_all_cookies(current_user['user_id'])
        if cookie_id not in user_cookies:
            raise HTTPException(status_code=403, detail='无权限操作该账号')

        if cookie_id not in cookie_manager.manager.cookies:
            raise HTTPException(status_code=404, detail='账号不存在')

        # 检查是否启用AI回复
        if not ai_reply_engine.is_ai_enabled(cookie_id):
            raise HTTPException(status_code=400, detail='该账号未启用 AI 回复，请先在 AI 回复页开启并点击保存配置')

        # 检查AI设置是否完整
        settings = db_manager.get_ai_reply_settings(cookie_id)
        if not settings.get('api_key'):
            raise HTTPException(status_code=400, detail='未配置API Key，请先在AI设置中配置API Key')
        if not settings.get('base_url'):
            raise HTTPException(status_code=400, detail='未配置API地址，请先在AI设置中配置API地址')

        # 构造测试数据
        test_message = test_data.get('message', '你好')
        test_item_info = {
            'title': test_data.get('item_title', '测试商品'),
            'price': test_data.get('item_price', 100),
            'desc': test_data.get('item_desc', '这是一个测试商品')
        }

        # 生成测试回复（跳过等待时间）
        reply = await ai_reply_engine.generate_reply_async(
            message=test_message,
            item_info=test_item_info,
            chat_id=f"test_{int(time.time())}",
            cookie_id=cookie_id,
            user_id="test_user",
            item_id="test_item",
            skip_wait=True  # 测试时跳过10秒等待
        )

        if reply:
            return {"message": "测试成功", "reply": reply}
        else:
            raise HTTPException(status_code=400, detail="AI回复生成失败，请检查API Key是否正确、API地址是否可访问")

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"测试AI回复异常: {e}")
        import traceback
        logger.error(f"详细错误: {traceback.format_exc()}")
        raise HTTPException(status_code=500, detail=f"服务器错误: {str(e)}")


# ==================== 日志管理API ====================

@app.get("/logs")
async def get_logs(lines: int = 200, level: str = None, source: str = None,
                   _: Dict[str, Any] = Depends(require_admin)):
    """获取实时系统日志"""
    try:
        # 获取文件日志收集器
        collector = get_file_log_collector()

        # 获取日志
        logs = collector.get_logs(lines=lines, level_filter=level, source_filter=source)

        return {"success": True, "logs": logs}

    except Exception as e:
        return {"success": False, "message": f"获取日志失败: {str(e)}", "logs": []}


@app.get("/risk-control-logs")
async def get_risk_control_logs(
    cookie_id: str = None,
    processing_status: str = None,
    limit: int = 100,
    offset: int = 0,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """获取当前用户账号的风控日志。"""
    try:
        limit = max(1, min(limit, 200))
        offset = max(0, offset)
        if processing_status not in (None, "", "processing", "success", "failed"):
            raise HTTPException(status_code=400, detail="无效的风控日志处理状态")
        processing_status = processing_status or None
        user_id = current_user["user_id"]
        log_with_user(
            'info',
            f"查询风控日志: cookie_id={cookie_id}, processing_status={processing_status}, "
            f"limit={limit}, offset={offset}",
            current_user,
        )

        logs = db_manager.get_risk_control_logs(
            cookie_id=cookie_id,
            limit=limit,
            offset=offset,
            user_id=user_id,
            processing_status=processing_status,
        )
        total_count = db_manager.get_risk_control_logs_count(
            cookie_id=cookie_id,
            user_id=user_id,
            processing_status=processing_status,
        )

        log_with_user('info', f"风控日志查询成功，共 {len(logs)} 条记录，总计 {total_count} 条", current_user)

        return {
            "success": True,
            "data": logs,
            "total": total_count,
            "limit": limit,
            "offset": offset
        }

    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"获取风控日志失败: {str(e)}", current_user)
        return {
            "success": False,
            "message": f"获取风控日志失败: {str(e)}",
            "data": [],
            "total": 0
        }


@app.delete("/risk-control-logs/{log_id}")
async def delete_risk_control_log(
    log_id: int,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """删除风控日志记录，普通用户只能删除自己账号的记录。"""
    try:
        log_with_user('info', f"删除风控日志记录: {log_id}", current_user)

        is_admin = current_user.get("is_admin", False) or current_user.get("username") == ADMIN_USERNAME
        success = db_manager.delete_risk_control_log(
            log_id,
            user_id=None if is_admin else current_user["user_id"],
        )

        if success:
            log_with_user('info', f"风控日志删除成功: {log_id}", current_user)
            return {"success": True, "message": "删除成功"}
        else:
            log_with_user('warning', f"风控日志删除失败或无权限: {log_id}", current_user)
            return {"success": False, "message": "删除失败，记录可能不存在或无权访问"}

    except Exception as e:
        log_with_user('error', f"删除风控日志失败: {log_id} - {str(e)}", current_user)
        return {"success": False, "message": f"删除失败: {str(e)}"}


@app.get("/logs/stats")
async def get_log_stats(_: Dict[str, Any] = Depends(require_admin)):
    """获取日志统计信息"""
    try:
        collector = get_file_log_collector()
        stats = collector.get_stats()

        return {"success": True, "stats": stats}

    except Exception as e:
        return {"success": False, "message": f"获取日志统计失败: {str(e)}", "stats": {}}


@app.post("/logs/clear")
async def clear_logs(_: Dict[str, Any] = Depends(require_admin)):
    """清空日志"""
    try:
        collector = get_file_log_collector()
        collector.clear_logs()

        return {"success": True, "message": "日志已清空"}

    except Exception as e:
        return {"success": False, "message": f"清空日志失败: {str(e)}"}


# ==================== 商品管理API ====================

@app.post("/items/get-all-from-account")
async def get_all_items_from_account(request: dict, current_user: Dict[str, Any] = Depends(get_current_user)):
    """从指定账号获取所有商品信息"""
    try:
        cookie_id = request.get('cookie_id')
        if not cookie_id:
            raise HTTPException(status_code=400, detail="缺少cookie_id参数")

        user_cookies = db_manager.get_all_cookies(current_user['user_id'])
        if cookie_id not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该账号")

        # 获取指定账号的cookie信息
        cookie_info = db_manager.get_cookie_by_id(cookie_id)
        if not cookie_info:
            raise HTTPException(status_code=404, detail="未找到指定的账号信息")

        cookies_str = cookie_info.get('cookies_str', '')
        if not cookies_str:
            raise HTTPException(status_code=400, detail="账号cookie信息为空")

        # 创建XianyuLive实例，传入正确的账号和用户归属
        from XianyuAutoAsync import XianyuLive
        xianyu_instance = XianyuLive(
            cookies_str=cookies_str,
            cookie_id=cookie_id,
            user_id=current_user['user_id']
        )

        try:
            logger.info(f"开始获取账号 {cookie_id} 的所有商品信息")
            result = await xianyu_instance.get_all_items()
        finally:
            await xianyu_instance.close_session()

        if not result.get('success'):
            error_message = result.get('error', '闲鱼商品同步失败，未返回明确原因')
            logger.error(f"获取商品信息失败: {error_message}")
            raise HTTPException(status_code=502, detail=error_message)

        total_count = result.get('total_count', 0)
        total_pages = result.get('total_pages', 1)
        saved_count = result.get('total_saved', 0)
        account_id = result.get('account_id') or cookie_id
        group_name = result.get('group_name', '在售')
        confirmed_empty = bool(result.get('confirmed_empty'))
        api_total_count = result.get('api_total_count', total_count)
        group_declared_count = result.get('group_declared_count', total_count)
        parsed_count = result.get('parsed_count', total_count)
        count_reconciled = bool(result.get('count_reconciled'))
        off_shelf_count = int(result.get('off_shelf_count') or 0)
        # 本次未被接口返回的商品已标为已下架，要告诉用户 —— 否则列表里少了几件
        # 会显得像数据丢了。
        off_shelf_note = (
            f"；另有 {off_shelf_count} 件本次未返回，已标记为已下架"
            if off_shelf_count else ""
        )

        if confirmed_empty:
            message = f"闲鱼接口确认账号 {account_id} 的“{group_name}”分组当前为 0 件"
            logger.info(message)
        elif count_reconciled:
            message = (
                f"成功同步 {parsed_count} 件商品；闲鱼接口计数异常"
                f"（接口 {api_total_count}、分组 {group_declared_count}），已按商品列表自动校准"
                f"{off_shelf_note}"
            )
            logger.warning(message)
        else:
            message = f"成功获取商品，共 {total_count} 件，保存 {saved_count} 件{off_shelf_note}"
            logger.info(
                f"成功获取账号 {cookie_id} 的 {total_count} 个商品"
                f"（共{total_pages}页），保存 {saved_count} 个，"
                f"标记已下架 {off_shelf_count} 个"
            )

        return {
            "success": True,
            "message": message,
            "total_count": total_count,
            "total_pages": total_pages,
            "saved_count": saved_count,
            "confirmed_empty": confirmed_empty,
            "api_total_count": api_total_count,
            "group_declared_count": group_declared_count,
            "parsed_count": parsed_count,
            "count_reconciled": count_reconciled,
            "group_name": group_name,
            "account_id": account_id
        }

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"获取账号商品信息异常: {str(e)}")
        raise HTTPException(status_code=500, detail=f"获取商品信息异常: {str(e)}")


@app.post("/items/get-by-page")
async def get_items_by_page(request: dict, current_user: Dict[str, Any] = Depends(get_current_user)):
    """从指定账号按页获取商品信息"""
    try:
        # 验证参数
        cookie_id = request.get('cookie_id')
        page_number = request.get('page_number', 1)
        page_size = request.get('page_size', 20)

        if not cookie_id:
            raise HTTPException(status_code=400, detail="缺少cookie_id参数")

        user_cookies = db_manager.get_all_cookies(current_user['user_id'])
        if cookie_id not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限操作该账号")

        # 验证分页参数
        try:
            page_number = int(page_number)
            page_size = int(page_size)
        except (ValueError, TypeError):
            raise HTTPException(status_code=400, detail="页码和每页数量必须是数字")

        if page_number < 1:
            raise HTTPException(status_code=400, detail="页码必须大于0")

        if page_size < 1 or page_size > 100:
            raise HTTPException(status_code=400, detail="每页数量必须在1-100之间")

        # 获取账号信息
        account = db_manager.get_cookie_by_id(cookie_id)
        if not account:
            raise HTTPException(status_code=404, detail="账号不存在")

        cookies_str = account['cookies_str']
        if not cookies_str:
            raise HTTPException(status_code=400, detail="账号cookies为空")

        # 创建XianyuLive实例，传入正确的账号和用户归属
        from XianyuAutoAsync import XianyuLive
        xianyu_instance = XianyuLive(
            cookies_str=cookies_str,
            cookie_id=cookie_id,
            user_id=current_user['user_id']
        )

        try:
            logger.info(f"开始获取账号 {cookie_id} 第{page_number}页商品信息（每页{page_size}条）")
            result = await xianyu_instance.get_item_list_info(page_number, page_size)
        finally:
            await xianyu_instance.close_session()

        if not result.get('success'):
            error_message = result.get('error', '闲鱼商品分页获取失败，未返回明确原因')
            logger.error(f"获取商品信息失败: {error_message}")
            raise HTTPException(status_code=502, detail=error_message)

        current_count = result.get('current_count', 0)
        account_id = result.get('account_id') or cookie_id
        group_name = result.get('group_name', '在售')
        confirmed_empty = bool(result.get('confirmed_empty'))
        message = (
            f"闲鱼接口确认账号 {account_id} 的“{group_name}”分组当前为 0 件"
            if confirmed_empty
            else f"成功获取第{page_number}页 {current_count} 个商品"
        )
        logger.info(
            f"成功获取账号 {cookie_id} 第{page_number}页 {current_count} 个商品，"
            f"confirmed_empty={confirmed_empty}"
        )
        return {
            "success": True,
            "message": message,
            "page_number": page_number,
            "page_size": page_size,
            "current_count": current_count,
            "confirmed_empty": confirmed_empty,
            "group_name": group_name,
            "account_id": account_id
        }

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"获取账号商品信息异常: {str(e)}")
        raise HTTPException(status_code=500, detail=f"获取商品信息异常: {str(e)}")


# ------------------------- 用户设置接口 -------------------------

@app.get('/user-settings')
def get_user_settings(current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取当前用户的设置"""
    from app.db_manager import db_manager
    try:
        user_id = current_user['user_id']
        settings = db_manager.get_user_settings(user_id)
        return settings
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.put('/user-settings/{key}')
def update_user_setting(key: str, setting_data: dict, current_user: Dict[str, Any] = Depends(get_current_user)):
    """更新用户设置"""
    from app.db_manager import db_manager
    try:
        user_id = current_user['user_id']
        value = setting_data.get('value')
        description = setting_data.get('description', '')

        log_with_user('info', f"更新用户设置: {key} = {value}", current_user)

        success = db_manager.set_user_setting(user_id, key, value, description)
        if success:
            log_with_user('info', f"用户设置更新成功: {key}", current_user)
            return {'msg': 'setting updated', 'key': key, 'value': value}
        else:
            log_with_user('error', f"用户设置更新失败: {key}", current_user)
            raise HTTPException(status_code=400, detail='更新失败')
    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"更新用户设置异常: {key} - {str(e)}", current_user)
        raise HTTPException(status_code=500, detail=str(e))

@app.get('/user-settings/{key}')
def get_user_setting(key: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取用户特定设置"""
    from app.db_manager import db_manager
    try:
        user_id = current_user['user_id']
        setting = db_manager.get_user_setting(user_id, key)
        if setting:
            return setting
        else:
            raise HTTPException(status_code=404, detail='设置不存在')
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ------------------------- 管理员专用接口 -------------------------

@app.get('/admin/users')
def get_all_users(admin_user: Dict[str, Any] = Depends(require_admin)):
    """获取所有用户信息（管理员专用）"""
    from app.db_manager import db_manager
    try:
        log_with_user('info', "查询所有用户信息", admin_user)
        users = db_manager.get_all_users()

        # 为每个用户添加统计信息
        for user in users:
            user_id = user['id']
            # 统计用户的Cookie数量
            user_cookies = db_manager.get_all_cookies(user_id)
            user['cookie_count'] = len(user_cookies)

            # 统计用户的卡券数量
            user_cards = db_manager.get_all_cards(user_id)
            user['card_count'] = len(user_cards) if user_cards else 0

            # 隐藏密码字段
            if 'password_hash' in user:
                del user['password_hash']

        log_with_user('info', f"返回用户信息，共 {len(users)} 个用户", admin_user)
        return {"users": users}
    except Exception as e:
        log_with_user('error', f"获取用户信息失败: {str(e)}", admin_user)
        raise HTTPException(status_code=500, detail=str(e))

@app.delete('/admin/users/{user_id}')
def delete_user(user_id: int, admin_user: Dict[str, Any] = Depends(require_admin)):
    """删除用户（管理员专用）"""
    from app.db_manager import db_manager
    try:
        # 不能删除管理员自己
        if user_id == admin_user['user_id']:
            log_with_user('warning', "尝试删除管理员自己", admin_user)
            raise HTTPException(status_code=400, detail="不能删除管理员自己")

        # 获取要删除的用户信息
        user_to_delete = db_manager.get_user_by_id(user_id)
        if not user_to_delete:
            raise HTTPException(status_code=404, detail="用户不存在")

        log_with_user('info', f"准备删除用户: {user_to_delete['username']} (ID: {user_id})", admin_user)

        # 删除用户及其相关数据
        success = db_manager.delete_user_and_data(user_id)

        if success:
            log_with_user('info', f"用户删除成功: {user_to_delete['username']} (ID: {user_id})", admin_user)
            return {"message": f"用户 {user_to_delete['username']} 删除成功"}
        else:
            log_with_user('error', f"用户删除失败: {user_to_delete['username']} (ID: {user_id})", admin_user)
            raise HTTPException(status_code=400, detail="删除失败")
    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"删除用户异常: {str(e)}", admin_user)
        raise HTTPException(status_code=500, detail=str(e))

@app.get('/admin/risk-control-logs')
async def get_admin_risk_control_logs(
    cookie_id: str = None,
    limit: int = 100,
    offset: int = 0,
    admin_user: Dict[str, Any] = Depends(require_admin)
):
    """获取风控日志（管理员专用）"""
    try:
        log_with_user('info', f"查询风控日志: cookie_id={cookie_id}, limit={limit}, offset={offset}", admin_user)

        # 获取风控日志
        logs = db_manager.get_risk_control_logs(cookie_id=cookie_id, limit=limit, offset=offset)
        total_count = db_manager.get_risk_control_logs_count(cookie_id=cookie_id)

        log_with_user('info', f"风控日志查询成功，共 {len(logs)} 条记录，总计 {total_count} 条", admin_user)

        return {
            "success": True,
            "data": logs,
            "total": total_count,
            "limit": limit,
            "offset": offset
        }

    except Exception as e:
        log_with_user('error', f"查询风控日志失败: {str(e)}", admin_user)
        return {"success": False, "message": f"查询失败: {str(e)}", "data": [], "total": 0}


@app.get('/admin/cookies')
def get_admin_cookies(admin_user: Dict[str, Any] = Depends(require_admin)):
    """获取所有Cookie信息（管理员专用）"""
    try:
        log_with_user('info', "查询所有Cookie信息", admin_user)

        if cookie_manager.manager is None:
            return {
                "success": True,
                "cookies": [],
                "message": "CookieManager 未就绪"
            }

        # 获取所有用户的cookies
        from app.db_manager import db_manager
        all_users = db_manager.get_all_users()
        all_cookies = []

        for user in all_users:
            user_id = user['id']
            user_cookies = db_manager.get_all_cookies(user_id)
            for cookie_id, cookie_value in user_cookies.items():
                # 获取cookie详细信息
                cookie_details = db_manager.get_cookie_details(cookie_id)
                cookie_info = {
                    'cookie_id': cookie_id,
                    'user_id': user_id,
                    'username': user['username'],
                    'nickname': cookie_details.get('remark', '') if cookie_details else '',
                    'enabled': cookie_manager.manager.get_cookie_status(cookie_id)
                }
                all_cookies.append(cookie_info)

        log_with_user('info', f"获取到 {len(all_cookies)} 个Cookie", admin_user)
        return {
            "success": True,
            "cookies": all_cookies,
            "total": len(all_cookies)
        }

    except Exception as e:
        log_with_user('error', f"获取Cookie信息失败: {str(e)}", admin_user)
        return {
            "success": False,
            "cookies": [],
            "message": f"获取失败: {str(e)}"
        }


@app.get('/admin/logs')
def get_system_logs(admin_user: Dict[str, Any] = Depends(require_admin),
                   lines: int = 100,
                   level: str = None):
    """获取系统日志（管理员专用）"""
    import os
    import glob

    try:
        log_with_user('info', f"查询系统日志，行数: {lines}, 级别: {level}", admin_user)

        # 查找日志文件
        log_files = glob.glob("logs/xianyu_*.log")
        logger.info(f"找到日志文件: {log_files}")

        if not log_files:
            logger.warning("未找到日志文件")
            return {"logs": [], "message": "未找到日志文件", "success": False}

        # 获取最新的日志文件
        latest_log_file = max(log_files, key=os.path.getctime)
        logger.info(f"使用最新日志文件: {latest_log_file}")

        logs = []
        try:
            with open(latest_log_file, 'r', encoding='utf-8') as f:
                all_lines = f.readlines()
                logger.info(f"读取到 {len(all_lines)} 行日志")

                # 如果指定了日志级别，进行过滤
                if level:
                    filtered_lines = [line for line in all_lines if f"| {level.upper()} |" in line]
                    logger.info(f"按级别 {level} 过滤后剩余 {len(filtered_lines)} 行")
                else:
                    filtered_lines = all_lines

                # 获取最后N行
                recent_lines = filtered_lines[-lines:] if len(filtered_lines) > lines else filtered_lines
                logger.info(f"取最后 {len(recent_lines)} 行日志")

                for line in recent_lines:
                    logs.append(line.strip())

        except Exception as e:
            logger.error(f"读取日志文件失败: {str(e)}")
            log_with_user('error', f"读取日志文件失败: {str(e)}", admin_user)
            return {"logs": [], "message": f"读取日志文件失败: {str(e)}", "success": False}

        log_with_user('info', f"返回日志记录 {len(logs)} 条", admin_user)
        logger.info(f"成功返回 {len(logs)} 条日志记录")

        return {
            "logs": logs,
            "log_file": latest_log_file,
            "total_lines": len(logs),
            "success": True
        }

    except Exception as e:
        logger.error(f"获取系统日志失败: {str(e)}")
        log_with_user('error', f"获取系统日志失败: {str(e)}", admin_user)
        return {"logs": [], "message": f"获取系统日志失败: {str(e)}", "success": False}

@app.get('/admin/log-files')
def list_log_files(admin_user: Dict[str, Any] = Depends(require_admin)):
    """列出所有可用的系统日志文件"""
    import os
    import glob
    from datetime import datetime

    try:
        log_with_user('info', "查询日志文件列表", admin_user)

        log_dir = "logs"
        if not os.path.exists(log_dir):
            logger.warning("日志目录不存在")
            return {"success": True, "files": []}

        log_pattern = os.path.join(log_dir, "xianyu_*.log")
        log_files = glob.glob(log_pattern)

        files_info = []
        for file_path in log_files:
            try:
                stat_info = os.stat(file_path)
                files_info.append({
                    "name": os.path.basename(file_path),
                    "size": stat_info.st_size,
                    "modified_at": datetime.fromtimestamp(stat_info.st_mtime).isoformat(),
                    "modified_ts": stat_info.st_mtime
                })
            except OSError as e:
                logger.warning(f"读取日志文件信息失败 {file_path}: {e}")

        # 按修改时间倒序排序
        files_info.sort(key=lambda item: item.get("modified_ts", 0), reverse=True)

        logger.info(f"返回日志文件列表，共 {len(files_info)} 个文件")
        return {"success": True, "files": files_info}

    except Exception as e:
        logger.error(f"获取日志文件列表失败: {str(e)}")
        log_with_user('error', f"获取日志文件列表失败: {str(e)}", admin_user)
        raise HTTPException(status_code=500, detail=str(e))

@app.get('/admin/logs/export')
def export_log_file(file: str, admin_user: Dict[str, Any] = Depends(require_admin)):
    """导出指定的日志文件"""
    import os
    from fastapi.responses import StreamingResponse

    try:
        if not file:
            raise HTTPException(status_code=400, detail="缺少文件参数")

        safe_name = os.path.basename(file)
        log_dir = os.path.abspath("logs")
        target_path = os.path.abspath(os.path.join(log_dir, safe_name))

        # 防止目录遍历
        if not target_path.startswith(log_dir):
            log_with_user('warning', f"尝试访问非法日志文件: {file}", admin_user)
            raise HTTPException(status_code=400, detail="非法的日志文件路径")

        if not os.path.exists(target_path):
            log_with_user('warning', f"日志文件不存在: {file}", admin_user)
            raise HTTPException(status_code=404, detail="日志文件不存在")

        log_with_user('info', f"导出日志文件: {safe_name}", admin_user)
        def iter_file(path: str):
            file_handle = open(path, 'rb')
            try:
                while True:
                    chunk = file_handle.read(8192)
                    if not chunk:
                        break
                    yield chunk
            finally:
                file_handle.close()

        headers = {
            "Content-Disposition": f'attachment; filename="{safe_name}"'
        }
        return StreamingResponse(
            iter_file(target_path),
            media_type='text/plain; charset=utf-8',
            headers=headers
        )

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"导出日志文件失败: {str(e)}")
        log_with_user('error', f"导出日志文件失败: {str(e)}", admin_user)
        raise HTTPException(status_code=500, detail=str(e))

@app.get('/admin/stats')
def get_system_stats(admin_user: Dict[str, Any] = Depends(require_admin)):
    """获取系统统计信息（管理员专用）"""
    from app.db_manager import db_manager
    try:
        log_with_user('info', "查询系统统计信息", admin_user)

        # 用户统计
        all_users = db_manager.get_all_users()
        total_users = len(all_users)

        # Cookie统计
        all_cookies = db_manager.get_all_cookies()
        total_cookies = len(all_cookies)
        
        # 活跃账号统计（启用状态的账号）
        active_cookies = 0
        for cookie_id in all_cookies.keys():
            status = db_manager.get_cookie_status(cookie_id)
            if status:
                active_cookies += 1

        # 卡券统计
        all_cards = db_manager.get_all_cards()
        total_cards = len(all_cards) if all_cards else 0

        # 关键词统计
        all_keywords = db_manager.get_all_keywords()
        total_keywords = sum(len(kw_list) for kw_list in all_keywords.values())

        # 订单统计
        total_orders = 0
        try:
            orders = db_manager.get_all_orders()
            total_orders = len(orders) if orders else 0
        except:
            pass

        stats = {
            "total_users": total_users,
            "total_cookies": total_cookies,
            "active_cookies": active_cookies,
            "total_cards": total_cards,
            "total_keywords": total_keywords,
            "total_orders": total_orders
        }

        log_with_user('info', f"系统统计信息查询完成: {stats}", admin_user)
        return stats

    except Exception as e:
        log_with_user('error', f"获取系统统计信息失败: {str(e)}", admin_user)
        raise HTTPException(status_code=500, detail=str(e))

# ------------------------- BI报表分析接口 -------------------------

@app.get('/analytics/orders')
def get_order_analytics(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """
    获取订单分析数据（BI报表）

    Args:
        start_date: 开始日期 (格式: YYYY-MM-DD)
        end_date: 结束日期 (格式: YYYY-MM-DD)
    """
    from app.db_manager import db_manager
    try:
        log_with_user('info', f"查询订单分析数据: {start_date} - {end_date}", current_user)

        # 获取当前用户的ID
        user_id = current_user['user_id']

        # 定义有效订单状态（只统计这几种状态）
        valid_statuses = ['pending_ship', 'shipped', 'completed']

        # 调用数据库分析函数，传入包含状态
        analytics_data = db_manager.get_order_analytics(
            start_date=start_date,
            end_date=end_date,
            user_id=user_id,
            include_statuses=valid_statuses
        )

        if 'error' in analytics_data:
            log_with_user('error', f"获取订单分析数据失败: {analytics_data['error']}", current_user)
            raise HTTPException(status_code=500, detail=analytics_data['error'])

        log_with_user('info', "订单分析数据查询成功", current_user)
        return analytics_data

    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"获取订单分析数据失败: {str(e)}", current_user)
        raise HTTPException(status_code=500, detail=str(e))

@app.get('/analytics/orders/valid')
def get_valid_orders(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """
    获取有效订单详情列表（用于统计中的订单明细）

    Args:
        start_date: 开始日期 (格式: YYYY-MM-DD)
        end_date: 结束日期 (格式: YYYY-MM-DD)
    """
    from app.db_manager import db_manager
    try:
        log_with_user('info', f"查询有效订单列表: {start_date} - {end_date}", current_user)

        # 获取当前用户的ID
        user_id = current_user['user_id']

        # 定义有效订单状态
        valid_statuses = ['pending_ship', 'shipped', 'completed']

        # 调用数据库函数获取有效订单
        orders = db_manager.get_orders_for_analytics(
            start_date=start_date,
            end_date=end_date,
            user_id=user_id,
            include_statuses=valid_statuses
        )

        log_with_user('info', f"查询到 {len(orders)} 个有效订单", current_user)
        return {"orders": orders}

    except Exception as e:
        log_with_user('error', f"获取有效订单列表失败: {str(e)}", current_user)
        raise HTTPException(status_code=500, detail=str(e))

# ------------------------- 指定商品回复接口 -------------------------

@app.get("/itemReplays")
def get_all_items(current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取当前用户的所有商品回复信息"""
    try:
        # 只返回当前用户的商品信息
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        all_items = []
        for cookie_id in user_cookies.keys():
            items = db_manager.get_itemReplays_by_cookie(cookie_id)
            all_items.extend(items)

        return {"items": all_items}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"获取商品回复信息失败: {str(e)}")

@app.get("/itemReplays/cookie/{cookie_id}")
def get_items_by_cookie(cookie_id: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取指定Cookie的商品信息"""
    try:
        # 检查cookie是否属于当前用户
        user_id = current_user['user_id']
        from app.db_manager import db_manager
        user_cookies = db_manager.get_all_cookies(user_id)

        if cookie_id not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限访问该Cookie")

        items = db_manager.get_itemReplays_by_cookie(cookie_id)
        return {"items": items}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"获取商品信息失败: {str(e)}")

@app.put("/item-reply/{cookie_id}/{item_id}")
def update_item_reply(
    cookie_id: str,
    item_id: str,
    data: dict,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """
    更新指定账号和商品的回复内容
    """
    try:
        user_id = current_user['user_id']
        from app.db_manager import db_manager

        # 验证cookie是否属于用户
        user_cookies = db_manager.get_all_cookies(user_id)
        if cookie_id not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限访问该Cookie")

        reply_content = data.get("reply_content", "").strip()
        if not reply_content:
            raise HTTPException(status_code=400, detail="回复内容不能为空")

        db_manager.update_item_reply(cookie_id=cookie_id, item_id=item_id, reply_content=reply_content)

        return {"message": "商品回复更新成功"}

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"更新商品回复失败: {str(e)}")

@app.delete("/item-reply/{cookie_id}/{item_id}")
def delete_item_reply(cookie_id: str, item_id: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """
    删除指定账号cookie_id和商品item_id的商品回复
    """
    try:
        user_id = current_user['user_id']
        user_cookies = db_manager.get_all_cookies(user_id)
        if cookie_id not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限访问该Cookie")

        success = db_manager.delete_item_reply(cookie_id, item_id)
        if not success:
            raise HTTPException(status_code=404, detail="商品回复不存在")

        return {"message": "商品回复删除成功"}

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"删除商品回复失败: {str(e)}")

class ItemToDelete(BaseModel):
    cookie_id: str
    item_id: str

class BatchDeleteRequest(BaseModel):
    items: List[ItemToDelete]

@app.delete("/item-reply/batch")
async def batch_delete_item_reply(
    req: BatchDeleteRequest,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """
    批量删除商品回复
    """
    user_id = current_user['user_id']
    from app.db_manager import db_manager

    # 先校验当前用户是否有权限删除每个cookie对应的回复
    user_cookies = db_manager.get_all_cookies(user_id)
    for item in req.items:
        if item.cookie_id not in user_cookies:
            raise HTTPException(status_code=403, detail=f"无权限访问Cookie {item.cookie_id}")

    result = db_manager.batch_delete_item_replies([item.dict() for item in req.items])
    return {
        "success_count": result["success_count"],
        "failed_count": result["failed_count"]
    }

@app.get("/item-reply/{cookie_id}/{item_id}")
def get_item_reply(cookie_id: str, item_id: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """
    获取指定账号cookie_id和商品item_id的商品回复内容
    """
    try:
        user_id = current_user['user_id']
        # 校验cookie_id是否属于当前用户
        user_cookies = db_manager.get_all_cookies(user_id)
        if cookie_id not in user_cookies:
            raise HTTPException(status_code=403, detail="无权限访问该Cookie")

        # 获取指定商品回复
        item_replies = db_manager.get_itemReplays_by_cookie(cookie_id)
        # 找对应item_id的回复
        item_reply = next((r for r in item_replies if r['item_id'] == item_id), None)

        if item_reply is None:
            raise HTTPException(status_code=404, detail="商品回复不存在")

        return item_reply

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"获取商品回复失败: {str(e)}")


# ------------------------- 数据库备份和恢复接口 -------------------------

@app.get('/admin/backup/download')
def download_database_backup(admin_user: Dict[str, Any] = Depends(require_admin)):
    """下载数据库备份文件（管理员专用）"""
    import os
    from fastapi.responses import FileResponse
    from datetime import datetime

    try:
        log_with_user('info', "请求下载数据库备份", admin_user)

        # 使用db_manager的实际数据库路径
        from app.db_manager import db_manager
        db_file_path = db_manager.db_path

        # 检查数据库文件是否存在
        if not os.path.exists(db_file_path):
            log_with_user('error', f"数据库文件不存在: {db_file_path}", admin_user)
            raise HTTPException(status_code=404, detail="数据库文件不存在")

        # 生成带时间戳的文件名
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        download_filename = f"xianyu_backup_{timestamp}.db"

        log_with_user('info', f"开始下载数据库备份: {download_filename}", admin_user)

        return FileResponse(
            path=db_file_path,
            filename=download_filename,
            media_type='application/octet-stream'
        )

    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"下载数据库备份失败: {str(e)}", admin_user)
        raise HTTPException(status_code=500, detail=str(e))

@app.post('/admin/backup/upload')
async def upload_database_backup(admin_user: Dict[str, Any] = Depends(require_admin),
                                backup_file: UploadFile = File(...)):
    """上传并恢复数据库备份文件（管理员专用）"""
    import os
    import shutil
    import sqlite3
    from datetime import datetime

    try:
        log_with_user('info', f"开始上传数据库备份: {backup_file.filename}", admin_user)

        # 验证文件类型
        if not backup_file.filename.endswith('.db'):
            log_with_user('warning', f"无效的备份文件类型: {backup_file.filename}", admin_user)
            raise HTTPException(status_code=400, detail="只支持.db格式的数据库文件")

        # 验证文件大小（限制100MB）
        content = await backup_file.read()
        if len(content) > 100 * 1024 * 1024:  # 100MB
            log_with_user('warning', f"备份文件过大: {len(content)} bytes", admin_user)
            raise HTTPException(status_code=400, detail="备份文件大小不能超过100MB")

        # 验证是否为有效的SQLite数据库文件
        temp_file_path = f"temp_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"

        try:
            # 保存临时文件
            with open(temp_file_path, 'wb') as temp_file:
                temp_file.write(content)

            # 验证数据库文件完整性
            conn = sqlite3.connect(temp_file_path)
            cursor = conn.cursor()
            cursor.execute("SELECT name FROM sqlite_master WHERE type='table';")
            tables = cursor.fetchall()
            conn.close()

            # 检查是否包含必要的表
            table_names = [table[0] for table in tables]
            required_tables = ['users', 'cookies']  # 最基本的表

            missing_tables = [table for table in required_tables if table not in table_names]
            if missing_tables:
                log_with_user('warning', f"备份文件缺少必要的表: {missing_tables}", admin_user)
                raise HTTPException(status_code=400, detail=f"备份文件不完整，缺少表: {', '.join(missing_tables)}")

            log_with_user('info', f"备份文件验证通过，包含 {len(table_names)} 个表", admin_user)

        except sqlite3.Error as e:
            log_with_user('error', f"备份文件验证失败: {str(e)}", admin_user)
            if os.path.exists(temp_file_path):
                os.remove(temp_file_path)
            raise HTTPException(status_code=400, detail="无效的数据库文件")

        # 备份当前数据库
        from app.db_manager import db_manager
        current_db_path = db_manager.db_path

        # 生成备份文件路径（与原数据库在同一目录）
        db_dir = os.path.dirname(current_db_path)
        backup_filename = f"xianyu_data_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
        backup_current_path = os.path.join(db_dir, backup_filename)

        if os.path.exists(current_db_path):
            shutil.copy2(current_db_path, backup_current_path)
            log_with_user('info', f"当前数据库已备份为: {backup_current_path}", admin_user)

        # 关闭当前数据库连接
        if hasattr(db_manager, 'conn') and db_manager.conn:
            db_manager.conn.close()
            log_with_user('info', "已关闭当前数据库连接", admin_user)

        # 替换数据库文件
        shutil.move(temp_file_path, current_db_path)
        log_with_user('info', f"数据库文件已替换: {current_db_path}", admin_user)

        # 重新初始化数据库连接（使用原有的db_path）
        db_manager.__init__(db_manager.db_path)
        log_with_user('info', "数据库连接已重新初始化", admin_user)

        # 验证新数据库
        try:
            test_users = db_manager.get_all_users()
            log_with_user('info', f"数据库恢复成功，包含 {len(test_users)} 个用户", admin_user)
        except Exception as e:
            log_with_user('error', f"数据库恢复后验证失败: {str(e)}", admin_user)
            # 如果验证失败，尝试恢复原数据库
            if os.path.exists(backup_current_path):
                shutil.copy2(backup_current_path, current_db_path)
                db_manager.__init__()
                log_with_user('info', "已恢复原数据库", admin_user)
            raise HTTPException(status_code=500, detail="数据库恢复失败，已回滚到原数据库")

        return {
            "success": True,
            "message": "数据库恢复成功",
            "backup_file": backup_current_path,
            "user_count": len(test_users)
        }

    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"上传数据库备份失败: {str(e)}", admin_user)
        # 清理临时文件
        if 'temp_file_path' in locals() and os.path.exists(temp_file_path):
            os.remove(temp_file_path)
        raise HTTPException(status_code=500, detail=str(e))

@app.get('/admin/backup/list')
def list_backup_files(admin_user: Dict[str, Any] = Depends(require_admin)):
    """列出服务器上的备份文件（管理员专用）"""
    import os
    import glob
    from datetime import datetime

    try:
        log_with_user('info', "查询备份文件列表", admin_user)

        # 查找备份文件（在data目录中）
        backup_files = glob.glob("data/xianyu_data_backup_*.db")

        backup_list = []
        for file_path in backup_files:
            try:
                stat = os.stat(file_path)
                backup_list.append({
                    'filename': os.path.basename(file_path),
                    'size': stat.st_size,
                    'size_mb': round(stat.st_size / (1024 * 1024), 2),
                    'created_time': datetime.fromtimestamp(stat.st_ctime).strftime('%Y-%m-%d %H:%M:%S'),
                    'modified_time': datetime.fromtimestamp(stat.st_mtime).strftime('%Y-%m-%d %H:%M:%S')
                })
            except Exception as e:
                log_with_user('warning', f"读取备份文件信息失败: {file_path} - {str(e)}", admin_user)

        # 按修改时间倒序排列
        backup_list.sort(key=lambda x: x['modified_time'], reverse=True)

        log_with_user('info', f"找到 {len(backup_list)} 个备份文件", admin_user)

        return {
            "backups": backup_list,
            "total": len(backup_list)
        }

    except Exception as e:
        log_with_user('error', f"查询备份文件列表失败: {str(e)}", admin_user)
        raise HTTPException(status_code=500, detail=str(e))


# ------------------------- 系统管理接口 -------------------------

@app.post('/admin/reload-cache')
async def reload_system_cache(admin_user: Dict[str, Any] = Depends(require_admin)):
    """刷新系统缓存（管理员专用）"""
    try:
        log_with_user('info', "刷新系统缓存", admin_user)
        
        # 这里可以添加实际的缓存刷新逻辑
        # 例如：重新加载配置、清理内存缓存等
        
        log_with_user('info', "系统缓存刷新成功", admin_user)
        return {"success": True, "message": "系统缓存已刷新"}
        
    except Exception as e:
        log_with_user('error', f"刷新系统缓存失败: {str(e)}", admin_user)
        raise HTTPException(status_code=500, detail=str(e))


# ------------------------- 数据管理接口 -------------------------

@app.get('/admin/data/{table_name}')
def get_table_data(table_name: str, admin_user: Dict[str, Any] = Depends(require_admin)):
    """获取指定表的所有数据（管理员专用）"""
    from app.db_manager import db_manager
    try:
        log_with_user('info', f"查询表数据: {table_name}", admin_user)

        # 验证表名安全性
        allowed_tables = [
            'users', 'cookies', 'cookie_status', 'keywords', 'default_replies', 'default_reply_records',
            'ai_reply_settings', 'ai_conversations', 'ai_item_cache', 'item_info',
            'message_notifications', 'cards', 'delivery_rules', 'notification_channels',
            'user_settings', 'system_settings', 'email_verifications', 'captcha_codes', 'orders', "item_replay",
            'risk_control_logs'
        ]

        if table_name not in allowed_tables:
            log_with_user('warning', f"尝试访问不允许的表: {table_name}", admin_user)
            raise HTTPException(status_code=400, detail="不允许访问该表")

        # 获取表数据
        data, columns = db_manager.get_table_data(table_name)

        log_with_user('info', f"表 {table_name} 查询成功，共 {len(data)} 条记录", admin_user)

        return {
            "success": True,
            "data": data,
            "columns": columns,
            "count": len(data)
        }

    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"查询表数据失败: {table_name} - {str(e)}", admin_user)
        raise HTTPException(status_code=500, detail=str(e))

@app.delete('/admin/data/{table_name}/{record_id}')
def delete_table_record(table_name: str, record_id: str, admin_user: Dict[str, Any] = Depends(require_admin)):
    """删除指定表的指定记录（管理员专用）"""
    from app.db_manager import db_manager
    try:
        log_with_user('info', f"删除表记录: {table_name}.{record_id}", admin_user)

        # 验证表名安全性
        allowed_tables = [
            'users', 'cookies', 'cookie_status', 'keywords', 'default_replies', 'default_reply_records',
            'ai_reply_settings', 'ai_conversations', 'ai_item_cache', 'item_info',
            'message_notifications', 'cards', 'delivery_rules', 'notification_channels',
            'user_settings', 'system_settings', 'email_verifications', 'captcha_codes', 'orders','item_replay'
        ]

        if table_name not in allowed_tables:
            log_with_user('warning', f"尝试删除不允许的表记录: {table_name}", admin_user)
            raise HTTPException(status_code=400, detail="不允许操作该表")

        # 特殊保护：不能删除管理员用户
        if table_name == 'users' and record_id == str(admin_user['user_id']):
            log_with_user('warning', "尝试删除管理员自己", admin_user)
            raise HTTPException(status_code=400, detail="不能删除管理员自己")

        # 删除记录
        success = db_manager.delete_table_record(table_name, record_id)

        if success:
            log_with_user('info', f"表记录删除成功: {table_name}.{record_id}", admin_user)
            return {"success": True, "message": "删除成功"}
        else:
            log_with_user('warning', f"表记录删除失败: {table_name}.{record_id}", admin_user)
            raise HTTPException(status_code=400, detail="删除失败，记录可能不存在")

    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"删除表记录异常: {table_name}.{record_id} - {str(e)}", admin_user)
        raise HTTPException(status_code=500, detail=str(e))

@app.delete('/admin/data/{table_name}')
def clear_table_data(table_name: str, admin_user: Dict[str, Any] = Depends(require_admin)):
    """清空指定表的所有数据（管理员专用）"""
    from app.db_manager import db_manager
    try:
        log_with_user('info', f"清空表数据: {table_name}", admin_user)

        # 验证表名安全性
        allowed_tables = [
            'cookies', 'cookie_status', 'keywords', 'default_replies', 'default_reply_records',
            'ai_reply_settings', 'ai_conversations', 'ai_item_cache', 'item_info',
            'message_notifications', 'cards', 'delivery_rules', 'notification_channels',
            'user_settings', 'system_settings', 'email_verifications', 'captcha_codes', 'orders', 'item_replay',
            'risk_control_logs'
        ]

        # 不允许清空用户表
        if table_name == 'users':
            log_with_user('warning', "尝试清空用户表", admin_user)
            raise HTTPException(status_code=400, detail="不允许清空用户表")

        if table_name not in allowed_tables:
            log_with_user('warning', f"尝试清空不允许的表: {table_name}", admin_user)
            raise HTTPException(status_code=400, detail="不允许清空该表")

        # 清空表数据
        success = db_manager.clear_table_data(table_name)

        if success:
            log_with_user('info', f"表数据清空成功: {table_name}", admin_user)
            return {"success": True, "message": "清空成功"}
        else:
            log_with_user('warning', f"表数据清空失败: {table_name}", admin_user)
            raise HTTPException(status_code=400, detail="清空失败")

    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"清空表数据异常: {table_name} - {str(e)}", admin_user)
        raise HTTPException(status_code=500, detail=str(e))


# 商品多规格管理API
@app.put("/items/{cookie_id}/{item_id}/multi-spec")
def update_item_multi_spec(cookie_id: str, item_id: str, spec_data: dict, current_user: Dict[str, Any] = Depends(get_current_user)):
    """更新商品的多规格状态"""
    try:
        from app.db_manager import db_manager
        if cookie_id not in db_manager.get_all_cookies(current_user['user_id']):
            raise HTTPException(status_code=403, detail="无权限操作该账号")

        is_multi_spec = spec_data.get('is_multi_spec', False)

        success = db_manager.update_item_multi_spec_status(cookie_id, item_id, is_multi_spec)

        if success:
            return {"message": f"商品多规格状态已{'开启' if is_multi_spec else '关闭'}"}
        else:
            raise HTTPException(status_code=404, detail="商品不存在")

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# 商品多数量发货管理API
@app.put("/items/{cookie_id}/{item_id}/multi-quantity-delivery")
def update_item_multi_quantity_delivery(cookie_id: str, item_id: str, delivery_data: dict, current_user: Dict[str, Any] = Depends(get_current_user)):
    """更新商品的多数量发货状态"""
    try:
        from app.db_manager import db_manager
        if cookie_id not in db_manager.get_all_cookies(current_user['user_id']):
            raise HTTPException(status_code=403, detail="无权限操作该账号")

        multi_quantity_delivery = delivery_data.get('multi_quantity_delivery', False)

        success = db_manager.update_item_multi_quantity_delivery_status(cookie_id, item_id, multi_quantity_delivery)

        if success:
            return {"message": f"商品多数量发货状态已{'开启' if multi_quantity_delivery else '关闭'}"}
        else:
            raise HTTPException(status_code=404, detail="商品不存在")

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))





# ==================== 订单管理接口 ====================

@app.get('/api/orders')
def get_user_orders(
    current_user: Dict[str, Any] = Depends(get_current_user),
    page: int = Query(1, ge=1, description="页码"),
    page_size: int = Query(20, ge=1, le=100, description="每页数量"),
    cookie_id: Optional[str] = Query(None, description="筛选Cookie ID"),
    status: Optional[str] = Query(None, description="筛选状态")
):
    """获取当前用户的订单信息（支持分页）"""
    try:
        from app.db_manager import db_manager

        user_id = current_user['user_id']
        log_with_user('info', f"查询用户订单信息 (page={page}, page_size={page_size})", current_user)

        # 获取用户的所有Cookie
        user_cookies = db_manager.get_all_cookies(user_id)

        # 如果指定了cookie_id筛选
        if cookie_id and cookie_id in user_cookies:
            user_cookies = {cookie_id: user_cookies[cookie_id]}

        # 获取所有订单数据
        all_orders = []
        # 各状态的全量计数，在状态筛选前累加
        status_counts: Dict[str, int] = {}
        # 先获取所有商品的 item_id 到 item_title 的映射
        item_titles = {}
        with db_manager.lock:
            cursor = db_manager.conn.cursor()
            cursor.execute('SELECT item_id, item_title FROM item_info')
            for row in cursor.fetchall():
                item_titles[row[0]] = row[1]

        for cid in user_cookies.keys():
            orders = db_manager.get_orders_by_cookie(cid, limit=1000)
            for order in orders:
                order['cookie_id'] = cid
                # 添加 item_title 字段
                order['item_title'] = item_titles.get(order.get('item_id'), '')
                # 状态计数在筛选之前统计，保证各标签数字始终是全量口径
                order_state = get_order_status(order)
                status_counts[order_state] = status_counts.get(order_state, 0) + 1
                status_counts['all'] = status_counts.get('all', 0) + 1
                # 状态筛选
                if status and order_state != status:
                    continue
                all_orders.append(order)

        # 按创建时间倒序排列
        all_orders.sort(key=lambda x: x.get('created_at', ''), reverse=True)

        # 分页处理
        total = len(all_orders)
        total_pages = (total + page_size - 1) // page_size
        start_idx = (page - 1) * page_size
        end_idx = start_idx + page_size
        paginated_orders = all_orders[start_idx:end_idx]

        log_with_user('info', f"用户订单查询成功，共 {total} 条记录，第 {page}/{total_pages} 页", current_user)
        return {
            "success": True,
            "data": paginated_orders,
            "total": total,
            "page": page,
            "page_size": page_size,
            "total_pages": total_pages,
            # 各状态的全量条数，供前端标签显示，不受分页和筛选影响
            "status_counts": status_counts
        }

    except Exception as e:
        log_with_user('error', f"查询用户订单失败: {str(e)}", current_user)
        raise HTTPException(status_code=500, detail=f"查询订单失败: {str(e)}")


@app.get('/api/seller-features')
@app.get('/api/orders/seller-features')  # 兼容旧前端缓存
async def get_seller_features(current_user: Dict[str, Any] = Depends(get_current_user)):
    """返回买家互动类功能的开关状态。

    这两项会对买家产生不可撤销的实际动作，订单页需要据此决定是否展示入口，
    避免用户点了才发现被后端拒绝。

    开关按账号存，所以返回逐账号的映射；顶层的两个布尔保留为「是否有任意账号
    开启」，供只需要粗粒度判断的地方使用（比如整块入口要不要出现）。
    """
    from app.db_manager import db_manager

    accounts = db_manager.get_all_cookies(current_user['user_id']) or {}
    per_account = {
        cid: db_manager.get_buyer_interaction_settings(cid)
        for cid in accounts
    }
    return {
        'accounts': per_account,
        'auto_rate_enabled': any(v['auto_rate_enabled'] for v in per_account.values()),
        'auto_flower_enabled': any(v['auto_flower_enabled'] for v in per_account.values()),
        'auto_thanks_enabled': any(v['auto_thanks_enabled'] for v in per_account.values()),
        'auto_rate_template': _rate_template(),
    }


class BuyerInteractionUpdate(BaseModel):
    auto_rate_enabled: Optional[bool] = None
    auto_flower_enabled: Optional[bool] = None
    auto_thanks_enabled: Optional[bool] = None


@app.put('/api/seller-features/{cookie_id}')
async def update_seller_features(
    cookie_id: str,
    payload: BuyerInteractionUpdate,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """按账号更新评价/求花开关。"""
    from app.db_manager import db_manager

    if cookie_id not in (db_manager.get_all_cookies(current_user['user_id']) or {}):
        raise HTTPException(status_code=403, detail="无权限操作该账号")

    db_manager.update_buyer_interaction_settings(
        cookie_id,
        auto_rate_enabled=payload.auto_rate_enabled,
        auto_flower_enabled=payload.auto_flower_enabled,
        auto_thanks_enabled=payload.auto_thanks_enabled,
    )
    log_with_user(
        'info',
        f"更新账号 {cookie_id} 买家互动开关: rate={payload.auto_rate_enabled}, "
        f"flower={payload.auto_flower_enabled}, thanks={payload.auto_thanks_enabled}",
        current_user
    )
    return {"success": True, **db_manager.get_buyer_interaction_settings(cookie_id)}


@app.get('/api/orders/{order_id}')
def get_order_detail(order_id: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """获取订单详情"""
    try:
        from app.db_manager import db_manager

        user_id = current_user['user_id']
        log_with_user('info', f"查询订单详情: {order_id}", current_user)

        # 获取用户的所有Cookie
        user_cookies = db_manager.get_all_cookies(user_id)

        # 在用户的订单中查找
        for cookie_id in user_cookies.keys():
            order = db_manager.get_order_by_id(order_id)
            if order and order.get('cookie_id') == cookie_id:
                log_with_user('info', f"订单详情查询成功: {order_id}", current_user)
                return {"success": True, "data": order}

        log_with_user('warning', f"订单不存在或无权访问: {order_id}", current_user)
        raise HTTPException(status_code=404, detail="订单不存在或无权访问")

    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"查询订单详情失败: {str(e)}", current_user)
        raise HTTPException(status_code=500, detail=f"查询订单详情失败: {str(e)}")


@app.delete('/api/orders/{order_id}')
def delete_order(order_id: str, current_user: Dict[str, Any] = Depends(get_current_user)):
    """删除订单"""
    try:
        from app.db_manager import db_manager

        user_id = current_user['user_id']
        log_with_user('info', f"删除订单: {order_id}", current_user)

        # 获取用户的所有Cookie
        user_cookies = db_manager.get_all_cookies(user_id)

        # 验证订单属于当前用户
        order = db_manager.get_order_by_id(order_id)
        if not order:
            raise HTTPException(status_code=404, detail="订单不存在")

        if order.get('cookie_id') not in user_cookies:
            raise HTTPException(status_code=403, detail="无权删除此订单")

        # 删除订单
        success = db_manager.delete_order(order_id)
        if success:
            log_with_user('info', f"订单删除成功: {order_id}", current_user)
            return {"success": True, "message": "删除成功"}
        else:
            raise HTTPException(status_code=500, detail="删除失败")

    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"删除订单失败: {str(e)}", current_user)
        raise HTTPException(status_code=500, detail=f"删除订单失败: {str(e)}")


@app.post('/api/orders/{order_id}/refresh')
async def refresh_single_order(
    order_id: str,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """刷新单条订单状态"""
    try:
        from app.db_manager import db_manager
        from utils.order_fetcher_optimized import process_orders_batch

        user_id = current_user['user_id']
        log_with_user('info', f"刷新单条订单: {order_id}", current_user)

        # 获取用户的所有Cookie
        user_cookies = db_manager.get_all_cookies(user_id)

        # 验证订单存在且属于当前用户
        order = db_manager.get_order_by_id(order_id)
        if not order:
            raise HTTPException(status_code=404, detail="订单不存在")

        cookie_id = order.get('cookie_id')
        if not cookie_id or cookie_id not in user_cookies:
            raise HTTPException(status_code=403, detail="无权刷新此订单")

        cookies_str = user_cookies[cookie_id]
        if not cookies_str:
            raise HTTPException(status_code=400, detail="Cookie无效")

        # 调用批量刷新函数处理单条订单
        batch_results = await process_orders_batch(
            order_ids=[order_id],
            cookie_id=cookie_id,
            cookie_string=cookies_str,
            max_concurrent=1,
            timeout=30,
            headless=True,
            use_pool=True,
            force_refresh=True
        )

        if not batch_results or len(batch_results) == 0:
            raise HTTPException(status_code=500, detail="刷新失败")

        result = batch_results[0]
        if result.get('error'):
            raise HTTPException(status_code=500, detail=f"刷新失败: {result.get('error')}")

        order_status = normalize_order_status(
            result.get('order_status'),
            result.get('status_text') or result.get('dom_status') or result.get('api_status'),
        )

        # 更新数据库
        db_manager.insert_or_update_order(
            order_id=order_id,
            item_id=result.get('item_id') or None,
            buyer_id=result.get('buyer_id') or None,
            spec_name=result.get('spec_name') or None,
            spec_value=result.get('spec_value') or None,
            quantity=result.get('quantity') or None,
            amount=result.get('amount') or None,
            order_status=order_status,
            cookie_id=cookie_id,
            receiver_name=result.get('receiver_name') or None,
            receiver_phone=result.get('receiver_phone') or None,
            receiver_address=result.get('receiver_address') or None,
        )

        log_with_user('info', f"订单刷新成功: {order_id}, 新状态: {order_status}", current_user)
        return JSONResponse({
            "success": True,
            "message": "订单刷新成功",
            "data": {
                "order_id": order_id,
                "order_status": order_status,
            }
        })

    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"刷新订单失败: {str(e)}", current_user)
        raise HTTPException(status_code=500, detail=f"刷新订单失败: {str(e)}")


def check_order_data_completeness(order: Dict[str, Any]) -> bool:
    """
    检查订单数据是否完整

    Args:
        order: 订单数据字典

    Returns:
        True表示数据完整，False表示需要刷新
    """
    # 检查关键字段是否为空或为'unknown'
    incomplete_conditions = [
        not order.get('receiver_name') or order.get('receiver_name') == 'unknown',
        not order.get('receiver_phone') or order.get('receiver_phone') == 'unknown',
        not order.get('receiver_address') or order.get('receiver_address') == 'unknown',
        order.get('order_status') == 'unknown',
        not order.get('buyer_id') or order.get('buyer_id') == 'unknown',
    ]

    return not any(incomplete_conditions)


@app.put('/api/orders/{order_id}')
async def update_order(
    order_id: str,
    update_data: dict,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """
    更新订单信息
    自动检查订单数据完整性，如数据不完整则通过 Playwright 从订单详情页获取最新完整数据
    获取完整信息包括：订单ID、商品ID、买家ID、规格、数量、金额、订单状态、收货人信息
    """
    try:
        from app.db_manager import db_manager
        from utils.order_fetcher_optimized import fetch_order_complete

        user_id = current_user['user_id']
        log_with_user('info', f"更新订单: {order_id}, 数据: {update_data}", current_user)

        # 获取用户的所有Cookie
        user_cookies = db_manager.get_all_cookies(user_id)

        # 验证订单属于当前用户
        order = db_manager.get_order_by_id(order_id)
        if not order:
            raise HTTPException(status_code=404, detail="订单不存在")

        if order.get('cookie_id') not in user_cookies:
            raise HTTPException(status_code=403, detail="无权修改此订单")

        # 检查订单数据完整性
        is_complete = check_order_data_completeness(order)

        if not is_complete:
            log_with_user('info', f"订单 {order_id} 数据不完整，开始使用Playwright获取完整数据", current_user)

            # 获取该订单对应的Cookie字符串
            cookie_id = order.get('cookie_id')
            cookie_string = user_cookies.get(cookie_id)

            if cookie_string:

                try:
                    # 使用优化后的合并函数：一次浏览器访问获取所有数据
                    log_with_user('info', f"使用优化方法获取订单 {order_id} 的完整数据", current_user)

                    complete_result = await fetch_order_complete(
                        order_id=order_id,
                        cookie_id=cookie_id,
                        cookie_string=cookie_string,
                        timeout=30,
                        headless=True,
                        use_pool=True  # 使用浏览器池
                    )

                    if complete_result:
                        log_with_user('info', f"成功获取订单 {order_id} 的完整数据（一次浏览器调用）", current_user)

                        order_status = normalize_order_status(
                            complete_result.get('order_status'),
                            complete_result.get('status_text')
                            or complete_result.get('dom_status')
                            or complete_result.get('api_status'),
                        )

                        # 构建要更新的完整数据
                        refresh_data = {
                            'order_id': order_id,
                            'item_id': complete_result.get('item_id') or order.get('item_id'),
                            'buyer_id': complete_result.get('buyer_id') or order.get('buyer_id'),
                            'order_status': order_status or order.get('order_status'),
                            'spec_name': complete_result.get('spec_name') or None,
                            'spec_value': complete_result.get('spec_value') or None,
                            'quantity': complete_result.get('quantity') or None,
                            'amount': complete_result.get('amount') or None,
                            'created_at': complete_result.get('order_time') or None,
                            'receiver_name': complete_result.get('receiver_name') or None,
                            'receiver_phone': complete_result.get('receiver_phone') or None,
                            'receiver_address': complete_result.get('receiver_address') or None,
                            'receiver_city': complete_result.get('receiver_city') or None
                        }

                        # 更新数据库
                        db_manager.insert_or_update_order(**refresh_data)
                        log_with_user('info', f"订单 {order_id} 完整数据已更新到数据库", current_user)
                    else:
                        log_with_user('warning', f"订单 {order_id} 详情获取失败，继续使用现有数据", current_user)

                except Exception as e:
                    log_with_user('error', f"获取订单 {order_id} 详情时出错: {str(e)}", current_user)
                    # 继续执行，即使刷新失败也允许用户手动更新
            else:
                log_with_user('warning', f"订单 {order_id} 的Cookie信息不完整，无法刷新", current_user)

        # 提取可更新的字段
        allowed_fields = {
            'item_id', 'buyer_id', 'spec_name', 'spec_value',
            'quantity', 'amount', 'order_status',
            'receiver_name', 'receiver_phone', 'receiver_address', 'receiver_city',
            'system_shipped', 'created_at'
        }

        # 只保留允许更新的字段
        filtered_data = {k: v for k, v in update_data.items() if k in allowed_fields}

        if not filtered_data:
            # 如果没有用户提供的更新数据
            if not is_complete:
                # 数据不完整，已经进行了自动刷新，返回刷新后的订单
                updated_order = db_manager.get_order_by_id(order_id)
                return {
                    "success": True,
                    "message": "订单数据已自动刷新",
                    "data": updated_order,
                    "refreshed": True
                }
            else:
                # 数据完整，直接返回当前订单信息
                updated_order = db_manager.get_order_by_id(order_id)
                return {
                    "success": True,
                    "message": "订单数据已是最新",
                    "data": updated_order,
                    "refreshed": False
                }

        # 应用用户提供的更新
        success = db_manager.insert_or_update_order(
            order_id=order_id,
            **filtered_data
        )

        if success:
            log_with_user('info', f"订单更新成功: {order_id}", current_user)
            # 返回更新后的订单
            updated_order = db_manager.get_order_by_id(order_id)
            return {
                "success": True,
                "message": "更新成功",
                "data": updated_order,
                "refreshed": not is_complete  # 标记是否进行了自动刷新
            }
        else:
            raise HTTPException(status_code=500, detail="更新失败")

    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"更新订单失败: {str(e)}", current_user)
        raise HTTPException(status_code=500, detail=f"更新订单失败: {str(e)}")


@app.post('/api/orders/refresh')
async def refresh_orders_status(
    cookie_id: Optional[str] = Form(None),
    status: Optional[str] = Form(None),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """
    智能刷新订单状态
    1. 从数据库获取订单列表（支持筛选）
    2. 对非'已发货'状态的订单，使用Playwright查询最新状态
    3. 更新数据库中有变化的订单
    """
    try:
        from app.db_manager import db_manager
        from utils.order_fetcher_optimized import process_orders_batch

        user_id = current_user['user_id']
        log_with_user('info', f"开始智能刷新订单状态（优化版：并发处理） (cookie_id={cookie_id}, status={status})", current_user)

        # 获取用户的所有Cookie
        user_cookies = db_manager.get_all_cookies(user_id)

        # 如果指定了cookie_id，只使用该Cookie
        if cookie_id:
            if cookie_id not in user_cookies:
                raise HTTPException(status_code=404, detail="Cookie不存在或无权访问")
            user_cookies = {cookie_id: user_cookies[cookie_id]}

        # 获取需要刷新的订单
        orders_to_refresh = []
        for cid in user_cookies.keys():
            # 获取该Cookie的所有订单
            orders = db_manager.get_orders_by_cookie(cid, limit=1000)

            # 筛选需要刷新的订单
            for order in orders:
                # 如果指定了状态筛选，只刷新该状态的订单
                current_order_status = get_order_status(order)
                if status and current_order_status != status:
                    continue

                # 稳定状态（已发货、交易成功、交易关闭）的订单不需要刷新。
                needs_refresh = not is_stable_order_status(current_order_status)

                if needs_refresh:
                    orders_to_refresh.append({
                        'order_id': order['order_id'],
                        'cookie_id': cid,
                        'current_status': current_order_status
                    })

        log_with_user('info', f"找到 {len(orders_to_refresh)} 个需要刷新的订单", current_user)

        if not orders_to_refresh:
            return JSONResponse({
                "success": True,
                "message": "没有需要刷新的订单",
                "summary": {
                    "total": 0,
                    "updated": 0,
                    "no_change": 0,
                    "failed": 0
                },
                "results": []
            })

        # 刷新订单信息（包括状态、买家ID、金额等）
        updated_count = 0
        failed_count = 0
        no_change_count = 0
        refresh_results = []

        # 按cookie_id分组订单（因为每个cookie需要单独的浏览器实例）
        orders_by_cookie = {}
        for order_info in orders_to_refresh:
            cid = order_info['cookie_id']
            if cid not in orders_by_cookie:
                orders_by_cookie[cid] = []
            orders_by_cookie[cid].append(order_info)

        # 对每个cookie的订单进行并发批量处理
        for cid, cookie_orders in orders_by_cookie.items():
            cookies_str = user_cookies[cid]
            if not cookies_str:
                log_with_user('warning', f"Cookie {cid} 的值为空，跳过", current_user)
                failed_count += len(cookie_orders)
                continue

            # 提取订单ID列表
            order_ids = [o['order_id'] for o in cookie_orders]
            log_with_user('info', f"使用并发处理Cookie {cid} 的 {len(order_ids)} 个订单", current_user)

            # 并发批量处理（一次浏览器调用获取所有数据）
            batch_results = await process_orders_batch(
                order_ids=order_ids,
                cookie_id=cid,
                cookie_string=cookies_str,
                max_concurrent=5,  # 并发5个
                timeout=30,
                headless=True,
                use_pool=True,  # 使用浏览器池
                force_refresh=True  # 强制刷新，跳过缓存检查
            )

            # 处理结果并更新数据库
            for i, result in enumerate(batch_results):
                order_info = cookie_orders[i]
                order_id = order_info['order_id']
                current_status = order_info['current_status']

                if result and not result.get('error'):
                    # 调试：打印API和DOM状态
                    api_status = result.get('api_status', 'N/A')
                    dom_status = result.get('dom_status', 'N/A')
                    log_with_user('debug', f"订单 {order_id} - API状态: {api_status}, DOM状态: {dom_status}", current_user)

                    order_status = normalize_order_status(
                        result.get('order_status'),
                        result.get('status_text') or dom_status or api_status,
                    )

                    # 更新数据库
                    success = db_manager.insert_or_update_order(
                        order_id=order_id,
                        item_id=result.get('item_id') or None,
                        buyer_id=result.get('buyer_id') or None,
                        spec_name=result.get('spec_name') or None,
                        spec_value=result.get('spec_value') or None,
                        quantity=result.get('quantity') or None,
                        amount=result.get('amount') or None,
                        order_status=order_status if order_status != current_status else None,
                        cookie_id=cid,
                        created_at=result.get('order_time') or None,
                        receiver_name=result.get('receiver_name') or None,
                        receiver_phone=result.get('receiver_phone') or None,
                        receiver_address=result.get('receiver_address') or None
                    )

                    if success:
                        # 检查是否有更新
                        has_changes = (
                            order_status != current_status or
                            result.get('buyer_id') or
                            result.get('amount')
                        )

                        if has_changes:
                            updated_count += 1
                            refresh_results.append({
                                'order_id': order_id,
                                'old_status': current_status,
                                'new_status': order_status,
                                'status_text': result.get('status_text', '')
                            })
                            log_with_user('info', f"订单 {order_id} 已更新 | {current_status} -> {order_status}", current_user)
                        else:
                            no_change_count += 1
                    else:
                        failed_count += 1
                        log_with_user('error', f"订单 {order_id} 更新失败", current_user)
                else:
                    failed_count += 1
                    error_msg = result.get('error', '未知错误') if result else '未知错误'
                    log_with_user('warning', f"订单 {order_id} 获取失败: {error_msg}", current_user)

        # 返回刷新结果
        log_with_user('info', f"订单刷新完成: 更新{updated_count}个, 无变化{no_change_count}个, 失败{failed_count}个", current_user)

        return JSONResponse({
            "success": True,
            "message": f"刷新完成: 更新{updated_count}个, 无变化{no_change_count}个, 失败{failed_count}个",
            "summary": {
                "total": len(orders_to_refresh),
                "updated": updated_count,
                "no_change": no_change_count,
                "failed": failed_count
            },
            "updated_orders": refresh_results
        })

    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"刷新订单状态失败: {str(e)}", current_user)
        raise HTTPException(status_code=500, detail=f"刷新订单状态失败: {str(e)}")


# 已取消：全量核对订单数据功能
# 现在使用更新订单状态接口进行单个订单的数据核查
# @app.post('/api/orders/verify-all')
# async def verify_all_orders(current_user: Dict[str, Any] = Depends(get_current_user)):
#     """
#     全量核对所有订单数据
#     通过 Playwright 访问每个订单的详情页，更新时间、收货人信息等
#     """
#     pass


@app.post('/api/orders/manual-ship')
async def manual_ship_orders(
    order_ids: List[str] = Body(..., description="订单ID列表"),
    ship_mode: str = Body(..., description="发货模式: status_only（仅修改发货状态）或 full_delivery（完整发货流程）"),
    custom_content: Optional[str] = Body(None, description="自定义发货内容（保留兼容）"),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """
    手动发货

    发货模式：
    - status_only: 仅在闲鱼标记为已发货（不发送卡券给买家）
    - full_delivery: 完整发货流程（匹配卡券、发送卡券给买家、标记发货状态）
    """
    try:
        from app.db_manager import db_manager
        from XianyuAutoAsync import XianyuLive
        import asyncio

        user_id = current_user['user_id']
        log_with_user('info', f"开始手动发货: 订单数量={len(order_ids)}, 模式={ship_mode}", current_user)

        # 验证发货模式
        if ship_mode not in ['status_only', 'full_delivery']:
            raise HTTPException(status_code=400, detail="发货模式必须是 status_only 或 full_delivery")

        # 获取用户的所有Cookie
        user_cookies = db_manager.get_all_cookies(user_id)

        success_count = 0
        failed_count = 0
        results = []

        # 遍历每个订单
        for order_id in order_ids:
            try:
                # 获取订单信息
                order = db_manager.get_order_by_id(order_id)
                if not order:
                    results.append({
                        'order_id': order_id,
                        'success': False,
                        'message': '订单不存在'
                    })
                    failed_count += 1
                    continue

                # 验证订单属于当前用户
                cookie_id = order.get('cookie_id')
                if cookie_id not in user_cookies:
                    results.append({
                        'order_id': order_id,
                        'success': False,
                        'message': '无权操作此订单'
                    })
                    failed_count += 1
                    continue

                item_id = order.get('item_id')
                buyer_id = order.get('buyer_id')

                if ship_mode == 'status_only':
                    # ====== 仅修改闲鱼发货状态 ======
                    if not item_id:
                        results.append({
                            'order_id': order_id,
                            'success': False,
                            'message': '订单缺少商品ID'
                        })
                        failed_count += 1
                        continue

                    # 获取cookies_str用于创建独立session
                    cookies_str = user_cookies.get(cookie_id)
                    if not cookies_str:
                        results.append({
                            'order_id': order_id,
                            'success': False,
                            'message': '无法获取账号Cookie信息'
                        })
                        failed_count += 1
                        continue

                    # 创建独立的aiohttp session（避免跨异步上下文问题）
                    import aiohttp
                    from app.secure_confirm import SecureConfirm

                    try:
                        async with aiohttp.ClientSession(
                            headers={'cookie': cookies_str},
                            timeout=aiohttp.ClientTimeout(total=30)
                        ) as session:
                            confirm = SecureConfirm(session, cookies_str, cookie_id, None)
                            confirm_result = await confirm.auto_confirm(order_id, item_id)

                        if confirm_result and confirm_result.get('success'):
                            # 更新本地数据库状态
                            db_manager.insert_or_update_order(
                                order_id=order_id,
                                order_status='shipped',
                                system_shipped=True
                            )
                            results.append({
                                'order_id': order_id,
                                'success': True,
                                'message': '已成功修改闲鱼发货状态'
                            })
                            success_count += 1
                        else:
                            error_msg = confirm_result.get('error', '未知错误') if confirm_result else '确认发货返回空结果'
                            results.append({
                                'order_id': order_id,
                                'success': False,
                                'message': f'修改发货状态失败: {error_msg}'
                            })
                            failed_count += 1
                    except Exception as e:
                        log_with_user('error', f"确认发货异常: {str(e)}", current_user)
                        results.append({
                            'order_id': order_id,
                            'success': False,
                            'message': f'确认发货异常: {str(e)}'
                        })
                        failed_count += 1

                elif ship_mode == 'full_delivery':
                    # ====== 完整发货流程：匹配卡券 + 发送卡券 + 修改状态 ======
                    if not item_id:
                        results.append({
                            'order_id': order_id,
                            'success': False,
                            'message': '订单缺少商品ID，无法匹配发货规则'
                        })
                        failed_count += 1
                        continue

                    if not buyer_id:
                        results.append({
                            'order_id': order_id,
                            'success': False,
                            'message': '订单缺少买家ID，无法发送卡券'
                        })
                        failed_count += 1
                        continue

                    # 必须有运行中的实例（需要WebSocket发送消息）
                    live_instance = XianyuLive.get_instance(cookie_id)
                    if not live_instance:
                        results.append({
                            'order_id': order_id,
                            'success': False,
                            'message': '该账号未在线运行，无法执行完整发货。请先启动账号。'
                        })
                        failed_count += 1
                        continue

                    if not live_instance.ws or live_instance.ws.closed:
                        results.append({
                            'order_id': order_id,
                            'success': False,
                            'message': '该账号WebSocket连接已断开，无法发送消息。请等待重连后重试。'
                        })
                        failed_count += 1
                        continue

                    # 查找与买家的chat_id（优先从订单记录获取，回退到AI对话记录）
                    chat_id = order.get('chat_id') or ''
                    if not chat_id:
                        chat_id = db_manager.find_chat_id_by_buyer(cookie_id, buyer_id)
                    if not chat_id:
                        results.append({
                            'order_id': order_id,
                            'success': False,
                            'message': '未找到与该买家的聊天记录，无法发送卡券消息。请等待买家发送消息后重试。'
                        })
                        failed_count += 1
                        continue

                    protection_result = await live_instance.apply_delivery_block_rules(
                        order_id=order_id,
                        buyer_id=buyer_id,
                        item_id=item_id,
                        chat_id=chat_id,
                        websocket=live_instance.ws,
                        owner_id=user_id,
                    )
                    if protection_result["action"] == "block":
                        results.append({
                            'order_id': order_id,
                            'success': False,
                            'message': (
                                f'发货已被规则“{protection_result["rule_name"]}”拦截：'
                                f'{protection_result["reason"]}'
                            )
                        })
                        failed_count += 1
                        continue
                    card_only_delivery = protection_result["action"] == "card_only"

                    # 检查多数量发货
                    quantity_to_send = 1
                    order_detail = None
                    multi_quantity_delivery = db_manager.get_item_multi_quantity_delivery_status(cookie_id, item_id)
                    if multi_quantity_delivery:
                        try:
                            order_detail = await live_instance.fetch_order_detail_info(order_id, item_id, buyer_id)
                            if order_detail and isinstance(order_detail, dict):
                                qty = int(order_detail.get('quantity') or 1)
                                if qty > 1:
                                    quantity_to_send = qty
                        except Exception as e:
                            log_with_user('warning', f"获取订单数量失败，使用默认数量1: {str(e)}", current_user)

                    # 获取卡券内容；确认发货必须在消息全部发送成功后执行。
                    delivery_contents = []
                    delivery_context = {}
                    try:
                        first_content = await live_instance._auto_delivery(
                            item_id,
                            '',
                            order_id,
                            buyer_id,
                            order_detail=order_detail,
                            delivery_context=delivery_context,
                            requested_item_quantity=quantity_to_send,
                        )
                    except Exception as e:
                        log_with_user('error', f"获取第1个卡券失败: {str(e)}", current_user)
                        first_content = None

                    if first_content:
                        delivery_contents.append(first_content)
                        quantity_to_send *= max(
                            1,
                            int(delivery_context.get("delivery_count") or 1),
                        )

                    for i in range(1, quantity_to_send):
                        if not delivery_contents:
                            break
                        try:
                            delivery_content = await live_instance._auto_delivery(
                                item_id,
                                '',
                                order_id,
                                buyer_id,
                                order_detail=order_detail,
                                delivery_context=delivery_context,
                            )
                            if delivery_content:
                                delivery_contents.append(delivery_content)
                        except Exception as e:
                            log_with_user('error', f"获取第{i+1}个卡券失败: {str(e)}", current_user)

                    if not delivery_contents:
                        results.append({
                            'order_id': order_id,
                            'success': False,
                            'message': '未匹配到发货规则或卡券获取失败'
                        })
                        failed_count += 1
                        continue

                    # 发送卡券内容给买家
                    sent_count = 0
                    send_errors = []
                    for idx, content in enumerate(delivery_contents):
                        try:
                            if content.startswith("__IMAGE_SEND__"):
                                image_data = content.replace("__IMAGE_SEND__", "")
                                card_id = None
                                if "|" in image_data:
                                    card_id_str, image_url = image_data.split("|", 1)
                                    try:
                                        card_id = int(card_id_str)
                                    except ValueError:
                                        card_id = None
                                else:
                                    image_url = image_data
                                await live_instance.send_image_msg(
                                    live_instance.ws, chat_id, buyer_id,
                                    image_url, card_id=card_id
                                )
                            else:
                                await live_instance.send_msg(
                                    live_instance.ws, chat_id, buyer_id, content
                                )

                            # 多条消息之间间隔1秒
                            if len(delivery_contents) > 1 and idx < len(delivery_contents) - 1:
                                await asyncio.sleep(1)
                            sent_count += 1
                        except Exception as e:
                            log_with_user('error', f"发送第{idx+1}条卡券消息失败: {str(e)}", current_user)
                            send_errors.append(f"第{idx + 1}条: {str(e)}")

                    acquired_all = len(delivery_contents) == quantity_to_send
                    sent_all = sent_count == len(delivery_contents)

                    if acquired_all and sent_all:
                        confirm_required = (
                            live_instance.is_auto_confirm_enabled()
                            and not card_only_delivery
                        )
                        platform_confirmed = False
                        confirm_error = None

                        if confirm_required:
                            try:
                                confirm_result = await live_instance.auto_confirm(order_id, item_id)
                                platform_confirmed = bool(confirm_result and confirm_result.get('success'))
                                if platform_confirmed:
                                    live_instance.confirmed_orders[order_id] = time.time()
                                else:
                                    confirm_error = (confirm_result or {}).get('error', '未知错误')
                            except Exception as e:
                                confirm_error = str(e)

                        live_instance.mark_delivery_sent(order_id, update_order_status=platform_confirmed)
                        update_values = {
                            'order_id': order_id,
                            'system_shipped': True
                        }
                        if platform_confirmed:
                            update_values['order_status'] = 'shipped'
                        db_manager.insert_or_update_order(**update_values)

                        if card_only_delivery:
                            results.append({
                                'order_id': order_id,
                                'success': True,
                                'message': (
                                    f'订单命中“{protection_result["rule_name"]}”并已关闭，'
                                    f'仅发送{sent_count}条卡券信息'
                                )
                            })
                            success_count += 1
                            continue

                        if confirm_required and not platform_confirmed:
                            results.append({
                                'order_id': order_id,
                                'success': False,
                                'message': f'卡券已全部发送，但闲鱼确认发货失败，请手动确认：{confirm_error}'
                            })
                            failed_count += 1
                            continue

                        results.append({
                            'order_id': order_id,
                            'success': True,
                            'message': f'完整发货成功，已发送{sent_count}条卡券信息给买家'
                        })
                        success_count += 1
                    else:
                        live_instance.delivery_blocked_orders.add(order_id)
                        live_instance.last_delivery_time[order_id] = time.time()
                        detail = f'应发{quantity_to_send}条，获取{len(delivery_contents)}条，成功发送{sent_count}条'
                        if send_errors:
                            detail += f'；{"；".join(send_errors)}'
                        results.append({
                            'order_id': order_id,
                            'success': False,
                            'message': f'完整发货未完成：{detail}。已停止自动重试，请人工核对'
                        })
                        failed_count += 1

            except Exception as e:
                results.append({
                    'order_id': order_id,
                    'success': False,
                    'message': str(e)
                })
                failed_count += 1
                log_with_user('error', f"发货订单 {order_id} 时发生异常: {str(e)}", current_user)

        log_with_user('info', f"手动发货完成: 成功{success_count}个, 失败{failed_count}个", current_user)

        return {
            "success": True,
            "message": f"发货完成: 成功{success_count}个, 失败{failed_count}个",
            "total": len(order_ids),
            "success_count": success_count,
            "failed_count": failed_count,
            "results": results
        }

    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"手动发货失败: {str(e)}", current_user)
        raise HTTPException(status_code=500, detail=f"手动发货失败: {str(e)}")


@app.post('/api/orders/import')
async def import_orders(
    orders: List[Dict[str, Any]] = Body(..., description="订单列表"),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """
    导入订单
    支持批量导入自定义订单数据
    """
    try:
        from app.db_manager import db_manager

        user_id = current_user['user_id']
        log_with_user('info', f"开始导入订单: 订单数量={len(orders)}", current_user)

        # 获取用户的所有Cookie
        user_cookies = db_manager.get_all_cookies(user_id)

        success_count = 0
        failed_count = 0
        results = []

        # 必需字段验证
        required_fields = ['order_id', 'cookie_id']
        optional_fields = [
            'item_id', 'item_title', 'item_price', 'item_image',
            'buyer_id',
            'receiver_name', 'receiver_phone', 'receiver_address', 'receiver_city',
            'status', 'status_text', 'order_time', 'pay_time',
            'quantity', 'amount'
        ]

        for order_data in orders:
            try:
                # 验证必需字段
                missing_fields = [f for f in required_fields if not order_data.get(f)]
                if missing_fields:
                    results.append({
                        'order_id': order_data.get('order_id', 'unknown'),
                        'success': False,
                        'message': f'缺少必需字段: {", ".join(missing_fields)}'
                    })
                    failed_count += 1
                    continue

                order_id = str(order_data['order_id'])
                cookie_id = str(order_data['cookie_id'])

                # 验证Cookie属于当前用户
                if cookie_id not in user_cookies:
                    results.append({
                        'order_id': order_id,
                        'success': False,
                        'message': '无权操作此账号的订单'
                    })
                    failed_count += 1
                    continue

                # 检查订单是否已存在
                existing_order = db_manager.get_order_by_id(order_id)

                # 准备订单数据，直接使用 insert_or_update_order 的参数名
                # 构建参数字典，只传递非 None 的值
                insert_params = {
                    'order_id': order_id,
                    'cookie_id': cookie_id
                }

                # 前端字段名 -> 数据库参数名映射
                param_mapping = {
                    'item_id': 'item_id',
                    'buyer_id': 'buyer_id',
                    'receiver_name': 'receiver_name',
                    'receiver_phone': 'receiver_phone',
                    'receiver_address': 'receiver_address',
                    'receiver_city': 'receiver_city',
                    'status': 'order_status',  # 注意：前端用 status，后端用 order_status
                    'status_text': 'status_text',
                    'order_time': 'order_time',
                    'pay_time': 'pay_time',
                    'quantity': 'quantity',
                    'amount': 'amount',
                    'item_title': 'item_title',
                    'item_price': 'item_price',
                    'item_image': 'item_image'
                }

                # 遍历订单数据，添加到参数字典
                for field, value in order_data.items():
                    if value is not None and field in param_mapping:
                        param_name = param_mapping[field]
                        insert_params[param_name] = value

                # 使用 insert_or_update_order 统一处理
                db_manager.insert_or_update_order(**insert_params)

                results.append({
                    'order_id': order_id,
                    'success': True,
                    'message': '订单已更新' if existing_order else '订单已导入'
                })

                success_count += 1

            except Exception as e:
                results.append({
                    'order_id': order_data.get('order_id', 'unknown'),
                    'success': False,
                    'message': str(e)
                })
                failed_count += 1
                log_with_user('error', f"导入订单时发生异常: {str(e)}", current_user)

        log_with_user('info', f"导入订单完成: 成功{success_count}个, 失败{failed_count}个", current_user)

        return {
            "success": True,
            "message": f"导入完成: 成功{success_count}个, 失败{failed_count}个",
            "total": len(orders),
            "success_count": success_count,
            "failed_count": failed_count,
            "results": results
        }

    except HTTPException:
        raise
    except Exception as e:
        log_with_user('error', f"导入订单失败: {str(e)}", current_user)
        raise HTTPException(status_code=500, detail=f"导入订单失败: {str(e)}")


# ==================== 前端 SPA Catch-All 路由 ====================
# 必须放在所有 API 路由之后，用于处理前端 SPA 的直接访问
# 这样用户直接访问 /dashboard、/accounts 等前端路由时，会返回 index.html
# 然后由 React Router 在客户端处理路由

# 定义后端 API 的一级路径。未匹配的 API 路径必须返回 JSON 404，不能回退到 SPA。
API_ROOTS = {
    'admin', 'ai-reply-settings', 'ai-reply-test', 'analytics', 'api', 'backup',
    'cards', 'change-admin-password', 'change-password', 'cookie', 'cookies',
    'blacklist', 'debug', 'default-replies', 'delivery-block-rules', 'delivery-rules', 'face-verification',
    'generate-captcha', 'geetest', 'health', 'item-reply', 'itemReplays', 'items',
    'item-delivery-configs',
    'keywords', 'keywords-export', 'keywords-import', 'keywords-with-item-id',
    'keywords-with-type', 'login', 'login-info-settings', 'login-info-status',
    'logout', 'logs', 'message-notifications', 'notification-channels',
    'password-login', 'qr-login', 'quick-phrases', 'register', 'registration-settings',
    'registration-status', 'risk-control-logs', 'send-message',
    'send-verification-code', 'static', 'system', 'system-settings',
    'upload-image', 'user-settings', 'verify', 'verify-captcha', 'xianyu'
}


# ------------------------- 卖家端订单同步与互动 -------------------------

# 评价和求花都会产生不可撤销的对外动作，默认关闭，需显式开启后才允许调用
AUTO_RATE_SETTING_KEY = 'auto_rate_enabled'
AUTO_FLOWER_SETTING_KEY = 'auto_flower_enabled'
# 默认评价文案。订单页打开评价框时预填，避免每单重复手打，仍可逐单改。
RATE_TEMPLATE_SETTING_KEY = 'auto_rate_template'
DEFAULT_RATE_TEMPLATE = '宝贝收到了，很满意，感谢老板，欢迎下次再来！'


def _seller_feature_enabled(key: str) -> bool:
    from app.db_manager import db_manager
    return str(db_manager.get_system_setting(key) or '').strip().lower() in ('1', 'true', 'yes')


@app.get('/api/announcement')
async def get_announcement(
    force: bool = False,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """全局公告与版本检查。

    数据来自系统设置里配置的公网 JSON 地址，后端代拉并缓存 10 分钟。
    远端不可用时沿用上次结果，保证页面不受影响。
    """
    from app.announcement import get_announcement_payload
    return await get_announcement_payload(force=force)


def _rate_template() -> str:
    """默认评价文案，没配过就用内置的一条。"""
    from app.db_manager import db_manager
    return str(db_manager.get_system_setting(RATE_TEMPLATE_SETTING_KEY) or '').strip() or DEFAULT_RATE_TEMPLATE


def _resolve_order_cookie(order_id: str, user_cookies: Dict[str, str]) -> tuple:
    """校验订单归属并返回 (cookie_id, cookies_str)。"""
    from app.db_manager import db_manager

    order = db_manager.get_order_by_id(str(order_id))
    if not order:
        raise HTTPException(status_code=404, detail=f"订单不存在: {order_id}")

    cookie_id = order.get('cookie_id')
    if cookie_id not in user_cookies:
        raise HTTPException(status_code=403, detail=f"无权操作订单: {order_id}")

    cookies_str = user_cookies.get(cookie_id)
    if not cookies_str:
        raise HTTPException(status_code=400, detail=f"账号缺少 Cookie: {cookie_id}")
    return cookie_id, cookies_str


@app.post('/api/orders/sync-sold')
async def sync_sold_orders(
    cookie_id: Optional[str] = Form(None),
    days: int = Form(7),
    include_refund_history: bool = Form(True),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """通过卖家端接口全量同步卖出订单。

    消息驱动只能捕获监听在线期间的订单，且金额按商品挂牌价推算并不准确；
    这里直接拉取接口的真实成交数据做兜底对账。

    不按时间筛选 —— 卖家端只保留近期订单，加时间条件会切掉边界订单。
    ``days`` 仅为兼容旧调用保留。开启 ``include_refund_history`` 时会用退款列表
    补全订单列表已查不到的历史归档单。
    """
    from app.db_manager import db_manager
    from utils.seller_order_sync import sync_account_orders

    user_id = current_user['user_id']
    user_cookies = db_manager.get_all_cookies(user_id)
    if cookie_id:
        if cookie_id not in user_cookies:
            raise HTTPException(status_code=404, detail="Cookie不存在或无权访问")
        user_cookies = {cookie_id: user_cookies[cookie_id]}

    summary = {
        'total': 0, 'saved': 0, 'failed': 0, 'from_refund': 0,
        'incomplete_accounts': [], 'accounts': {},
    }

    for cid, cookies_str in user_cookies.items():
        if not cookies_str:
            continue
        try:
            res = await sync_account_orders(
                cid, cookies_str, include_refund_history=include_refund_history
            )
        except Exception as e:
            log_with_user('error', f"账号 {cid} 卖出订单同步异常: {e}", current_user)
            summary['accounts'][cid] = {'error': str(e)}
            continue

        summary['total'] += res['total']
        summary['saved'] += res['saved']
        summary['failed'] += res['failed']
        summary['from_refund'] += res['from_refund']
        if not res['complete']:
            summary['incomplete_accounts'].append(cid)
        summary['accounts'][cid] = {
            'total': res['total'],
            'saved': res['saved'],
            'failed': res['failed'],
            'from_refund': res['from_refund'],
            'expected': res['expected'],
            'complete': res['complete'],
        }

    log_with_user(
        'info',
        f"卖出订单同步完成: 共 {summary['total']} 单，成功 {summary['saved']}，失败 {summary['failed']}",
        current_user,
    )

    message = f"同步完成: 共 {summary['total']} 单，成功 {summary['saved']}"
    if summary['from_refund']:
        message += f"（含历史退款单 {summary['from_refund']} 单）"
    if summary['failed']:
        message += f"，失败 {summary['failed']}"
    if summary['incomplete_accounts']:
        # 拉取数量少于服务端角标，提示用户而不是静默少数据
        message += f"；部分账号可能未拉全: {', '.join(summary['incomplete_accounts'])}"

    return JSONResponse({
        'success': True,
        'message': message,
        'summary': summary,
    })


@app.get('/api/orders/{order_id}/refund-record')
async def get_order_refund_record(
    order_id: str,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """查询订单的退款记录（只读）。"""
    from app.db_manager import db_manager
    from utils.xianyu_seller_api import XianyuSellerAPI, SellerApiError

    user_cookies = db_manager.get_all_cookies(current_user['user_id'])
    cid, cookies_str = _resolve_order_cookie(order_id, user_cookies)

    api = XianyuSellerAPI(cid, cookies_str)
    try:
        data = await api.get_refund_record(order_id)
        return JSONResponse({'success': True, 'data': data})
    except SellerApiError as exc:
        return JSONResponse({'success': False, 'message': str(exc)}, status_code=502)
    finally:
        await api.close()


# ------------------------- 快捷短语 -------------------------

@app.get('/quick-phrases')
def list_quick_phrases(
    include_disabled: bool = Query(False),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """获取快捷短语列表，供人工客服快速插入常用话术。"""
    from app.db_manager import db_manager
    return {'success': True, 'data': db_manager.get_quick_phrases(include_disabled)}


@app.post('/quick-phrases')
def create_quick_phrase(
    title: str = Form(...),
    content: str = Form(...),
    category: str = Form('默认'),
    sort_order: int = Form(0),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    from app.db_manager import db_manager

    if not title.strip() or not content.strip():
        raise HTTPException(status_code=400, detail="标题和内容不能为空")

    phrase_id = db_manager.create_quick_phrase(
        title.strip(), content.strip(), category.strip() or '默认', sort_order
    )
    if not phrase_id:
        raise HTTPException(status_code=500, detail="新增快捷短语失败")
    log_with_user('info', f"新增快捷短语: {title}", current_user)
    return {'success': True, 'id': phrase_id}


@app.put('/quick-phrases/{phrase_id}')
def update_quick_phrase(
    phrase_id: int,
    title: Optional[str] = Form(None),
    content: Optional[str] = Form(None),
    category: Optional[str] = Form(None),
    sort_order: Optional[int] = Form(None),
    enabled: Optional[bool] = Form(None),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    from app.db_manager import db_manager

    ok = db_manager.update_quick_phrase(
        phrase_id, title=title, content=content, category=category,
        sort_order=sort_order, enabled=enabled,
    )
    if not ok:
        raise HTTPException(status_code=404, detail="快捷短语不存在或无字段更新")
    return {'success': True}


@app.delete('/quick-phrases/{phrase_id}')
def delete_quick_phrase(
    phrase_id: int,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    from app.db_manager import db_manager

    if not db_manager.delete_quick_phrase(phrase_id):
        raise HTTPException(status_code=404, detail="快捷短语不存在")
    log_with_user('info', f"删除快捷短语: {phrase_id}", current_user)
    return {'success': True}


@app.post('/quick-phrases/{phrase_id}/use')
def use_quick_phrase(
    phrase_id: int,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """记录一次使用，用于统计高频短语。"""
    from app.db_manager import db_manager
    return {'success': db_manager.increment_quick_phrase_usage(phrase_id)}


@app.post('/api/captcha/manual-session')
async def start_manual_captcha(
    cookie_id: str = Form(...),
    timeout: int = Form(300),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """开启人工验证会话。

    自动识别虽然能把滑块拖到目标位置，但服务端拦的是行为特征，重试再多次也
    不会通过。这里把滑块画面推到前端弹窗，用户手动完成后取回新 Cookie。

    只有账号处于风控状态时才允许调用 —— 非风控时闲鱼不会下发惩罚页，
    开会话只会白等一场，还可能因为多余请求把账号推向风控。
    """
    from app.db_manager import db_manager
    from utils.manual_captcha import open_manual_session
    from utils import risk_control

    user_cookies = db_manager.get_all_cookies(current_user['user_id'])
    if cookie_id not in user_cookies:
        raise HTTPException(status_code=404, detail="账号不存在或无权访问")

    guard = risk_control.registry.get(cookie_id)
    if not guard.is_blocked:
        raise HTTPException(
            status_code=409,
            detail="该账号当前不处于风控状态，无需人工验证",
        )

    timeout = max(60, min(int(timeout or 300), 900))
    result = await open_manual_session(
        cookie_id, user_cookies[cookie_id], timeout=timeout
    )

    if result['success']:
        # 拿到 x5sec 后必须清掉 x5secdata 等挑战标记，否则闲鱼会认为验证仍未完成，
        # 继续返回 FAIL_SYS_USER_VALIDATE —— 表现为"滑块过了但账号还是用不了"
        from utils.xianyu_utils import drop_stale_captcha_challenge
        cleaned_cookies = drop_stale_captcha_challenge(result['cookies_str'])
        if cleaned_cookies != result['cookies_str']:
            log_with_user('info', f"账号 {cookie_id} 已清除过期的验证挑战标记", current_user)
        result['cookies_str'] = cleaned_cookies

        # 保存新 Cookie 并解除风控熔断，让账号能立刻重连
        db_manager.save_cookie(cookie_id, result['cookies_str'])

        # 运行中的实例仍持有旧 Cookie，不同步会继续用旧值打接口并立刻再次熔断
        try:
            manager = cookie_manager.manager
            if manager is not None:
                manager.cookies[cookie_id] = result['cookies_str']
                instance = manager.instances.get(cookie_id)
                if instance is not None:
                    instance.cookies_str = result['cookies_str']
                    # 清掉失效令牌，强制下次请求重新获取
                    instance.current_token = None
                    log_with_user('info', f"账号 {cookie_id} 运行实例已同步新 Cookie", current_user)
        except Exception as exc:
            log_with_user('warning', f"同步实例 Cookie 失败: {exc}", current_user)

        try:
            from utils import risk_control

            risk_control.registry.get(cookie_id).reset()
        except Exception as exc:
            log_with_user('warning', f"重置风控状态失败: {exc}", current_user)
        log_with_user('info', f"账号 {cookie_id} 人工验证完成，已更新 Cookie", current_user)

    return JSONResponse({
        'success': result['success'],
        'message': result['message'],
        'session_id': result['session_id'],
    })


def classify_verification_event(detail: str, blocked: bool = False, reason: str = '') -> Tuple[str, str]:
    text = str(detail or '')
    # 已经处理成功的事件不算"当前需要验证"。这里的 detail 里含 processing_result，
    # 一条「滑块验证成功」的日志同样带着"滑块"二字，只按关键词匹配会让刚恢复的
    # 账号继续显示需要验证。
    if '成功' in text and ('滑块验证成功' in text or '验证成功' in text):
        return ('risk_control', reason or '闲鱼限制了当前账号请求') if blocked else ('none', '')
    if 'action=captcha' in text or 'slider_captcha' in text or '滑块' in text:
        return 'slider', 'Token 刷新被闲鱼重定向到滑块验证页'
    if '人脸' in text or 'face' in text.lower() or 'iframeRedirect' in text:
        return 'face', '闲鱼要求手机完成人脸或安全验证'
    if '扫码' in text or 'qr' in text.lower():
        return 'qr', '闲鱼要求重新扫码登录'
    if blocked:
        return 'risk_control', reason or '闲鱼限制了当前账号请求'
    return 'none', ''


@app.post('/api/risk-control/{cookie_id}/fresh-captcha-url')
async def get_fresh_captcha_url(
    cookie_id: str,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    """实时取一个新鲜的滑块验证链接。

    punish 链接里的 x5secdata 只能用一次、且约 1 小时后失效，过期后打开只会看到
    「抱歉，页面访问出现了问题」。风控日志里存的是历史链接，直接拿来用基本都是死链，
    所以这里凭账号 Cookie 重新请求一次 Token 接口 —— 仍被风控就拿到新链接，
    风控已解除则直接返回可用状态，让用户不必白跑一趟验证。
    """
    user_cookies = db_manager.get_all_cookies(current_user['user_id'])
    if cookie_id not in user_cookies:
        raise HTTPException(status_code=403, detail='无权限操作该账号')

    import aiohttp
    from app.config import API_ENDPOINTS
    from utils.xianyu_utils import trans_cookies, generate_sign, generate_device_id

    cookies_str = user_cookies[cookie_id]
    try:
        cookie_dict = trans_cookies(cookies_str)
    except ValueError:
        raise HTTPException(status_code=400, detail='账号 Cookie 为空，请重新扫码登录')

    token = (cookie_dict.get('_m_h5_tk') or '').split('_')[0]
    if not token:
        raise HTTPException(
            status_code=409,
            detail='账号 Cookie 缺少 _m_h5_tk 字段，请重新扫码登录',
        )

    timestamp = str(int(time.time() * 1000))
    device_id = generate_device_id(cookie_dict.get('unb', ''))
    data_value = (
        '{"appKey":"444e9908a51d1cb236a27862abc769c9","deviceId":"' + device_id + '"}'
    )
    params = {
        'jsv': '2.7.2', 'appKey': '34839810', 't': timestamp,
        'sign': generate_sign(timestamp, token, data_value),
        'v': '1.0', 'type': 'originaljson', 'accountSite': 'xianyu',
        'dataType': 'json', 'timeout': '20000',
        'api': 'mtop.taobao.idlemessage.pc.login.token',
        'sessionOption': 'AutoLoginOnly',
        'dangerouslySetWindvaneParams': '%5Bobject%20Object%5D',
        'smToken': 'token', 'queryToken': 'sm', 'sm': 'sm',
        'spm_cnt': 'a21ybx.im.0.0',
        'spm_pre': 'a21ybx.home.sidebar.1.4c053da6vYwnmf',
        'log_id': '4c053da6vYwnmf',
    }
    headers = {
        'accept': 'application/json',
        'content-type': 'application/x-www-form-urlencoded',
        'user-agent': (
            'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
            '(KHTML, like Gecko) Chrome/139.0.0.0 Safari/537.36'
        ),
        'referer': 'https://www.goofish.com/',
        'origin': 'https://www.goofish.com',
        'cookie': cookies_str,
    }

    try:
        async with aiohttp.ClientSession() as session:
            async with session.post(
                API_ENDPOINTS.get('token'),
                params=params, data={'data': data_value}, headers=headers,
                timeout=aiohttp.ClientTimeout(total=25),
            ) as response:
                payload = await response.json(content_type=None)
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f'请求闲鱼接口失败: {exc}') from exc

    data = payload.get('data') if isinstance(payload, dict) else None
    data = data if isinstance(data, dict) else {}

    if data.get('accessToken'):
        log_with_user('info', f'账号 {cookie_id} 风控已解除，无需验证', current_user)
        return {
            'success': True,
            'need_verify': False,
            'message': '该账号风控已解除，无需再做验证',
        }

    url = data.get('url') or ''
    if url and ('punish' in url or 'action=captcha' in url):
        log_with_user('info', f'账号 {cookie_id} 已取得新的验证链接', current_user)
        return {
            'success': True,
            'need_verify': True,
            'verification_url': url,
            'message': '已取得新的验证链接，请在浏览器中完成滑块',
        }

    raise HTTPException(
        status_code=502,
        detail=f"闲鱼未返回验证链接: {payload.get('ret') if isinstance(payload, dict) else payload}",
    )


@app.get('/api/risk-control/status')
def get_risk_control_status(current_user: Dict[str, Any] = Depends(get_current_user)):
    """查询各账号的风控熔断状态。

    命中平台风控后账号会进入冷却期，期间所有主动请求被拦截 —— 继续请求
    只会让风控持续更久。这个接口让用户能看到还要等多久。
    """
    from app.db_manager import db_manager
    from utils import risk_control

    user_cookies = db_manager.get_all_cookies(current_user['user_id'])
    snapshot = risk_control.registry.snapshot()
    manager = cookie_manager.manager

    accounts = []
    for cid in user_cookies.keys():
        state = snapshot.get(cid) or {
            'cookie_id': cid, 'blocked': False,
            'remaining_seconds': 0, 'consecutive_hits': 0, 'reason': '',
        }

        # 账号已经跑起来了就不该再提示需要验证。下面的判定依据是风控日志，
        # 而那是历史记录 —— 滑块过完、Token 已拿到之后，旧事件仍留在表里，
        # 只看日志会让恢复正常的账号一直挂着「拿不到令牌」的提示。
        running = False
        if manager is not None:
            try:
                running = manager.get_task_state(cid) == 'running'
            except Exception:
                running = False
        if running and not state.get('blocked'):
            accounts.append({
                **state,
                'verification_type': 'none',
                'verification_message': '',
                'verification_url': '',
                'latest_event': '',
                'latest_event_at': None,
            })
            continue

        recent_logs = db_manager.get_risk_control_logs(
            cookie_id=cid, limit=1, user_id=current_user['user_id']
        )
        latest = recent_logs[0] if recent_logs else {}

        # 风控日志是历史记录，过期的事件不能当成"当前仍需验证"。
        # punish 链接里的 x5secdata 本身也只有约 1 小时有效期，
        # 超过这个窗口的事件既没参考价值，链接也打不开了。
        event_is_fresh = True
        created_at = latest.get('created_at')
        if created_at:
            try:
                from datetime import datetime
                text = str(created_at).replace('T', ' ').split('.')[0]
                event_time = datetime.strptime(text, '%Y-%m-%d %H:%M:%S')
                event_is_fresh = (datetime.now() - event_time).total_seconds() <= 3600
            except Exception:
                event_is_fresh = True

        if not event_is_fresh and not state.get('blocked'):
            accounts.append({
                **state,
                'verification_type': 'none',
                'verification_message': '',
                'verification_url': '',
                'latest_event': latest.get('event_description') or '',
                'latest_event_at': created_at,
            })
            continue

        detail = ' '.join(str(latest.get(key) or '') for key in (
            'event_type', 'event_description', 'processing_result', 'error_message'
        ))
        verification_type, verification_message = classify_verification_event(
            detail, blocked=bool(state.get('blocked')), reason=state.get('reason') or ''
        )
        # 把惩罚页地址单独抽出来。服务器端用 CDP 拖滑块实测通过率很低
        # （风控查的是合并前子事件密度，自动化派发的事件每帧只有一个），
        # 让用户在自己浏览器里用真实鼠标过一次要可靠得多。
        verification_url = ''
        event_text = latest.get('event_description') or ''
        url_match = re.search(r'URL:\s*(\S+)', event_text)
        if url_match:
            candidate = url_match.group(1)
            if 'punish' in candidate or 'action=captcha' in candidate:
                verification_url = candidate

        state = {
            **state,
            'verification_type': verification_type,
            'verification_message': verification_message,
            'verification_url': verification_url,
            'latest_event': event_text,
            'latest_event_at': latest.get('created_at'),
        }
        accounts.append(state)

    return {
        'success': True,
        'blocked_count': sum(1 for a in accounts if a['blocked']),
        'accounts': accounts,
    }


@app.post('/api/items/polish')
async def polish_items(
    cookie_id: Optional[str] = Form(None),
    item_ids: Optional[str] = Form(None),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """手动擦亮商品。

    擦亮会把商品重新推到搜索和推荐前列，是平台提供的免费曝光手段。
    平台对每个商品每天的擦亮次数有限制，超出时接口返回业务错误，不影响其他商品。

    Args:
        item_ids: 逗号分隔的商品 ID；为空时擦亮该账号的全部商品。
    """
    from app.db_manager import db_manager
    from utils.item_polish import polish_account_items

    user_cookies = db_manager.get_all_cookies(current_user['user_id'])
    if cookie_id:
        if cookie_id not in user_cookies:
            raise HTTPException(status_code=404, detail="Cookie不存在或无权访问")
        user_cookies = {cookie_id: user_cookies[cookie_id]}

    targets = [i.strip() for i in str(item_ids or '').split(',') if i.strip()] or None
    summary = {'total': 0, 'success': 0, 'failed': 0, 'accounts': {}}

    for cid, cookies_str in user_cookies.items():
        if not cookies_str:
            continue
        try:
            res = await polish_account_items(cid, cookies_str, item_ids=targets)
        except Exception as e:
            log_with_user('error', f"账号 {cid} 商品擦亮异常: {e}", current_user)
            summary['accounts'][cid] = {'error': str(e)}
            continue

        summary['total'] += res['total']
        summary['success'] += res['success']
        summary['failed'] += res['failed']
        summary['accounts'][cid] = {
            'total': res['total'],
            'success': res['success'],
            'failed': res['failed'],
            'details': res['details'],
        }

    log_with_user(
        'info',
        f"商品擦亮完成: 共 {summary['total']} 个，成功 {summary['success']}，失败 {summary['failed']}",
        current_user,
    )
    return JSONResponse({
        'success': True,
        'message': f"擦亮完成: 共 {summary['total']} 个，成功 {summary['success']}，失败 {summary['failed']}",
        'summary': summary,
    })


@app.get('/api/orders/{order_id}/logistics')
async def get_order_logistics(
    order_id: str,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """查询订单的物流轨迹（只读）。

    项目原先没有物流能力，实物订单的运单状态只能人工去闲鱼查看。
    """
    from app.db_manager import db_manager
    from utils.xianyu_seller_api import (
        XianyuSellerAPI, SellerApiError, parse_logistics_trace,
    )

    user_cookies = db_manager.get_all_cookies(current_user['user_id'])
    cid, cookies_str = _resolve_order_cookie(order_id, user_cookies)

    api = XianyuSellerAPI(cid, cookies_str)
    try:
        data = parse_logistics_trace(await api.get_logistics_trace(order_id))
        return JSONResponse({'success': True, 'data': data})
    except SellerApiError as exc:
        return JSONResponse({'success': False, 'message': str(exc)}, status_code=502)
    finally:
        await api.close()


@app.get('/api/orders/{order_id}/consign-info')
async def get_order_consign_info(
    order_id: str,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """查询发货所需信息（只读）。

    返回规格、收货信息、支持的发货方式和是否必须填序列号。
    规格只能从这个接口拿 —— 订单列表的 itemInfoLines 是空的。
    """
    from app.db_manager import db_manager
    from utils.xianyu_seller_api import (
        XianyuSellerAPI, SellerApiError, parse_consign_render,
    )

    user_cookies = db_manager.get_all_cookies(current_user['user_id'])
    cid, cookies_str = _resolve_order_cookie(order_id, user_cookies)

    api = XianyuSellerAPI(cid, cookies_str)
    try:
        data = parse_consign_render(await api.get_consign_render(order_id))
        return JSONResponse({'success': True, 'data': data})
    except SellerApiError as exc:
        return JSONResponse({'success': False, 'message': str(exc)}, status_code=502)
    finally:
        await api.close()


@app.post('/api/orders/{order_id}/require-flower')
async def require_order_flower(
    order_id: str,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """向买家索要小红花。

    会给买家发送一条消息，需先开启 auto_flower_enabled 开关。
    """
    from app.db_manager import db_manager
    from utils.xianyu_seller_api import XianyuSellerAPI, SellerApiError

    user_cookies = db_manager.get_all_cookies(current_user['user_id'])
    cid, cookies_str = _resolve_order_cookie(order_id, user_cookies)

    # 开关按账号判：先解析出订单归属账号，再看该账号有没有开
    if not db_manager.get_buyer_interaction_settings(cid)['auto_flower_enabled']:
        raise HTTPException(
            status_code=403,
            detail="该账号未开启求花功能，请先在买家互动里为它启用（会向买家发送消息）",
        )

    api = XianyuSellerAPI(cid, cookies_str)
    try:
        data = await api.require_flower(order_id)
        log_with_user('info', f"订单 {order_id} 已发送求花请求", current_user)
        return JSONResponse({'success': True, 'data': data})
    except SellerApiError as exc:
        log_with_user('warning', f"订单 {order_id} 求花失败: {exc}", current_user)
        return JSONResponse({'success': False, 'message': str(exc)}, status_code=502)
    finally:
        await api.close()


@app.post('/api/orders/rate')
async def rate_orders(
    order_ids: str = Form(...),
    feedback: str = Form(...),
    rate: int = Form(1),
    anonymous: bool = Form(False),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """给买家提交评价，支持批量。

    评价提交后无法撤销，需先开启 auto_rate_enabled 开关。
    同一批订单必须属于同一个账号，因为接口按登录态区分卖家身份。
    """
    from app.db_manager import db_manager
    from utils.xianyu_seller_api import XianyuSellerAPI, SellerApiError

    id_list = [item.strip() for item in str(order_ids or '').split(',') if item.strip()]
    if not id_list:
        raise HTTPException(status_code=400, detail="订单列表不能为空")
    if not (feedback or '').strip():
        raise HTTPException(status_code=400, detail="评价内容不能为空")
    if rate not in (1, 0, -1):
        raise HTTPException(status_code=400, detail="评价等级只能是 1（好）/ 0（中）/ -1（差）")

    user_cookies = db_manager.get_all_cookies(current_user['user_id'])
    resolved = {_resolve_order_cookie(order_id, user_cookies)[0] for order_id in id_list}
    if len(resolved) > 1:
        raise HTTPException(status_code=400, detail="批量评价的订单必须属于同一个账号")

    cid = resolved.pop()

    # 开关按账号判：先解析出订单归属账号，再看该账号有没有开
    if not db_manager.get_buyer_interaction_settings(cid)['auto_rate_enabled']:
        raise HTTPException(
            status_code=403,
            detail="该账号未开启评价功能，请先在买家互动里为它启用（评价提交后不可撤销）",
        )

    api = XianyuSellerAPI(cid, user_cookies[cid])
    try:
        result = await api.create_rate(
            id_list, feedback=feedback, rate=rate, anonymous=anonymous
        )
        log_with_user(
            'info',
            f"已提交评价: 成功 {len(result['success_order_ids'])} 单，"
            f"失败 {len(result['fail_orders'])} 单",
            current_user,
        )
        return JSONResponse({'success': result['success'], 'data': result})
    except SellerApiError as exc:
        log_with_user('warning', f"提交评价失败: {exc}", current_user)
        return JSONResponse({'success': False, 'message': str(exc)}, status_code=502)
    finally:
        await api.close()


@app.get('/{path:path}', response_class=HTMLResponse)
async def catch_all_route(path: str):
    """
    Catch-all 路由：处理所有未匹配的 GET 请求
    如果是 API 请求，返回 404；否则返回前端 index.html
    """
    full_path = f'/{path}'
    root_segment = path.split('/', 1)[0]
    if root_segment in API_ROOTS:
        raise HTTPException(status_code=404, detail="Not Found")
    
    # 返回前端页面
    return await serve_frontend()


# 移除自动启动，由Start.py或手动启动
# if __name__ == "__main__":
#     uvicorn.run(app, host="0.0.0.0", port=8080)
